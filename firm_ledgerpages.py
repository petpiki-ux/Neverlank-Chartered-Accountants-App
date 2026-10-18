"""Routes for Firm Finance > Ledger: journals, chart of accounts, ledger per account, trial balance, reconciliation to the
statements and the period lock. The postings and rules are in firm_ledger.py."""
import io
from datetime import date, timedelta

from flask import render_template, redirect, url_for, request, flash, send_file
from flask_login import current_user

from extensions import db
from models import FirmAccount, FirmJournal, FirmJournalLine, FirmYearEnd, FIRM_ACCOUNT_TYPES
from finance import finance_required, _parse_date, _parse_float
from firm_finance import firmfs_bp, _pick_year_end, _year_ends
import firm_ledger as fl
import financials


def _statement_lines():
    """[(code, label)] of the statement lines an account can be mapped to, grouped for a select."""
    groups = {}
    for code, label, stmt, section, side, _cf in financials.FS_CATEGORIES:
        if code == "excluded":
            continue
        groups.setdefault(section, []).append((code, label))
    return groups


def _journal_accounts():
    """Accounts a journal line may use: active, and not a cash/bank account or one the app alone posts to."""
    fl.ensure_chart()
    out = []
    for a in FirmAccount.query.order_by(FirmAccount.code).all():
        if not a.is_active or a.fs_category == "cash" or a.system_key in fl.JOURNAL_REFUSED_KEYS:
            continue
        out.append(a)
    return out


def _page(template, active, **ctx):
    ye, yes = _pick_year_end()
    return render_template(template, active=active, ye=ye, year_ends=yes, lock=fl.current_lock(), **ctx)


# ------------------------------------------------------------------ journals

@firmfs_bp.route("/ledger")
@finance_required
def ledger_index():
    fl.ensure_chart()
    status = request.args.get("status", "")
    q = (request.args.get("q") or "").strip()
    d_from, d_to = _parse_date(request.args.get("from")), _parse_date(request.args.get("to"))
    query = FirmJournal.query
    if status in ("draft", "posted"):
        query = query.filter(FirmJournal.status == status)
    if d_from:
        query = query.filter(FirmJournal.journal_date >= d_from)
    if d_to:
        query = query.filter(FirmJournal.journal_date <= d_to)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(FirmJournal.number.ilike(like), FirmJournal.description.ilike(like), FirmJournal.reference.ilike(like)))
    journals = query.order_by(FirmJournal.journal_date.desc(), FirmJournal.id.desc()).limit(500).all()
    drafts = FirmJournal.query.filter_by(status="draft").count()
    return _page("firmfs/ledger_journals.html", "ledger_journals", journals=journals, status=status, q=q, d_from=d_from, d_to=d_to,
                 drafts=drafts)


def _read_journal_form(j):
    """Copy the posted form onto journal j (header and lines). Returns the list of problems with what was typed."""
    j.journal_date = _parse_date(request.form.get("journal_date"))
    j.description = " ".join((request.form.get("description") or "").split())[:250]
    j.reference = " ".join((request.form.get("reference") or "").split())[:120] or None
    accts = request.form.getlist("line_account")
    drs, crs = request.form.getlist("line_debit"), request.form.getlist("line_credit")
    narrs = request.form.getlist("line_narration")
    problems = []
    for ln in list(j.lines):
        j.lines.remove(ln)
    order = 0
    for i, raw_acct in enumerate(accts):
        dr_raw = (drs[i] if i < len(drs) else "").strip()
        cr_raw = (crs[i] if i < len(crs) else "").strip()
        if not raw_acct and not dr_raw and not cr_raw:
            continue                     # a blank row
        dr = _parse_float(dr_raw, 0.0) if dr_raw else 0.0
        cr = _parse_float(cr_raw, 0.0) if cr_raw else 0.0
        if dr is None or cr is None:
            problems.append(f"Line {order + 1}: amounts must be numbers.")
            dr, cr = 0.0, 0.0
        acct = FirmAccount.query.get(int(raw_acct)) if raw_acct.isdigit() else None
        line = FirmJournalLine(account_id=acct.id if acct else None, debit=fl.r2(dr), credit=fl.r2(cr),
                               narration=(narrs[i].strip()[:250] if i < len(narrs) else None) or None, order=order)
        line.account = acct
        j.lines.append(line)
        order += 1
    return problems


def _save_journal(j, is_new):
    """Shared by new/edit: save as a draft, or save and post. Returns a redirect, or None to re-show the form."""
    action = request.form.get("action", "save")
    problems = _read_journal_form(j)
    post = action == "post"
    problems += fl.journal_problems(j, for_post=post, year_ends=_year_ends())
    if not j.lines:
        problems.append("Add at least one line.")
    if problems:
        db.session.rollback()
        for p in dict.fromkeys(problems):
            flash(p, "danger")
        return None
    if is_new:
        j.number = fl.next_journal_number(j.journal_date)
        j.created_by_id = current_user.id
        j.status = "draft"
        db.session.add(j)
    if post:
        fl.post_journal(j, current_user)
    db.session.commit()
    flash(f"{j.number} posted to the ledger." if post else f"{j.number} saved as a draft.", "success")
    return redirect(url_for("firmfs.journal_view", journal_id=j.id))


def _form_state(j):
    """What the journal form shows: the typed values if a save was refused, else the journal's own."""
    return {
        "journal_date": request.form.get("journal_date") or (j.journal_date.isoformat() if j and j.journal_date else (request.args.get("date") or date.today().isoformat())),
        "description": request.form.get("description") if request.method == "POST" else (j.description if j else request.args.get("description", "")),
        "reference": request.form.get("reference") if request.method == "POST" else (j.reference if j and j.reference else ""),
    }


def _form_lines(j):
    if request.method == "POST":
        accts = request.form.getlist("line_account")
        drs, crs, nr = request.form.getlist("line_debit"), request.form.getlist("line_credit"), request.form.getlist("line_narration")
        rows = [{"account_id": int(a) if a.isdigit() else None, "debit": drs[i] if i < len(drs) else "", "credit": crs[i] if i < len(crs) else "",
                 "narration": nr[i] if i < len(nr) else ""} for i, a in enumerate(accts)]
    elif j is not None:
        rows = [{"account_id": l.account_id, "debit": f"{l.debit:.2f}" if l.debit else "", "credit": f"{l.credit:.2f}" if l.credit else "",
                 "narration": l.narration or ""} for l in j.lines]
    else:
        rows = []
        pre = request.args.get("acct", type=int)
        if pre:                                   # a line pre-filled from a link (e.g. from the clearing page)
            rows.append({"account_id": None, "debit": request.args.get("dr", ""), "credit": request.args.get("cr", ""),
                         "narration": request.args.get("narr", "")})
            rows.append({"account_id": pre, "debit": "", "credit": "", "narration": ""})
    while len(rows) < 4:
        rows.append({"account_id": None, "debit": "", "credit": "", "narration": ""})
    return rows


@firmfs_bp.route("/ledger/journal/new", methods=["GET", "POST"])
@finance_required
def journal_new():
    if request.method == "POST":
        got = _save_journal(FirmJournal(), True)
        if got:
            return got
    return _page("firmfs/ledger_journal_form.html", "ledger_new", j=None, form=_form_state(None), rows=_form_lines(None),
                 accounts=_journal_accounts())


@firmfs_bp.route("/ledger/journal/<int:journal_id>/edit", methods=["GET", "POST"])
@finance_required
def journal_edit(journal_id):
    j = FirmJournal.query.get_or_404(journal_id)
    if j.status != "draft":
        flash("A posted journal is never changed. Reverse it and post a corrected one.", "danger")
        return redirect(url_for("firmfs.journal_view", journal_id=j.id))
    if request.method == "POST":
        got = _save_journal(j, False)
        if got:
            return got
    return _page("firmfs/ledger_journal_form.html", "ledger_journals", j=j, form=_form_state(j), rows=_form_lines(j),
                 accounts=_journal_accounts())


@firmfs_bp.route("/ledger/journal/<int:journal_id>")
@finance_required
def journal_view(journal_id):
    j = FirmJournal.query.get_or_404(journal_id)
    problems = fl.journal_problems(j, for_post=True, year_ends=_year_ends()) if j.status == "draft" else []
    return _page("firmfs/ledger_journal.html", "ledger_journals", j=j, problems=problems, warnings=fl.journal_warnings(j),
                 can_reverse=(j.status == "posted" and j.reverses_id is None and j.reversed_by is None), today=date.today(),
                 cat_label={c[0]: c[1] for c in financials.FS_CATEGORIES})


@firmfs_bp.route("/ledger/journal/<int:journal_id>/post", methods=["POST"])
@finance_required
def journal_post(journal_id):
    j = FirmJournal.query.get_or_404(journal_id)
    if j.status != "draft":
        flash("That journal is already posted.", "info")
        return redirect(url_for("firmfs.journal_view", journal_id=j.id))
    problems = fl.journal_problems(j, for_post=True, year_ends=_year_ends())
    if problems:
        for p in dict.fromkeys(problems):
            flash(p, "danger")
        return redirect(url_for("firmfs.journal_view", journal_id=j.id))
    fl.post_journal(j, current_user)
    db.session.commit()
    flash(f"{j.number} posted to the ledger.", "success")
    return redirect(url_for("firmfs.journal_view", journal_id=j.id))


@firmfs_bp.route("/ledger/journal/<int:journal_id>/delete", methods=["POST"])
@finance_required
def journal_delete(journal_id):
    j = FirmJournal.query.get_or_404(journal_id)
    if j.status != "draft":
        flash("A posted journal cannot be deleted. Reverse it instead.", "danger")
        return redirect(url_for("firmfs.journal_view", journal_id=j.id))
    number = j.number
    db.session.delete(j)
    db.session.commit()
    flash(f"Draft {number} deleted.", "success")
    return redirect(url_for("firmfs.ledger_index"))


@firmfs_bp.route("/ledger/journal/<int:journal_id>/reverse", methods=["POST"])
@finance_required
def journal_reverse(journal_id):
    j = FirmJournal.query.get_or_404(journal_id)
    if j.status != "posted" or j.reverses_id is not None or j.reversed_by is not None:
        flash("Only a posted journal that has not been reversed (and is not itself a reversal) can be reversed.", "danger")
        return redirect(url_for("firmfs.journal_view", journal_id=j.id))
    on = _parse_date(request.form.get("reverse_date")) or date.today()
    rev = fl.make_reversal(j, on, current_user)
    problems = fl.journal_problems(rev, for_post=True, year_ends=_year_ends())
    if problems:
        for p in dict.fromkeys(problems):
            flash("The reversal cannot be posted: " + p, "danger")
        return redirect(url_for("firmfs.journal_view", journal_id=j.id))
    db.session.add(rev)
    db.session.commit()
    flash(f"{j.number} reversed by {rev.number}.", "success")
    return redirect(url_for("firmfs.journal_view", journal_id=rev.id))


@firmfs_bp.route("/ledger/journal/<int:journal_id>/copy", methods=["POST"])
@finance_required
def journal_copy(journal_id):
    j = FirmJournal.query.get_or_404(journal_id)
    today = date.today()
    new = FirmJournal(number=fl.next_journal_number(today), journal_date=today, description=j.description, reference=j.reference,
                      status="draft", created_by_id=current_user.id)
    for i, ln in enumerate(j.lines):
        new.lines.append(FirmJournalLine(account_id=ln.account_id, debit=ln.debit, credit=ln.credit, narration=ln.narration, order=i))
    db.session.add(new)
    db.session.commit()
    flash(f"Copied to draft {new.number}. Set the date and check the amounts.", "success")
    return redirect(url_for("firmfs.journal_edit", journal_id=new.id))


# ------------------------------------------------------------------ chart of accounts

@firmfs_bp.route("/ledger/accounts")
@finance_required
def ledger_accounts():
    fl.ensure_chart()
    accounts = FirmAccount.query.order_by(FirmAccount.code).all()
    used = {aid: n for aid, n in db.session.query(FirmJournalLine.account_id, db.func.count(FirmJournalLine.id)).group_by(FirmJournalLine.account_id).all()}
    edit = FirmAccount.query.get(request.args.get("edit", type=int) or 0)
    cat_label = {c[0]: c[1] for c in financials.FS_CATEGORIES}
    return _page("firmfs/ledger_accounts.html", "ledger_accounts", accounts=accounts, used=used, edit=edit, types=FIRM_ACCOUNT_TYPES,
                 groups=_statement_lines(), cat_label=cat_label, form=None)


@firmfs_bp.route("/ledger/accounts/save", methods=["POST"])
@finance_required
def ledger_account_save():
    acct = FirmAccount.query.get(request.form.get("id", type=int) or 0)
    code = (request.form.get("code") or "").strip()[:12]
    name = " ".join((request.form.get("name") or "").split())[:120]
    typ = request.form.get("account_type")
    cat = (request.form.get("fs_category") or "").strip() or fl.EXCLUDED
    note = (request.form.get("description") or "").strip()[:250] or None
    if not code or not name or typ not in FIRM_ACCOUNT_TYPES:
        flash("Enter the account code and name, and choose its type.", "danger")
    elif cat != fl.EXCLUDED and cat not in fl.CAT:
        flash("Choose a statement line from the list.", "danger")
    elif FirmAccount.query.filter(FirmAccount.code == code, FirmAccount.id != (acct.id if acct else 0)).first():
        flash(f"Code {code} is already used by another account.", "danger")
    else:
        if acct is None:
            acct = FirmAccount(order=FirmAccount.query.count() + 1000)
            db.session.add(acct)
        if acct.system_key:
            # the app posts to this account and the statements derive its figure: only the wording and number may change
            acct.code, acct.name, acct.description = code, name, note or acct.description
        else:
            changed = acct.id is not None and (acct.fs_category or fl.EXCLUDED) != cat and FirmJournalLine.query.filter_by(account_id=acct.id).first()
            acct.code, acct.name, acct.account_type, acct.fs_category, acct.description = code, name, typ, cat, note
            if changed:
                flash("Statement line changed. Posted journals on this account now feed the new line in the financial statements.", "info")
        db.session.commit()
        flash(f"Account {code} {name} saved.", "success")
        return redirect(url_for("firmfs.ledger_accounts"))
    return redirect(url_for("firmfs.ledger_accounts", edit=acct.id if acct and acct.id else None))


@firmfs_bp.route("/ledger/accounts/<int:account_id>/toggle", methods=["POST"])
@finance_required
def ledger_account_toggle(account_id):
    a = FirmAccount.query.get_or_404(account_id)
    if a.system_key and a.is_active:
        flash(f"{a.label} cannot be switched off: the app posts to it.", "danger")
    else:
        a.is_active = not a.is_active
        db.session.commit()
        flash(f"{a.label} switched {'on' if a.is_active else 'off'}.", "success")
    return redirect(url_for("firmfs.ledger_accounts"))


@firmfs_bp.route("/ledger/accounts/<int:account_id>/delete", methods=["POST"])
@finance_required
def ledger_account_delete(account_id):
    a = FirmAccount.query.get_or_404(account_id)
    if a.system_key:
        flash(f"{a.label} cannot be deleted: the app posts to it.", "danger")
    elif FirmJournalLine.query.filter_by(account_id=a.id).first():
        flash(f"{a.label} has journals on it. Switch it off instead of deleting it.", "danger")
    else:
        db.session.delete(a)
        db.session.commit()
        flash(f"Account {a.label} deleted.", "success")
    return redirect(url_for("firmfs.ledger_accounts"))


@firmfs_bp.route("/ledger/accounts/sync", methods=["POST"])
@finance_required
def ledger_accounts_sync():
    added = fl.ensure_chart()
    flash("Accounts added for new cash book accounts, expense categories or services." if added else "The chart is already up to date with the cash book, expense categories and services.", "success" if added else "info")
    return redirect(url_for("firmfs.ledger_accounts"))


# ------------------------------------------------------------------ ledger per account, trial balance, reconciliation

def _build():
    ye, yes = _pick_year_end()
    return fl.build_ledger(year_ends=yes), ye, yes


def _as_at(L, ye):
    """(as_at, period_start) from ?date= or the chosen year-end."""
    d = _parse_date(request.args.get("date"))
    if d:
        return d, L.period_start_of(d)
    if ye is not None:
        col = next((c for c in L.cols if c["ye_id"] == ye.id), None)
        return ye.period_end, col["period_start"] if col else L.period_start_of(ye.period_end)
    today = date.today()
    return today, date(today.year, 1, 1)


@firmfs_bp.route("/ledger/account")
@finance_required
def ledger_account():
    L, ye, yes = _build()
    as_at, ps = _as_at(L, ye)
    acct = L.accounts.get(request.args.get("account", type=int) or 0)
    start = _parse_date(request.args.get("from")) or ps
    end = _parse_date(request.args.get("to")) or as_at
    gl = fl.general_ledger(L, acct, start, end) if acct else None
    return render_template("firmfs/ledger_account.html", active="ledger_gl", ye=ye, year_ends=yes, lock=fl.current_lock(), L=L, acct=acct,
                           gl=gl, start=start, end=end, accounts=sorted(L.accounts.values(), key=lambda a: a.code), warnings=L.warnings)


@firmfs_bp.route("/ledger/trial-balance")
@finance_required
def ledger_trial_balance():
    L, ye, yes = _build()
    as_at, ps = _as_at(L, ye)
    tb = fl.trial_balance(L, as_at, ps)
    return render_template("firmfs/ledger_tb.html", active="ledger_tb", ye=ye, year_ends=yes, lock=fl.current_lock(), L=L, tb=tb, as_at=as_at,
                           custom=bool(request.args.get("date")), warnings=L.warnings,
                           cat_label={c[0]: c[1] for c in financials.FS_CATEGORIES})


@firmfs_bp.route("/ledger/trial-balance.xlsx")
@finance_required
def ledger_tb_export():
    L, ye, yes = _build()
    as_at, ps = _as_at(L, ye)
    tb = fl.trial_balance(L, as_at, ps)
    import firm_statements as fs
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    bold, head = Font(bold=True), PatternFill("solid", fgColor="E8E2D0")
    num = '#,##0.00;(#,##0.00);"-"'
    wb = Workbook()
    ws = wb.active
    ws.title = "Trial balance"
    ws.append([fs.FIRM_NAME]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Trial balance as at {as_at.strftime('%d %B %Y')} (income and expenses from {ps.strftime('%d %b %Y')})"]); ws["A2"].font = bold
    ws.append([])
    ws.append(["Code", "Account", "Type", "Statement line", "Debit (USD)", "Credit (USD)"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, head
    for r in tb["rows"]:
        if r.get("group_total") or r.get("heading"):
            continue
        a = r["account"]
        ws.append([a.code if a else "", a.name if a else r["label"], a.account_type if a else "", (financials.category_label(a.fs_category) if a and a.fs_category in fl.CAT else "Not in the statements") if a else "",
                   r["debit"] or None, r["credit"] or None])
    ws.append(["", "Total", "", "", tb["debit"], tb["credit"]])
    for c in ws[ws.max_row]:
        c.font = bold
    for row in ws.iter_rows(min_row=5, min_col=5, max_col=6):
        for c in row:
            c.number_format = num
    for col, w in zip("ABCDEF", (10, 58, 12, 34, 16, 16)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Postings")
    ws2.append(["Date", "Source", "Reference", "Description", "Account", "Narration", "Debit (USD)", "Credit (USD)"])
    for c in ws2[1]:
        c.font, c.fill = bold, head
    for ln in L.lines:
        if ps <= ln["date"] <= as_at:
            a = L.accounts[ln["account_id"]]
            ws2.append([ln["date"], ln["source"], ln["ref"], ln["description"], a.label, ln["narr"], ln["dr"] or None, ln["cr"] or None])
    for row in ws2.iter_rows(min_row=2):
        row[0].number_format = "dd mmm yyyy"
        row[6].number_format = num
        row[7].number_format = num
    for col, w in zip("ABCDEFGH", (13, 20, 16, 50, 46, 30, 15, 15)):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"trial_balance_{as_at.isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@firmfs_bp.route("/ledger/reconcile")
@finance_required
def ledger_reconcile():
    ye, yes = _pick_year_end()
    if ye is None:
        flash("Add a year-end first: the ledger is compared with the statements one year-end at a time.", "info")
        return redirect(url_for("firmfs.index"))
    L = fl.build_ledger(year_ends=yes)
    col = next(c for c in L.cols if c["ye_id"] == ye.id)
    rec = fl.reconcile(L, col)
    return render_template("firmfs/ledger_reconcile.html", active="ledger_rec", ye=ye, year_ends=yes, lock=fl.current_lock(), L=L, rec=rec,
                           col=col, warnings=L.warnings)


# ------------------------------------------------------------------ period lock

@firmfs_bp.route("/ledger/lock", methods=["GET", "POST"])
@finance_required
def ledger_lock():
    if request.method == "POST":
        if request.form.get("action") == "unlock":
            fl.set_lock(None, request.form.get("note"), current_user)
            db.session.commit()
            flash("The books are open again: journals can be dated in any period.", "success")
        else:
            d = _parse_date(request.form.get("locked_through"))
            if d is None:
                flash("Choose the date the books are closed through.", "danger")
                return redirect(url_for("firmfs.ledger_lock"))
            fl.set_lock(d, request.form.get("note"), current_user)
            db.session.commit()
            stuck = FirmJournal.query.filter(FirmJournal.status == "draft", FirmJournal.journal_date <= d).count()
            flash(f"The books are closed through {d.strftime('%d %b %Y')}." + (f" {stuck} draft journal(s) are dated on or before it and cannot be posted until re-dated." if stuck else ""), "success")
        return redirect(url_for("firmfs.ledger_lock"))
    ye, yes = _pick_year_end()
    drafts_locked = 0
    if fl.current_lock():
        drafts_locked = FirmJournal.query.filter(FirmJournal.status == "draft", FirmJournal.journal_date <= fl.current_lock()).count()
    return render_template("firmfs/ledger_lock.html", active="ledger_lock", ye=ye, year_ends=yes, lock=fl.current_lock(),
                           history=fl.lock_history(), drafts_locked=drafts_locked)


# ------------------------------------------------------------------ clear the clearing accounts

def _alloc_plan():
    yes = _year_ends()
    L = fl.build_ledger(year_ends=yes)
    return L, yes, fl.allocation_proposals(L)


@firmfs_bp.route("/ledger/clearing")
@finance_required
def ledger_clearing():
    ye, yes = _pick_year_end()
    L, yes, plan = _alloc_plan()
    susp = [(L.by_key[k], plan["balances"].get(k, 0.0)) for k in fl.SUSPENSE_KEYS if k in L.by_key]
    accounts = [a for a in sorted(L.accounts.values(), key=lambda a: a.code)
                if a.is_active and a.system_key not in fl.SUSPENSE_KEYS and a.system_key not in fl.JOURNAL_REFUSED_KEYS and a.fs_category != "cash"]
    return render_template("firmfs/ledger_clearing.html", active="ledger_clear", ye=ye, year_ends=yes, lock=fl.current_lock(), plan=plan,
                           susp=susp, accounts=accounts, warnings=L.warnings)


@firmfs_bp.route("/ledger/clearing/apply", methods=["POST"])
@finance_required
def ledger_clearing_apply():
    ye_id = request.form.get("ye_id", type=int)
    back = url_for("firmfs.ledger_clearing", ye=ye_id)
    chosen = set(request.form.getlist("apply"))
    if not chosen:
        flash("Tick at least one item.", "info")
        return redirect(back)
    post = request.form.get("mode") != "draft"
    L, yes, plan = _alloc_plan()                       # recomputed here: the page's amounts are never trusted
    made, errors, skipped = [], [], 0
    work = [(p, None) for p in plan["proposals"] if p["key"] in chosen]
    for c in plan["choose"]:
        if c["key"] in chosen:
            acct = request.form.get("choose_" + c["key"], type=int)
            if not acct:
                skipped += 1
                continue
            work.append((c, acct))
    for p, acct in work:
        if p.get("blocked"):
            errors.append(p["blocked"])
            continue
        j, problems = fl.apply_proposal(p, current_user, post=post, choice=acct)
        if problems:
            errors.extend(problems)
            continue
        made.append(j)
    if made:
        db.session.commit()
        flash(f"{len(made)} journal(s) {'posted' if post else 'saved as drafts'}: " + ", ".join(j.number for j in made) + ".", "success")
    else:
        db.session.rollback()
        flash("Nothing was posted: those items are no longer needed." if not errors and not skipped else "Nothing was posted.", "info")
    for e in dict.fromkeys(errors):
        flash(e, "danger")
    if skipped:
        flash(f"{skipped} item(s) left alone because no account was chosen.", "info")
    return redirect(back)
