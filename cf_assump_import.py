"""Firm-specific forecast assumptions in Excel: download a model's assumptions as a plain workbook, change them, upload them back.

Two kinds of workbook can be read:

* the workbook downloaded with "Assumptions workbook" (every row starts with a short key in column A, so nothing is guessed);
* any other workbook the firm keeps its standard assumptions in (a house template, last year's model, the live-formula model
  export): rows are matched by their label (for example "Tax rate", "Receivable days", "Capex"), and revenue / expense lines are matched
  by the names the model already uses. Rows that cannot be matched are listed so nothing is lost silently.

Nothing is changed by reading a workbook: the caller shows the differences and the preparer confirms them. Plain Python
(openpyxl only), no database and no Flask, so the tests can drive it directly.
"""
import copy
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL

import cf_engine as E

MAX_BYTES = 8 * 1024 * 1024
SHEET = "Assumptions"

NAVY = "1F3A5F"
F_HEAD = Font(bold=True, color="FFFFFF")
FILL_HEAD = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
FILL_SEC = PatternFill(start_color="E9EEF5", end_color="E9EEF5", fill_type="solid")
FILL_IN = PatternFill(start_color="FFF9E0", end_color="FFF9E0", fill_type="solid")
F_IN = Font(color="0000FF")
F_KEY = Font(color="9AA5B1", size=8)


class ImportError_(Exception):
    """A problem the user must fix (message is shown as is)."""


SCALARS = [  # key, label, percent?
    ("tax_rate", "Income tax rate (%)", True), ("dep_rate", "Depreciation rate (% of cost a year)", True),
    ("deposit_rate", "Interest earned on positive cash (% a year)", True), ("overdraft_rate", "Overdraft interest rate (% a year)", True),
]
VECTORS = [  # key, label, percent?
    ("cogs_pct", "Cost of sales (% of revenue)", True), ("capex", "Capital expenditure", False), ("rec_days", "Receivable days", False),
    ("inv_days", "Inventory days", False), ("pay_days", "Payable days", False), ("equity_issue", "Share capital issued", False),
    ("dividend_payout", "Dividend payout (% of profit)", True),
]
OPENING = [
    ("ppe_cost", "Property, plant and equipment - cost"), ("ppe_accdep", "Property, plant and equipment - accumulated depreciation"),
    ("inventory", "Inventory"), ("receivables", "Receivables"), ("other_assets", "Other assets"), ("cash", "Cash and bank"),
    ("payables", "Payables"), ("other_liabilities", "Other liabilities"), ("share_capital", "Share capital"),
    ("tax_loss_bf", "Tax losses brought forward"),
]
COVENANTS = [
    ("min_dscr", "Minimum debt service cover (x)"), ("min_interest_cover", "Minimum interest cover (x)"), ("max_debt_equity", "Maximum debt to equity (x)"),
    ("min_current_ratio", "Minimum current ratio (x)"), ("max_net_debt_ebitda", "Maximum net debt to EBITDA (x)"),
]
OPEN_LABEL = dict(OPENING)
COV_LABEL = dict(COVENANTS)
SCALAR_LABEL = {k: lab for k, lab, _ in SCALARS}
VECTOR_LABEL = {k: lab for k, lab, _ in VECTORS}
PERCENT_KEYS = {k for k, _, p in SCALARS if p} | {k for k, _, p in VECTORS if p}

# phrases (normalised) that identify a row in a workbook that is not the downloaded template
SYNONYMS = {
    "tax_rate": ["tax rate", "income tax rate", "corporate tax rate", "company tax rate"],
    "dep_rate": ["depreciation rate", "depreciation rate of gross cost", "depreciation rate of cost a year"],
    "deposit_rate": ["interest earned on positive cash", "deposit rate", "interest on cash", "interest income rate", "interest earned on cash"],
    "overdraft_rate": ["overdraft interest rate", "overdraft rate", "overdraft interest"],
    "years": ["forecast years", "number of years", "years"],
    "start_month": ["forecast start", "first forecast month", "start month", "model start"],
    "currency": ["currency"],
    "cogs_pct": ["cost of sales", "cost of sales of revenue", "cogs", "cogs of revenue", "cost of goods sold", "direct costs of revenue"],
    "capex": ["capital expenditure", "capex", "capital expenditure capex"],
    "rec_days": ["receivable days", "debtor days", "receivables days", "days sales outstanding", "dso"],
    "inv_days": ["inventory days", "stock days", "inventory holding days", "days inventory"],
    "pay_days": ["payable days", "creditor days", "payables days", "days payable", "dpo"],
    "equity_issue": ["share capital issued", "equity issue", "equity raised", "new equity", "share issue"],
    "dividend_payout": ["dividend payout", "dividend payout of profit", "payout ratio", "dividend payout ratio"],
    "op_ppe_cost": ["property plant and equipment cost", "ppe cost", "fixed assets cost", "opening ppe cost", "property plant and equipment at cost"],
    "op_ppe_accdep": ["accumulated depreciation", "ppe accumulated depreciation", "property plant and equipment accumulated depreciation"],
    "op_inventory": ["opening inventory", "inventory opening", "opening stock"],
    "op_receivables": ["opening receivables", "opening debtors", "trade receivables", "trade and other receivables"],
    "op_other_assets": ["other assets", "other current assets"],
    "op_cash": ["cash and bank", "cash and cash equivalents", "opening cash", "cash at bank"],
    "op_payables": ["opening payables", "opening creditors", "trade payables", "trade and other payables"],
    "op_other_liabilities": ["other liabilities", "other current liabilities"],
    "op_share_capital": ["share capital", "issued share capital", "opening share capital"],
    "op_tax_loss_bf": ["tax losses brought forward", "assessed loss", "assessed loss brought forward", "tax loss brought forward"],
    "cov_min_dscr": ["minimum debt service cover", "min dscr", "minimum dscr", "dscr covenant"],
    "cov_min_interest_cover": ["minimum interest cover", "min interest cover"],
    "cov_max_debt_equity": ["maximum debt to equity", "max debt equity", "maximum debt equity"],
    "cov_min_current_ratio": ["minimum current ratio", "min current ratio"],
    "cov_max_net_debt_ebitda": ["maximum net debt to ebitda", "max net debt ebitda"],
}
SYN_INDEX = {}
for _k, _phrases in SYNONYMS.items():
    for _p in _phrases:
        SYN_INDEX.setdefault(_p, _k)


# --------------------------------------------------------------------------- small helpers

def _norm(text):
    t = re.sub(r"\(.*?\)", " ", str(text or "").lower())
    t = re.sub(r"[^a-z0-9% ]+", " ", t)
    t = t.replace("%", " ")
    return " ".join(t.split())


def _name_key(text):
    return " ".join(str(text or "").lower().split())


def _num(cell, percent=False):
    """Cell -> float or None. A cell formatted as a percentage holds a fraction, so it is turned into percentage points."""
    v = cell.value if hasattr(cell, "value") else cell
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        if percent and hasattr(cell, "number_format") and "%" in (cell.number_format or ""):
            x *= 100.0
        return x
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace(" ", "")
        if not s:
            return None
        try:
            x = float(s.rstrip("%"))
        except ValueError:
            return None
        return x
    return None


def _text(cell):
    v = cell.value if hasattr(cell, "value") else cell
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m")
    return str(v).strip()


def fmt_val(v):
    if isinstance(v, (list, tuple)):
        return ", ".join(fmt_val(x) for x in v)
    if v is None:
        return "none"
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (int, float)):
        if abs(v - round(v)) < 0.005:
            return f"{v:,.0f}"
        return f"{v:,.2f}"
    return str(v)


# --------------------------------------------------------------------------- the downloadable workbook

def build_workbook(meta, a):
    """The model's current assumptions as a plain workbook (blue-on-yellow inputs, keys in column A)."""
    a = E.normalise(a)
    n = a["years"]
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 50
    for c in range(3, 4 + max(n * 2 + 4, 14)):
        ws.column_dimensions[CL(c)].width = 13
    row = 1
    ws.cell(row=row, column=2, value=f"Forecast assumptions - {meta.get('name', '')}").font = Font(bold=True, size=14, color=NAVY)
    row += 1
    ws.cell(row=row, column=2, value=f"{meta.get('client', '')} - {meta.get('number', '')} - amounts in {a['currency']}, excluding VAT").font = Font(color="6B7785")
    row += 1
    ws.cell(row=row, column=2, value=("Change the yellow cells and upload this file on the model's Assumptions tab. Column A holds the key the app reads: "
                                      "do not edit or delete it. Percentages are typed as numbers (25.75 means 25.75%). "
                                      "You can add or remove revenue streams, expense lines and loans by adding or deleting rows in their tables."))
    ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 48
    row += 1
    ws.cell(row=row, column=1, value="model").font = F_KEY
    ws.cell(row=row, column=2, value="Model number")
    ws.cell(row=row, column=3, value=meta.get("number", ""))
    row += 2

    def section(text):
        nonlocal row
        for c in range(1, 4 + max(n * 2 + 4, 14)):
            ws.cell(row=row, column=c).fill = FILL_SEC
        ws.cell(row=row, column=2, value=text).font = Font(bold=True, color=NAVY)
        row += 1

    def head(values, start=2):
        nonlocal row
        for i, v in enumerate(values):
            c = ws.cell(row=row, column=start + i, value=v)
            c.font, c.fill = F_HEAD, FILL_HEAD
            c.alignment = Alignment(horizontal="center" if i else "left", wrap_text=True, vertical="center")
        row += 1

    def put(key, label, values, fmt=None, start=3):
        nonlocal row
        ws.cell(row=row, column=1, value=key).font = F_KEY
        ws.cell(row=row, column=2, value=label)
        for i, v in enumerate(values):
            c = ws.cell(row=row, column=start + i, value=v)
            c.font, c.fill = F_IN, FILL_IN
            if fmt and isinstance(v, (int, float)):
                c.number_format = fmt
        row += 1

    section("General")
    put("years", "Forecast years (1 to 10)", [n], "0")
    put("start_month", "First forecast month (YYYY-MM)", [a["start_month"]])
    put("currency", "Currency", [a["currency"]])
    for k, lab, _ in SCALARS:
        put(k, lab, [a[k]], "0.00")
    row += 1

    section("Year 1 revenue phasing - share of the year's revenue in each month (%, totals 100)")
    start = E.parse_start_month(a["start_month"])
    head(["Month"] + [E.MONTH_ABBR[E.add_months(start, m).month - 1] for m in range(12)])
    put("seasonality", "Share of year 1 revenue (%)", a["seasonality"], "0.00")
    row += 1

    section("Yearly drivers")
    head([""] + [E.year_label(a, y) for y in range(1, n + 1)])
    for k, lab, _ in VECTORS:
        put(k, lab, a[k], "#,##0.00" if k in PERCENT_KEYS or k.endswith("days") else "#,##0")
    row += 1

    section("Revenue streams")
    head(["Name", "Year 1 revenue"] + [f"Growth % {E.year_label(a, y)}" for y in range(2, n + 1)])
    for s in a["revenue"]:
        put("rev", s["name"], [s["y1"]] + s["growth"][1:], "#,##0.00")
        ws.cell(row=row - 1, column=2).fill = FILL_IN
        ws.cell(row=row - 1, column=2).font = F_IN
        ws.cell(row=row - 1, column=3).number_format = "#,##0"
    row += 1

    section("Operating expenses")
    head(["Name", "Type (fixed / pct)", "Paid through creditors (Yes / No)", "Year 1 amount"] + [f"Growth % {E.year_label(a, y)}" for y in range(2, n + 1)]
         + [f"% of revenue {E.year_label(a, y)}" for y in range(1, n + 1)])
    for o in a["opex"]:
        put("ox", o["name"], [o["kind"], "Yes" if o["creditors"] else "No", o["y1"]] + o["growth"][1:] + o["pct"], "#,##0.00")
        ws.cell(row=row - 1, column=2).fill = FILL_IN
        ws.cell(row=row - 1, column=2).font = F_IN
        ws.cell(row=row - 1, column=6).number_format = "#,##0"
    row += 1

    section("Opening balance sheet (the position the forecast starts from)")
    for k, lab in OPENING:
        put(f"op_{k}", lab, [a["opening"][k]], "#,##0")
    row += 1

    section("Loans")
    head(["Name", "Opening balance", "New drawdown", "Drawdown year", "Interest rate %", "Repayment (straight / bullet)", "Repayment years", "Grace years", "Maturity year"])
    for L in a["loans"]:
        put("ln", L["name"], [L["opening_balance"], L["drawdown_amount"], L["drawdown_year"], L["rate"], L["repay_type"], L["repay_years"],
                              L["grace_years"], L["maturity_year"]], "#,##0.00")
        ws.cell(row=row - 1, column=2).fill = FILL_IN
        ws.cell(row=row - 1, column=2).font = F_IN
    row += 1

    section("Covenants (leave blank for none)")
    for k, lab in COVENANTS:
        put(f"cov_{k}", lab, [a["covenants"].get(k)], "0.00")
    ws.freeze_panes = "C5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- reading

class _Reader:
    def __init__(self, base):
        self.base = E.normalise(base)
        self.a = copy.deepcopy(self.base)
        self.n = self.base["years"]
        self.warnings = []
        self.notes = []        # how rows were matched when not by key
        self.unmatched = []
        self.sources = {}      # item -> where it came from
        self.rev_rows, self.ox_rows, self.ln_rows = [], [], []
        self.matched = 0

    def warn(self, text):
        if text not in self.warnings:
            self.warnings.append(text)

    # ---- scalars / vectors
    def set_scalar(self, key, cell, where):
        if key in SCALAR_LABEL:
            x = _num(cell, True)
            if x is None:
                return self.warn(f"{where}: {SCALAR_LABEL[key]} is not a number, so it was left as it is.")
            self.a[key] = x
        elif key in OPEN_LABEL.keys() or key.startswith("op_"):
            k = key[3:]
            x = _num(cell)
            if x is None:
                return self.warn(f"{where}: {OPEN_LABEL.get(k, k)} is not a number, so it was left as it is.")
            self.a["opening"][k] = x
        elif key.startswith("cov_"):
            k = key[4:]
            raw = cell.value if hasattr(cell, "value") else cell
            self.a["covenants"][k] = None if raw in (None, "") else (_num(cell) if _num(cell) is not None else self.a["covenants"].get(k))
        self.sources[key] = where
        self.matched += 1

    def set_vector(self, key, cells, where, single_to_all=False):
        pct = key in PERCENT_KEYS
        vals = [_num(c, pct) for c in cells]
        cur = list(self.a[key])
        got = [v for v in vals if v is not None]
        if not got:
            return self.warn(f"{where}: no numbers found for {VECTOR_LABEL.get(key, key)}, so it was left as it is.")
        if len(got) == 1 and single_to_all:
            cur = [got[0]] * self.n
            self.notes.append(f"{VECTOR_LABEL[key]}: one value ({fmt_val(got[0])}) found at {where}; it was applied to every year.")
        else:
            for i in range(self.n):
                if i < len(vals) and vals[i] is not None:
                    cur[i] = vals[i]
            if len([v for v in vals if v is not None]) < self.n:
                self.notes.append(f"{VECTOR_LABEL[key]}: only {len(got)} of {self.n} years had a figure at {where}; the other years were left as they are.")
        self.a[key] = cur
        self.sources[key] = where
        self.matched += 1


def _row_cells(ws, r, c0, c1):
    return [ws.cell(row=r, column=c) for c in range(c0, c1 + 1)]


def _read_template(ws, rd, wsname):
    """Rows keyed in column A. Returns True when the sheet is a template sheet."""
    keyed = False
    maxr, maxc = ws.max_row, ws.max_column
    # the number of years decides the width of every yearly row, so read it first
    for r in range(1, maxr + 1):
        if _text(ws.cell(row=r, column=1)).lower() == "years":
            x = _num(ws.cell(row=r, column=3))
            if x is not None:
                n_new = E.clamp_years(x, rd.n)
                if n_new != rd.n:
                    rd.notes.append(f"The workbook has {n_new} forecast years; the model has {rd.n}. The model was resized to {n_new} years.")
                    rd.a = E.normalise(dict(rd.a, years=n_new))
                    rd.n = n_new
            break
    n = rd.n
    for r in range(1, maxr + 1):
        key = _text(ws.cell(row=r, column=1)).lower()
        if not key:
            continue
        where = f"{wsname}!row {r}"
        label = _text(ws.cell(row=r, column=2))
        if key == "model":
            rd.model_number = _text(ws.cell(row=r, column=3))
            keyed = True
        elif key == "years":
            keyed = True
            rd.sources["years"] = where
        elif key == "start_month":
            keyed = True
            raw = ws.cell(row=r, column=3).value
            txt = _text(ws.cell(row=r, column=3))
            if isinstance(raw, (datetime, date)) or re.match(r"^\d{4}-\d{2}", txt):
                rd.a["start_month"] = E.parse_start_month(raw if isinstance(raw, (datetime, date)) else txt).isoformat()[:7]
                rd.sources[key] = where
                rd.matched += 1
            else:
                rd.warn(f"{where}: the first forecast month must look like 2026-10, so it was left as it is.")
        elif key == "currency":
            keyed = True
            txt = _text(ws.cell(row=r, column=3)).upper()
            if txt:
                rd.a["currency"] = txt[:6]
                rd.sources[key] = where
                rd.matched += 1
        elif key in SCALAR_LABEL or key.startswith("op_") or key.startswith("cov_"):
            keyed = True
            if key.startswith("op_") and key[3:] not in OPEN_LABEL:
                continue
            if key.startswith("cov_") and key[4:] not in COV_LABEL:
                continue
            rd.set_scalar(key, ws.cell(row=r, column=3), where)
        elif key == "seasonality":
            keyed = True
            vals = [_num(c) for c in _row_cells(ws, r, 3, 14)]
            if all(v is not None for v in vals) and sum(vals) > 0:
                rd.a["seasonality"] = vals
                rd.sources[key] = where
                rd.matched += 1
                if abs(sum(vals) - 100) > 0.5:
                    rd.notes.append(f"The monthly revenue phasing adds up to {sum(vals):,.1f}, not 100, so it was scaled to 100%.")
            else:
                rd.warn(f"{where}: the 12 monthly phasing figures must all be numbers, so the phasing was left as it is.")
        elif key in VECTOR_LABEL:
            keyed = True
            rd.set_vector(key, _row_cells(ws, r, 3, 2 + n), where)
        elif key == "rev":
            keyed = True
            rd.rev_rows.append((r, label, [ws.cell(row=r, column=c) for c in range(3, 3 + n)]))
        elif key == "ox":
            keyed = True
            rd.ox_rows.append((r, label, [ws.cell(row=r, column=c) for c in range(3, 3 + 3 + (n - 1) + n)]))
        elif key == "ln":
            keyed = True
            rd.ln_rows.append((r, label, [ws.cell(row=r, column=c) for c in range(3, 11)]))
    return keyed


def _apply_tables(rd, wsname):
    n = rd.n
    base = rd.a
    old_rev = {_name_key(s["name"]): s for s in rd.base["revenue"]}
    old_ox = {_name_key(o["name"]): o for o in rd.base["opex"]}
    old_ln = {_name_key(L["name"]): L for L in rd.base["loans"]}

    if rd.rev_rows:
        out = []
        for r, name, cells in rd.rev_rows:
            if not name:
                continue
            prev = old_rev.get(_name_key(name))
            y1 = _num(cells[0])
            if y1 is None:
                rd.warn(f"{wsname}!row {r}: Year 1 revenue for \"{name}\" is not a number, so the line was skipped." if not prev else
                        f"{wsname}!row {r}: Year 1 revenue for \"{name}\" is not a number, so the current figure was kept.")
                if not prev:
                    continue
                y1 = prev["y1"]
            g = [0.0] + [(_num(cells[i], True) if _num(cells[i], True) is not None else (prev["growth"][i] if prev and i < len(prev["growth"]) else 0.0))
                         for i in range(1, n)]
            if y1 < 0:
                rd.warn(f"{wsname}!row {r}: negative revenue for \"{name}\" was set to zero.")
            out.append({"name": name, "y1": max(0.0, y1), "growth": g})
        if out:
            base["revenue"] = out
            rd.matched += len(out)

    if rd.ox_rows:
        out = []
        for r, name, cells in rd.ox_rows:
            if not name:
                continue
            prev = old_ox.get(_name_key(name))
            kind_txt = _text(cells[0]).lower()
            kind = "pct" if kind_txt.startswith("p") or "%" in kind_txt else ("fixed" if kind_txt.startswith("f") else (prev["kind"] if prev else "fixed"))
            cred_txt = _text(cells[1]).lower()
            cred = cred_txt not in ("no", "n", "false", "0") if cred_txt else (prev["creditors"] if prev else True)
            y1 = _num(cells[2])
            if y1 is None:
                y1 = prev["y1"] if prev else 0.0
            growth = [0.0] + [(_num(cells[2 + i], True) if _num(cells[2 + i], True) is not None else (prev["growth"][i] if prev and i < len(prev["growth"]) else 0.0))
                              for i in range(1, n)]
            pct = [(_num(cells[n + 2 + i], True) if _num(cells[n + 2 + i], True) is not None else (prev["pct"][i] if prev and i < len(prev["pct"]) else 0.0))
                   for i in range(n)]
            out.append({"name": name, "kind": kind, "y1": max(0.0, y1), "growth": growth, "pct": pct, "creditors": cred})
        base["opex"] = out
        rd.matched += len(out)

    if rd.ln_rows:
        out = []
        for r, name, cells in rd.ln_rows:
            if not name:
                continue
            prev = old_ln.get(_name_key(name))
            nums = [_num(c) for c in cells]
            pick = lambda i, key, dflt: nums[i] if nums[i] is not None else (prev[key] if prev else dflt)
            rt = _text(cells[4]).lower()
            out.append({"name": name, "opening_balance": pick(0, "opening_balance", 0.0), "drawdown_amount": pick(1, "drawdown_amount", 0.0),
                        "drawdown_year": pick(2, "drawdown_year", 1), "rate": pick(3, "rate", 12.0),
                        "repay_type": "bullet" if rt.startswith("b") else ("straight" if rt.startswith("s") else (prev["repay_type"] if prev else "straight")),
                        "repay_years": pick(5, "repay_years", 5), "grace_years": pick(6, "grace_years", 0), "maturity_year": pick(7, "maturity_year", n)})
        base["loans"] = out
        rd.matched += len(out)


def _read_by_label(wb, rd):
    """A workbook that is not the downloaded template: match rows by label."""
    n = rd.n
    old_rev = {_name_key(s["name"]): s for s in rd.base["revenue"]}
    old_ox = {_name_key(o["name"]): o for o in rd.base["opex"]}
    seen = set()
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and len(wb.worksheets) > 1:
            continue
        for r in range(1, min(ws.max_row, 600) + 1):
            label_cell = None
            for c in range(1, 5):
                cell = ws.cell(row=r, column=c)
                if isinstance(cell.value, str) and cell.value.strip() and _num(cell) is None:
                    label_cell = (c, cell)
                    break
            if not label_cell:
                continue
            c0, cell = label_cell
            label = cell.value.strip()
            if label.startswith("="):
                continue
            nums = []
            for c in range(c0 + 1, min(ws.max_column, c0 + 24) + 1):
                v = ws.cell(row=r, column=c)
                if isinstance(v.value, (int, float)) and not isinstance(v.value, bool):
                    nums.append(v)
            where = f"{ws.title}!row {r}"
            key = SYN_INDEX.get(_norm(label))
            lname = _name_key(label)
            if lname in old_rev or lname in old_ox:
                if not nums:
                    continue
                is_rev = lname in old_rev
                prev = old_rev[lname] if is_rev else old_ox[lname]
                vals = [abs(float(v.value)) for v in nums[:n]]
                if ("rev", lname) in seen or ("ox", lname) in seen:
                    continue
                seen.add(("rev" if is_rev else "ox", lname))
                target = next(x for x in (rd.a["revenue"] if is_rev else rd.a["opex"]) if _name_key(x["name"]) == lname)
                if target.get("kind") == "pct":
                    rd.notes.append(f"\"{label}\" is a percentage-of-revenue expense in the model, so the workbook row at {where} was not used.")
                    continue
                target["y1"] = vals[0]
                if len(vals) == n and n > 1:
                    g = [0.0]
                    for i in range(1, n):
                        g.append(round((vals[i] / vals[i - 1] - 1) * 100.0, 4) if vals[i - 1] > 0 else prev["growth"][i])
                    target["growth"] = g
                    rd.notes.append(f"\"{label}\" ({where}): the yearly amounts were read as year 1 plus the growth between years.")
                else:
                    rd.notes.append(f"\"{label}\" ({where}): only the year 1 amount was used; growth was left as it is.")
                rd.sources[("rev_" if is_rev else "ox_") + lname] = where
                rd.matched += 1
                continue
            if not key or key in seen:
                if not key and nums:
                    rd.unmatched.append(f"{where}: {label}")
                continue
            if not nums and key not in ("start_month", "currency"):
                continue
            seen.add(key)
            if key == "years":
                x = E.clamp_years(nums[0].value, n)
                if x != n:
                    rd.warn(f"{where}: the workbook shows {x} forecast years but the model has {n}; the number of years was not changed (resize the model on the Overview tab).")
            elif key == "start_month":
                raw = ws.cell(row=r, column=c0 + 1).value
                if isinstance(raw, (datetime, date)) or (isinstance(raw, str) and re.match(r"^\d{4}-\d{2}", raw.strip())):
                    rd.a["start_month"] = E.parse_start_month(raw if isinstance(raw, (datetime, date)) else raw.strip()).isoformat()[:7]
                    rd.sources[key] = where
                    rd.matched += 1
            elif key == "currency":
                raw = ws.cell(row=r, column=c0 + 1).value
                if isinstance(raw, str) and raw.strip():
                    rd.a["currency"] = raw.strip().upper()[:6]
                    rd.sources[key] = where
                    rd.matched += 1
            elif key in SCALAR_LABEL or key.startswith("op_") or key.startswith("cov_"):
                rd.set_scalar(key, nums[0], where)
                if key in PERCENT_KEYS and 0 < rd.a[key] <= 1.0:
                    rd.warn(f"{where}: {SCALAR_LABEL[key]} was read as {rd.a[key]:g}%. If it is a decimal fraction (0.2575 for 25.75%), format the cell as a percentage or type 25.75.")
            elif key in VECTOR_LABEL:
                rd.set_vector(key, nums[:n], where, single_to_all=True)
                if key in PERCENT_KEYS and 0 < max(rd.a[key]) <= 1.0:
                    rd.warn(f"{where}: {VECTOR_LABEL[key]} was read as {max(rd.a[key]):g}% or less. If the figures are decimal fractions (0.55 for 55%), format the cells as percentages or type 55.")
                rd.notes.append(f"\"{label}\" at {where} was read as {VECTOR_LABEL[key]}.")
    return


def read_workbook(raw, base):
    """Parse an uploaded workbook against the model's current assumptions.
    Returns {"assumptions": new (normalised), "changes": [...], "notes", "warnings", "unmatched", "mode", "matched", "model_number"}."""
    if len(raw) > MAX_BYTES:
        raise ImportError_("That file is too large to be an assumptions workbook (limit 8 MB).")
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        raise ImportError_("That file could not be opened. Upload an .xlsx workbook (for example the one downloaded with 'Assumptions workbook').")
    rd = _Reader(base)
    rd.model_number = ""
    mode = "template"
    keyed = False
    sheet = wb[SHEET] if SHEET in wb.sheetnames else None
    if sheet is not None:
        keyed = _read_template(sheet, rd, SHEET)
        if keyed:
            _apply_tables(rd, SHEET)
    if not keyed:
        mode = "labels"
        rd = _Reader(base)
        rd.model_number = ""
        _read_by_label(wb, rd)
    if rd.matched == 0:
        raise ImportError_("No assumptions were recognised in that workbook. Use the file downloaded with 'Assumptions workbook', or a workbook whose rows are labelled "
                           "(for example 'Tax rate', 'Receivable days', 'Capex' or the names of the model's revenue and expense lines).")
    # formulas without a stored result are read as blanks - tell the user
    try:
        wb_f = load_workbook(io.BytesIO(raw), data_only=False)
        blank = 0
        for ws in wb_f.worksheets:
            ws_v = wb[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and ws_v[c.coordinate].value is None:
                        blank += 1
        if blank:
            rd.warn(f"{blank} cell(s) contain formulas without a saved result (the file was not saved by Excel), so they were read as blank. Open and save the file in Excel, then upload it again.")
    except Exception:
        pass
    new = E.normalise(rd.a)
    old = E.normalise(base)
    return {"assumptions": new, "changes": diff(old, new), "notes": rd.notes, "warnings": rd.warnings, "unmatched": rd.unmatched[:40],
            "mode": mode, "matched": rd.matched, "model_number": getattr(rd, "model_number", ""), "sources": {str(k): v for k, v in rd.sources.items()}}


# --------------------------------------------------------------------------- differences

def _eq(a, b):
    if isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)):
        return len(a) == len(b) and all(_eq(x, y) for x, y in zip(a, b))
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) < 0.005
    return a == b


def diff(old, new):
    """What differs between two normalised assumption sets, as [{group, item, old, new}] with display strings."""
    out = []

    def add(group, item, o, n):
        if not _eq(o, n):
            out.append({"group": group, "item": item, "old": fmt_val(o), "new": fmt_val(n)})

    add("General", "Forecast years", old["years"], new["years"])
    add("General", "First forecast month", old["start_month"], new["start_month"])
    add("General", "Currency", old["currency"], new["currency"])
    for k, lab, _ in SCALARS:
        add("General", lab, old[k], new[k])
    add("Revenue phasing", "Share of year 1 revenue by month (%)", [round(x, 2) for x in old["seasonality"]], [round(x, 2) for x in new["seasonality"]])
    for k, lab, _ in VECTORS:
        add("Yearly drivers", lab + " (by year)", old[k], new[k])

    def by_name(lst):
        return {_name_key(x["name"]): x for x in lst}

    for group, sec, fields in (("Revenue streams", "revenue", [("y1", "Year 1 revenue"), ("growth", "Growth % (years 2+)")]),
                               ("Operating expenses", "opex", [("kind", "Type"), ("creditors", "Paid through creditors"), ("y1", "Year 1 amount"),
                                                               ("growth", "Growth % (years 2+)"), ("pct", "% of revenue (by year)")]),
                               ("Loans", "loans", [("opening_balance", "Opening balance"), ("drawdown_amount", "New drawdown"), ("drawdown_year", "Drawdown year"),
                                                   ("rate", "Interest rate %"), ("repay_type", "Repayment type"), ("repay_years", "Repayment years"),
                                                   ("grace_years", "Grace years"), ("maturity_year", "Maturity year")])):
        o, nw = by_name(old[sec]), by_name(new[sec])
        for k, x in nw.items():
            if k not in o:
                out.append({"group": group, "item": f"{x['name']} (new line)", "old": "not in the model", "new": "added"})
                continue
            for f_, lab in fields:
                ov, nv = o[k][f_], x[f_]
                if f_ == "growth":
                    ov, nv = ov[1:], nv[1:]
                add(group, f"{x['name']} - {lab}", ov, nv)
        for k, x in o.items():
            if k not in nw:
                out.append({"group": group, "item": f"{x['name']} (removed)", "old": "in the model", "new": "removed"})
    for k, lab in OPENING:
        add("Opening balance sheet", lab, old["opening"][k], new["opening"][k])
    for k, lab in COVENANTS:
        add("Covenants", lab, old["covenants"].get(k), new["covenants"].get(k))
    return out
