"""The firm's document design for Word and Excel files - the same look as the
printed invoice: a navy letterhead band with the logo, firm name and gold
strapline, a details strip, navy-header tables, gold headings and a navy
footer carrying the firm name, website and page numbers.

One place to change the look of every generated .docx / .xlsx. Used by
workpapers.py (working papers, financial statements, reports, checklists),
payroll_documents.py, whistleblower.py and training.py, and by the tool that
restyles the Word / Excel template library.
"""
import os

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Mm

NAVY = "0F2942"
NAVY_DARK = "0A1C2E"
GOLD = "B8934A"
GOLD_LIGHT = "D4AF6A"
INK = "1F2933"
MUTED = "6B7684"
LINE = "E3E8EE"
TINT = "F6F8FB"
SOFT_TEXT = "CDD8E4"

LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "img", "neverlank-logo.png")
BODY_FONT = "Calibri"
HEAD_FONT = "Georgia"


def rgb(hex_colour):
    return RGBColor.from_string(hex_colour)


def firm_info():
    """The firm's letterhead details (same dictionary the printed invoice uses)."""
    from invoicing import FIRM_LETTERHEAD
    return FIRM_LETTERHEAD


# ------------------------------------------------------------------ the Managing Partner's e-signature
# The image lives in brand_assets/ (NOT under static/, so it is never served as a public file). Every place that
# signs for the firm goes through these helpers, so replacing brand_assets/managing_partner_signature.png changes
# every document at once.

SIGNATURE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "brand_assets", "managing_partner_signature.png")
MP_NAME = "Petros Piki"
MP_CREDENTIALS = "CA(Z), MSc, RPA, RTA"
MP_TITLE = "Managing Partner"
# The signing block under an Independent Auditor's / Reviewer's Report (workpapers._audit_opinion_section).
AUDITOR_SIGNOFF_NAME = "Petros Piki, CA(Z), MSc, RPA"
AUDITOR_PAAB_CERT_NO = "0625"
AUDITOR_ICAZ_NO = "254069"
FIRM_PAAB_NO = "Z10110"
_sig_cache = {"mtime": None, "bytes": None, "uri": None}


def signature_bytes():
    """The signature image as PNG bytes, or None when the file is missing (documents then simply carry no signature)."""
    import base64
    try:
        mtime = os.path.getmtime(SIGNATURE_PATH)
    except OSError:
        return None
    if _sig_cache["mtime"] != mtime:
        with open(SIGNATURE_PATH, "rb") as fh:
            data = fh.read()
        _sig_cache.update(mtime=mtime, bytes=data, uri="data:image/png;base64," + base64.b64encode(data).decode("ascii"))
    return _sig_cache["bytes"]


def signature_data_uri():
    """The signature as a data: URI for the printed HTML documents (an empty string when there is none)."""
    return _sig_cache["uri"] if signature_bytes() else ""


def content_key(doc):
    """A short fingerprint of what the Managing Partner signs: the document number and amount (invoices, quotations, credit notes,
    receipts, advance receipts), or the subject, wording and signatory of a letter. If it changes after signing, the signature
    stops showing until the document is signed again."""
    import hashlib
    if doc is None:
        return ""
    cls = type(doc).__name__
    if cls == "ZimraReply":
        parts = [doc.subject, doc.body, doc.signatory, doc.addressee, doc.letter_date]
    elif cls == "WorkpaperNarrative":
        parts = [doc.kind, doc.body]
    else:
        number = next((getattr(doc, a) for a in ("invoice_number", "quote_number", "credit_note_number", "receipt_number", "advance_number", "number")
                       if getattr(doc, a, None)), "")
        amount = getattr(doc, "total", None)
        if amount is None:
            amount = getattr(doc, "amount", None)
        parts = [number, f"{amount:.2f}" if amount is not None else ""]
    return hashlib.sha1("|".join(str(x or "") for x in parts).encode("utf-8")).hexdigest()


def signature_valid(doc):
    """True when the Managing Partner has signed this document and its content is still what he signed."""
    return bool(doc is not None and getattr(doc, "mp_signed_at", None) and getattr(doc, "mp_signed_key", None)
                and doc.mp_signed_key == content_key(doc) and signature_bytes())


def signature_uri_for(doc):
    """The signature as a data: URI for a printed page - only when the document is signed (empty string otherwise)."""
    return signature_data_uri() if signature_valid(doc) else ""


def mp_full_name():
    return f"{MP_NAME} {MP_CREDENTIALS}"


def is_managing_partner(name):
    """True when a 'signed by' name is the Managing Partner (credentials and title are ignored), so a letter signed by
    another partner never carries his signature."""
    return bool(name) and MP_NAME.lower() in " ".join(str(name).lower().replace(",", " ").split())


def add_signature_picture(paragraph, width_in=1.45):
    """Put the signature image in a paragraph of a Word document. Returns False when there is no image."""
    import io
    data = signature_bytes()
    if not data:
        return False
    paragraph.add_run().add_picture(io.BytesIO(data), width=Inches(width_in))
    return True


def add_signature_block(doc, firm_name=None, width_in=1.45, signed_on=None):
    """The signature, a rule, the Managing Partner's name and title and 'For <firm>', as a run of paragraphs."""
    firm_name = firm_name or firm_info().get("name", "")
    p = doc.add_paragraph()
    _tight(p, before=4, after=0)
    has_image = add_signature_picture(p, width_in)
    line = doc.add_paragraph()
    _tight(line, before=0, after=0)
    if not has_image:
        _run(line, "_______________________________")
    else:
        para_border(line, edge="top", colour=NAVY, size=6, space=1)
    _run(line, mp_full_name(), size=10, bold=True)
    for txt in (MP_TITLE, f"For {firm_name}" if firm_name else "", f"Signed {signed_on.strftime('%d %b %Y')}" if signed_on else ""):
        if txt:
            q = doc.add_paragraph()
            _tight(q, before=0, after=0)
            _run(q, txt, size=9.5, colour=MUTED)


def pdf_signature_flowables(width=None, firm_name=None, align="LEFT", signed_on=None):
    """The signature block for the reportlab PDFs: image, name, title, 'For <firm>'. Empty list when there is no image."""
    import io
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm as _mm
    from reportlab.platypus import Image as RLImage, Paragraph, Spacer
    from xml.sax.saxutils import escape
    data = signature_bytes()
    if not data:
        return []
    firm_name = firm_name or firm_info().get("name", "")
    base = ParagraphStyle("mpsig", fontName="Helvetica", fontSize=9, leading=12, textColor=rgb_tuple(INK))
    img = RLImage(io.BytesIO(data), width=36 * _mm, height=36 * _mm * 178 / 294)
    img.hAlign = align
    out = [Spacer(1, 2 * _mm), img,
           Paragraph(f"<b>{escape(mp_full_name())}</b>", base),
           Paragraph(escape(MP_TITLE), ParagraphStyle("mpsig2", parent=base, textColor=rgb_tuple(MUTED)))]
    if firm_name:
        out.append(Paragraph(f"For {escape(firm_name)}", ParagraphStyle("mpsig3", parent=base, textColor=rgb_tuple(MUTED))))
    if signed_on:
        out.append(Paragraph(f"Signed {signed_on.strftime('%d %b %Y')}", ParagraphStyle("mpsig4", parent=base, textColor=rgb_tuple(MUTED))))
    return out


def rgb_tuple(hex_colour):
    from reportlab.lib import colors
    return colors.HexColor("#" + hex_colour)


def _cell_has_picture(cell):
    return bool(cell._tc.findall(".//" + qn("w:drawing")))


def sign_template_rows(doc, width_in=1.2):
    """In a library template, put the signature in the Signature cell of every row where Neverlank signs (the row's first
    cell mentions Neverlank). Returns how many were signed; rows already carrying a picture are left alone."""
    import io
    data = signature_bytes()
    if not data:
        return 0
    signed = 0
    for table in doc.tables:
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        if "signature" not in header:
            continue
        col = header.index("signature")
        for row in table.rows[1:]:
            if "neverlank" not in row.cells[0].text.lower():
                continue
            cell = row.cells[col]
            if _cell_has_picture(cell):
                continue
            para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
            para.add_run().add_picture(io.BytesIO(data), width=Inches(width_in))
            signed += 1
    return signed


# ------------------------------------------------------------------ low-level XML helpers

def shade(cell, hex_colour):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour)
    tc_pr.append(shd)


def cell_borders(cell, **edges):
    """cell_borders(cell, bottom=("single", 24, GOLD), left=("nil",)) - sz in eighths of a point."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, spec in edges.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), spec[0])
        if len(spec) > 1:
            el.set(qn("w:sz"), str(spec[1]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec[2])


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def table_borders(table, colour=LINE, size=4, inside=True):
    """Thin light borders on the whole table (applies to rows added later too)."""
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    names = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for name in names:
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    tbl_pr.append(borders)


def table_cell_margins(table, top=50, bottom=50, left=90, right=90):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblCellMar"))
    if old is not None:
        tbl_pr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def para_border(paragraph, edge="bottom", colour=GOLD, size=8, space=2):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), colour)
    borders.append(el)


def para_shade(paragraph, hex_colour):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour)
    p_pr.append(shd)


def _field_runs(paragraph, instruction, size=8, colour=INK):
    """Insert a live field (PAGE / NUMPAGES) as the begin / instruction / separate / result / end runs Word expects."""
    def piece(build):
        r = _run(paragraph, "", size=size, colour=colour)
        build(r._r)
    def fld(kind):
        def build(r):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r.append(el)
        return build
    def instr(r):
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = f" {instruction} "
        r.append(el)
    def result(r):
        el = OxmlElement("w:t")
        el.text = "1"
        r.append(el)
    for build in (fld("begin"), instr, fld("separate"), result, fld("end")):
        piece(build)


def _fixed_widths(table, widths_in):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def _run(paragraph, text, size=10, bold=False, italic=False, colour=INK, font=BODY_FONT):
    r = paragraph.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.color.rgb = rgb(colour)
    return r


def _tight(paragraph, before=0, after=0, line=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def usable_width_in(doc):
    s = doc.sections[0]
    return (s.page_width - s.left_margin - s.right_margin) / 914400.0


# ------------------------------------------------------------------ document set-up

def setup_document(doc, margins_in=0.9, base_size=10.5):
    """Fonts, heading styles and margins for a branded document."""
    for section in doc.sections:
        section.left_margin = Inches(margins_in)
        section.right_margin = Inches(margins_in)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.9)
        section.footer_distance = Inches(0.0)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(base_size)
    normal.font.color.rgb = rgb(INK)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    for name, size, colour in (("Heading 1", 15, NAVY), ("Heading 2", 12.5, NAVY), ("Heading 3", 11, GOLD)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = HEAD_FONT if name != "Heading 3" else BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(colour)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        face = HEAD_FONT if name != "Heading 3" else BODY_FONT
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), face)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True



def statement_page_header(section, lines, rule_colour="2E75B6", size=13):
    """A plain running page header for one section of a document: the given lines (entity name, statement title, period) one under the other
    in the body font, with a thin coloured rule under the last. Unlinks the section's header from the previous section so only this
    section (and any later one that links to it) carries it."""
    header = section.header
    header.is_linked_to_previous = False
    for p in list(header.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = header.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, text in enumerate(lines):
        p = first if i == 0 else header.add_paragraph(style=first.style)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _tight(p, before=0, after=2 if i < len(lines) - 1 else 0)
        _run(p, text, size=size, colour=INK)
        if i == len(lines) - 1:
            para_border(p, "bottom", rule_colour, size=6, space=4)
    section.header_distance = Inches(0.35)


def blank_page_header(section):
    """An empty header for this section (it does not carry on the previous section's header)."""
    header = section.header
    header.is_linked_to_previous = False
    for p in list(header.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(header.paragraphs[0].runs):
        r._r.getparent().remove(r._r)


def branded_footer(doc, firm=None, center=None):
    """A navy footer strip with a gold rule: firm name, website (or `center` text, e.g. a
    template reference) and 'Page X of Y'."""
    firm = firm or firm_info()
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        for p in list(footer.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        p = footer.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        width = section.page_width - section.left_margin - section.right_margin
        p.paragraph_format.tab_stops.add_tab_stop(width // 2, WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        para_border(p, "top", GOLD, size=18, space=4)
        para_shade(p, NAVY_DARK)
        _tight(p, before=0, after=0)
        _run(p, f" {firm.get('name', '')}", size=8, colour=SOFT_TEXT)
        _run(p, "\t", size=8)
        _run(p, center if center is not None else (firm.get("web", "") or ""), size=8, colour=GOLD_LIGHT)
        _run(p, "\tPage ", size=8, colour=SOFT_TEXT)
        _field_runs(p, "PAGE", size=8, colour=SOFT_TEXT)
        _run(p, " of ", size=8, colour=SOFT_TEXT)
        _field_runs(p, "NUMPAGES", size=8, colour=SOFT_TEXT)
        _run(p, " ", size=8)


# ------------------------------------------------------------------ the letterhead band

def letterhead_band(doc, title, sub=None, number=None, firm=None):
    """The invoice-style header: navy band with logo, firm name, gold strapline and
    address on the left; the document title (gold), its number and a sub-line on the right."""
    firm = firm or firm_info()
    width = usable_width_in(doc)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    logo_w, right_w = 1.1, 1.95
    widths = [logo_w, width - logo_w - right_w, right_w]
    _fixed_widths(table, widths)
    table_borders(table, colour=NAVY, size=0, inside=False)
    logo_cell, firm_cell, doc_cell = table.rows[0].cells
    for c in (logo_cell, firm_cell, doc_cell):
        shade(c, NAVY)
        cell_margins(c, top=150, bottom=150, left=140, right=140)
        cell_borders(c, bottom=("single", 36, GOLD))
        c.vertical_alignment = 1  # centre

    lp = logo_cell.paragraphs[0]
    _tight(lp)
    if os.path.exists(LOGO_PATH):
        try:
            lp.add_run().add_picture(LOGO_PATH, height=Inches(0.95))
        except Exception:
            pass

    p = firm_cell.paragraphs[0]
    _tight(p)
    _run(p, firm.get("name", ""), size=13.5, bold=True, colour="FFFFFF", font=HEAD_FONT)
    p = firm_cell.add_paragraph()
    _tight(p, before=2)
    _run(p, (firm.get("strapline") or firm.get("tagline") or "").upper(), size=7.5, bold=True, colour=GOLD_LIGHT)
    for text in (firm.get("address"),
                 "  ·  ".join(x for x in (firm.get("web"), firm.get("email"), firm.get("phone")) if x),
                 "  ·  ".join(x for x in (f"VAT No: {firm['vat_number']}" if firm.get("vat_number") else "",
                                           f"TIN: {firm['tin']}" if firm.get("tin") else "") if x)):
        if text:
            p = firm_cell.add_paragraph()
            _tight(p, before=1)
            _run(p, text, size=8, colour=SOFT_TEXT)

    p = doc_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _tight(p)
    longest = max((len(w) for w in title.split()), default=0)
    _run(p, title.upper(), size=10.5 if longest >= 13 else (15 if len(title) < 22 else 12), bold=True, colour=GOLD_LIGHT, font=HEAD_FONT)
    if number:
        p = doc_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _tight(p, before=3)
        _run(p, number, size=11, bold=True, colour="FFFFFF")
    if sub:
        p = doc_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _tight(p, before=1)
        _run(p, sub, size=8.5, colour=SOFT_TEXT)

    spacer = doc.add_paragraph()
    _tight(spacer, after=4)
    return table


def meta_strip(doc, items):
    """A row of label/value cells on a light tint (the invoice's no. / date / due date / currency strip)."""
    items = [(k, v) for k, v in items if k]
    if not items:
        return None
    width = usable_width_in(doc)
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed_widths(table, [width / len(items)] * len(items))
    table_borders(table, colour=LINE, size=4, inside=True)
    for cell, (label, value) in zip(table.rows[0].cells, items):
        shade(cell, TINT)
        cell_margins(cell, top=80, bottom=80, left=140, right=100)
        p = cell.paragraphs[0]
        _tight(p)
        _run(p, str(label).upper(), size=7, bold=True, colour=MUTED)
        p = cell.add_paragraph()
        _tight(p, before=2)
        _run(p, str(value if value not in (None, "") else "—"), size=10, bold=True, colour=NAVY)
    spacer = doc.add_paragraph()
    _tight(spacer, after=2)
    return table


def party_block(doc, label, name, lines=()):
    """The 'Bill to' style block: tint panel with a gold bar on the left."""
    width = usable_width_in(doc)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed_widths(table, [width])
    cell = table.rows[0].cells[0]
    shade(cell, TINT)
    cell_margins(cell, top=100, bottom=100, left=200, right=120)
    cell_borders(cell, left=("single", 36, GOLD), top=("nil",), bottom=("nil",), right=("nil",))
    p = cell.paragraphs[0]
    _tight(p)
    _run(p, str(label).upper(), size=7, bold=True, colour=GOLD)
    p = cell.add_paragraph()
    _tight(p, before=2)
    _run(p, name or "", size=11.5, bold=True, colour=NAVY)
    for line in lines:
        if line:
            p = cell.add_paragraph()
            _tight(p, before=1)
            _run(p, line, size=9, colour="555555")
    spacer = doc.add_paragraph()
    _tight(spacer, after=2)
    return table


# ------------------------------------------------------------------ headings, tables, text

def heading(doc, text, level=1):
    """A heading in the firm style (navy serif with a gold rule under level 1)."""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = rgb(NAVY if level < 3 else GOLD)
        run.font.name = HEAD_FONT if level < 3 else BODY_FONT
    if level == 1:
        para_border(h, "bottom", GOLD, size=8, space=3)
    return h


def style_table(table, header_rows=1, numeric_from=None, zebra=False):
    """Navy header row with white bold text, light grid, comfortable padding.
    `numeric_from`: index of the first column to right-align (money columns)."""
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table, colour=LINE, size=4, inside=True)
    table_cell_margins(table)
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            is_head = ri < header_rows
            if is_head:
                shade(cell, NAVY)
            elif zebra and (ri - header_rows) % 2 == 1:
                shade(cell, "FAFBFD")
            for p in cell.paragraphs:
                _tight(p, before=1, after=1)
                if numeric_from is not None and ci >= numeric_from:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    if is_head:
                        r.bold = True
                        r.font.color.rgb = rgb("FFFFFF")
    return table


def note_box(doc, text, title=None):
    """A small tinted note panel (the invoice's 'Notes' block)."""
    width = usable_width_in(doc)
    table = doc.add_table(rows=1, cols=1)
    _fixed_widths(table, [width])
    cell = table.rows[0].cells[0]
    shade(cell, TINT)
    cell_margins(cell, top=90, bottom=90, left=180, right=120)
    cell_borders(cell, left=("single", 24, GOLD), top=("nil",), bottom=("nil",), right=("nil",))
    p = cell.paragraphs[0]
    _tight(p)
    if title:
        _run(p, title.upper(), size=7, bold=True, colour=GOLD)
        p = cell.add_paragraph()
        _tight(p, before=2)
    _run(p, text, size=9.5, colour=INK)
    doc.add_paragraph()
    return table


# ------------------------------------------------------------------ Excel

def xlsx_title_band(ws, title, subtitle_lines=(), ncols=8, firm=None):
    """Navy band across the top of a sheet (logo, firm name, strapline, address) with the sheet
    title in row 5 and a details line in row 6. Returns the first free row (8)."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    firm = firm or firm_info()
    navy = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    ncols = max(ncols, 6)
    for r in range(1, 5):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = navy
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 14
    ws.row_dimensions[3].height = 13
    ws.row_dimensions[4].height = 13
    ws["B1"] = firm.get("name", "")
    ws["B1"].font = Font(name="Georgia", bold=True, size=15, color="FFFFFF")
    ws["B2"] = (firm.get("strapline") or firm.get("tagline") or "").upper()
    ws["B2"].font = Font(bold=True, size=8, color=GOLD_LIGHT)
    contact = " · ".join(x for x in (firm.get("address"), firm.get("web")) if x)
    ws["B3"] = contact
    ws["B3"].font = Font(size=8, color=SOFT_TEXT)
    gold = Side(style="thick", color=GOLD)
    for c in range(1, ncols + 1):
        ws.cell(row=4, column=c).border = Border(bottom=gold)
    ws["A5"] = title
    ws["A5"].font = Font(name="Georgia", bold=True, size=13, color=NAVY)
    ws.row_dimensions[5].height = 20
    if subtitle_lines:
        ws["A6"] = "  ·  ".join(x for x in subtitle_lines if x)
        ws["A6"].font = Font(size=9.5, color=MUTED)
    if os.path.exists(LOGO_PATH):
        try:
            img = XLImage(LOGO_PATH)
            img.width, img.height = 54, 54
            ws.add_image(img, "A1")
        except Exception:
            pass
    return 8


def xlsx_header_style():
    """(font, fill, alignment, border) for column-header cells."""
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    thin = Side(style="thin", color=LINE)
    return (Font(bold=True, color="FFFFFF"), PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid"),
            Alignment(horizontal="center", vertical="center", wrap_text=True),
            Border(left=thin, right=thin, top=thin, bottom=thin))


def xlsx_page_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.oddFooter.left.text = "&8Neverlank Chartered Accountants"
    ws.oddFooter.right.text = "&8Page &P of &N"
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------ PDF (reportlab)

def pdf_letterhead(width, title, number=None, sub=None, firm=None):
    """The invoice-style navy letterhead band as a reportlab Table of `width` points."""
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, Spacer, Table, TableStyle

    firm = firm or firm_info()
    base = ParagraphStyle("lh", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.white)
    right = ParagraphStyle("lhr", parent=base, alignment=TA_RIGHT)
    contact = [firm.get("address"), "  ·  ".join(x for x in (firm.get("web"), firm.get("email"), firm.get("phone")) if x)]
    ids = "  ·  ".join(x for x in (f"VAT No: {firm['vat_number']}" if firm.get("vat_number") else "",
                                   f"TIN: {firm['tin']}" if firm.get("tin") else "") if x)
    if ids:
        contact.append(ids)
    firm_cell = [
        Paragraph(f"<font name='Helvetica-Bold' size='15'>{escape(firm.get('name', ''))}</font>", ParagraphStyle("lhn", parent=base, leading=19)),
        Paragraph(f"<font color='#{GOLD_LIGHT}' size='7.5'>{escape((firm.get('strapline') or firm.get('tagline') or '').upper())}</font>", ParagraphStyle("lhs", parent=base, leading=11)),
        Spacer(1, 3),
    ] + [Paragraph(f"<font color='#{SOFT_TEXT}' size='8'>{escape(c)}</font>", base) for c in contact if c]
    size = 20 if len(title) <= 14 else (15 if len(title) <= 24 else 12)
    doc_cell = [Paragraph(f"<font name='Helvetica-Bold' size='{size}' color='#{GOLD_LIGHT}'>{escape(title.upper())}</font>",
                          ParagraphStyle("lht", parent=right, leading=size + 4))]
    if number:
        doc_cell.append(Paragraph(f"<font name='Helvetica-Bold' size='11'>{escape(number)}</font>", right))
    if sub:
        doc_cell.append(Paragraph(f"<font color='#{SOFT_TEXT}' size='8'>{escape(sub)}</font>", right))
    logo = Image(LOGO_PATH, 24 * mm, 24 * mm) if os.path.exists(LOGO_PATH) else ""
    right_w = 58 * mm
    band = Table([[logo, firm_cell, doc_cell]], colWidths=[30 * mm, width - 30 * mm - right_w, right_w])
    band.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{NAVY}")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LINEBELOW", (0, 0), (-1, -1), 3, colors.HexColor(f"#{GOLD}")),
    ]))
    return band


def pdf_meta_strip(width, items):
    """Label/value cells on a tint, like the invoice's details strip."""
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle
    base = ParagraphStyle("ms", fontName="Helvetica", fontSize=9, leading=12)
    cells = [[Paragraph(f"<font size='7' color='#{MUTED}'><b>{escape(str(k).upper())}</b></font>", base),
              Paragraph(f"<font name='Helvetica-Bold' size='10' color='#{NAVY}'>{escape(str(v if v not in (None, '') else '—'))}</font>", base)]
             for k, v in items]
    t = Table([cells], colWidths=[width / len(cells)] * len(cells))
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{TINT}")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{LINE}")),
        ("LINEAFTER", (0, 0), (-2, -1), 0.5, colors.HexColor(f"#{LINE}")), ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 9), ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return t


def pdf_footer(firm=None, with_page=True):
    """A reportlab onFirstPage/onLaterPages callback drawing the navy footer strip with a gold rule."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    firm = firm or firm_info()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(f"#{NAVY_DARK}"))
        canvas.rect(0, 0, A4[0], 9 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor(f"#{GOLD}"))
        canvas.rect(0, 9 * mm, A4[0], 0.8 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor(f"#{SOFT_TEXT}"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(16 * mm, 3.4 * mm, firm.get("name", ""))
        right = firm.get("web", "") or ""
        if with_page:
            right = (right + "   ·   " if right else "") + f"Page {doc.page}"
        canvas.drawRightString(A4[0] - 16 * mm, 3.4 * mm, right)
        canvas.restoreState()
    return footer


def pdf_table_style(header_rows=1):
    """TableStyle commands: navy header, light grid (extend with your own alignment commands)."""
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), colors.HexColor(f"#{NAVY}")),
        ("TEXTCOLOR", (0, 0), (-1, header_rows - 1), colors.white),
        ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, header_rows), (-1, -1), 0.5, colors.HexColor(f"#{LINE}")),
        ("ROWBACKGROUNDS", (0, header_rows), (-1, -1), [colors.white, colors.HexColor("#FAFBFD")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ])


def total_panel(doc, label, value, align_right=True, width_in=3.2):
    """The invoice's navy 'Total due' bar: label in white, amount in gold (Word)."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT if align_right else WD_TABLE_ALIGNMENT.CENTER
    _fixed_widths(table, [width_in * 0.5, width_in * 0.5])
    left, right = table.rows[0].cells
    for c in (left, right):
        shade(c, NAVY)
        cell_margins(c, top=100, bottom=100, left=160, right=160)
    p = left.paragraphs[0]
    _tight(p)
    _run(p, label, size=11, bold=True, colour="FFFFFF")
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _tight(p)
    _run(p, value, size=12, bold=True, colour=GOLD_LIGHT)
    return table


def letter_firm(name=None, strapline=None, address=None):
    """A copy of the firm letterhead dictionary with some lines replaced (e.g. a client's name
    on a payslip the firm processes for them)."""
    firm = dict(firm_info())
    if name is not None:
        firm["name"] = name
    if strapline is not None:
        firm["strapline"] = strapline
    if address is not None:
        firm["address"] = address
    return firm


# ------------------------------------------------------------------ schema order

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_ORDER = {
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
             "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
              "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
              "tblDescription"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
            "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "pBdr": ["top", "left", "start", "bottom", "right", "end", "between", "bar"],
    "tcBorders": ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV", "tl2br", "tr2bl"],
    "tblBorders": ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV"],
    "tcMar": ["top", "left", "start", "bottom", "right", "end"],
    "tblCellMar": ["top", "left", "start", "bottom", "right", "end"],
}


def _sort_children(el, order):
    rank = {name: i for i, name in enumerate(order)}
    kids = list(el)
    # a property may only appear once: where it was set twice, the later setting wins
    seen = {}
    for k in kids:
        seen[k.tag] = k
    if len(seen) != len(kids):
        for k in kids:
            if seen[k.tag] is not k:
                el.remove(k)
        kids = list(el)
    ordered = sorted(kids, key=lambda k: rank.get(k.tag.replace(_W, ""), len(order)))
    if ordered != kids:
        for k in kids:
            el.remove(k)
        for k in ordered:
            el.append(k)


def fix_schema_order(root):
    """Put the children of table / cell / paragraph property elements into the order the Word schema
    requires (Word refuses to open some files whose properties are in the wrong order). Safe to call
    repeatedly. `root` is an lxml element (document body, a footer, a styles part...)."""
    for tag, order in _ORDER.items():
        for el in root.iter(_W + tag):
            _sort_children(el, order)


def fix_document(doc):
    """fix_schema_order over the body, every header/footer and the styles of a python-docx Document."""
    fix_schema_order(doc.element)
    fix_schema_order(doc.styles.element)
    try:  # python-docx's default settings carry <w:zoom/> without its required percent
        for z in doc.settings.element.iter(_W + "zoom"):
            if z.get(_W + "percent") is None:
                z.set(_W + "percent", "100")
    except Exception:
        pass
    for section in doc.sections:
        for part in (section.footer, section.header, section.first_page_footer, section.first_page_header,
                     section.even_page_footer, section.even_page_header):
            try:
                if part.is_linked_to_previous is False or part._element is not None:
                    fix_schema_order(part._element)
            except Exception:
                pass
