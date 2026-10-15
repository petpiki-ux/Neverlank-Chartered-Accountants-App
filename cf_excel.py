"""Corporate Finance & Business Advisory - Excel exports with LIVE formulas.

The forecast workbook rebuilds the engine (cf_engine.run) in Excel formulas: change an assumption on the
Assumptions sheet and the Calc, Annual Summary, Monthly Year 1 and Valuation sheets recalculate. Structural
choices (number of years, revenue streams, expense lines, loans, loan repayment type, expense kind) are fixed
when the workbook is exported; all amounts, rates, days, growth, seasonality and valuation inputs are live.
Inputs are blue on a pale-yellow fill. The test-suite recalculates the workbook in LibreOffice and compares every
line with the Python engine.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.properties import PageSetupProperties

import brand_docs as brand
import cf_engine as E
import cf_tools as T

NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT_IN = '0.0'
MULT = '0.00"x"'
DAYS = '0.0'

NAVY, GOLD, LINE, TINT, MUTED = brand.NAVY, brand.GOLD, brand.LINE, brand.TINT, brand.MUTED
F_INPUT = Font(color="0000FF")
FILL_INPUT = PatternFill(start_color="FFF9E0", end_color="FFF9E0", fill_type="solid")
FILL_HEAD = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
FILL_SEC = PatternFill(start_color="E9EEF5", end_color="E9EEF5", fill_type="solid")
FILL_TOT = PatternFill(start_color=TINT, end_color=TINT, fill_type="solid")
THIN = Side(style="thin", color=LINE)
TOP = Border(top=Side(style="thin", color=NAVY))
FIRST_PERIOD_COL = 3  # column C on Calc


# --------------------------------------------------------------------------- small helpers

def _setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _section(ws, row, text, ncols):
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = FILL_SEC
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, color=NAVY)


def _head(ws, row, values, start_col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), FILL_HEAD
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)


def _inp(cell, value, fmt=None):
    cell.value = value
    cell.font, cell.fill = F_INPUT, FILL_INPUT
    if fmt:
        cell.number_format = fmt
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fmt(cell, fmt, bold=False, total=False):
    cell.number_format = fmt
    if bold or total:
        cell.font = Font(bold=True)
    if total:
        cell.fill = FILL_TOT
        cell.border = TOP


def _meta_lines(meta):
    return [x for x in (meta.get("client"), meta.get("name"), meta.get("currency")) if x]


# --------------------------------------------------------------------------- Assumptions sheet

def _assumptions_sheet(wb, a, meta):
    ws = wb.create_sheet("Assumptions")
    n = a["years"]
    ycols = [4 + i for i in range(n)]  # D..
    ncols = max(4 + n, 15)
    row = brand.xlsx_title_band(ws, "Assumptions", _meta_lines(meta) + ["Blue cells on yellow are inputs; everything else is calculated"], ncols=ncols)
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 46, "B": 14, "C": 16})
    for c in ycols:
        ws.column_dimensions[CL(c)].width = 13
    for c in range(4 + n, 16):
        ws.column_dimensions[CL(c)].width = 11
    R = {"row": {}, "op": {}, "loan": [], "season": [], "rev_annual": [], "rev_y1": [], "opex_annual": [], "opex_pct": [],
         "opex_growth": [], "creditor": [], "ycols": ycols}

    def last_col():
        return CL(ycols[-1])

    def yr(name):
        r = R["row"][name]
        return f"Assumptions!$D${r}:${last_col()}${r}"

    R["yr"] = yr

    # ---- general
    _section(ws, row, "General", ncols)
    row += 1
    gen = [("Forecast start (first month)", "start_month", a["start_month"], None),
           ("Forecast years (1-10)", "years", n, "0"),
           ("Currency", "currency", a["currency"], None),
           ("Tax rate (%)", "tax", a["tax_rate"], PCT_IN),
           ("Depreciation rate (% of gross cost p.a.)", "dep_rate", a["dep_rate"], PCT_IN),
           ("Interest earned on positive cash (% p.a.)", "deposit", a["deposit_rate"], PCT_IN),
           ("Overdraft interest (% p.a.)", "overdraft", a["overdraft_rate"], PCT_IN)]
    for label, key, val, fmt in gen:
        ws.cell(row=row, column=1, value=label)
        _inp(ws.cell(row=row, column=3), val, fmt)
        R[key] = f"Assumptions!$C${row}"
        row += 1
    row += 1

    # ---- seasonality
    _section(ws, row, "Year 1 revenue phasing - share of the year's revenue in each month (%, should total 100)", ncols)
    row += 1
    start = E.parse_start_month(a["start_month"])
    for m in range(12):
        d = E.add_months(start, m)
        ws.cell(row=row, column=3 + m, value=f"{E.MONTH_ABBR[d.month - 1]}-{str(d.year)[2:]}").font = Font(bold=True, color=NAVY)
        _inp(ws.cell(row=row + 1, column=3 + m), a["seasonality"][m], '0.00')
        R["season"].append(f"Assumptions!${CL(3 + m)}${row + 1}")
    ws.cell(row=row + 1, column=1, value="Share of year 1 revenue (%)")
    ws.cell(row=row + 1, column=2, value=f"=SUM(C{row + 1}:N{row + 1})").number_format = '0.00'
    R["season_total"] = f"Assumptions!$B${row + 1}"
    row += 3

    # ---- yearly header
    _section(ws, row, "Yearly drivers", ncols)
    row += 1
    _head(ws, row, ["", "", ""] + [f"Year {y + 1}" for y in range(n)])
    row += 1
    for i, y in enumerate(range(1, n + 1)):
        ws.cell(row=row, column=4 + i, value=E.year_label(a, y)).font = Font(size=8, color=MUTED)
        ws.cell(row=row, column=4 + i).alignment = Alignment(horizontal="center")
    row += 1

    def yearly(label, key, values, fmt=PCT_IN, note=None, first_blank=False):
        nonlocal row
        ws.cell(row=row, column=1, value=label)
        if note:
            ws.cell(row=row, column=2, value=note).font = Font(size=8, color=MUTED)
        for i in range(n):
            cell = ws.cell(row=row, column=4 + i)
            if first_blank and i == 0:
                cell.value = None
                continue
            _inp(cell, values[i], fmt)
        R["row"][key] = row
        row += 1
        return row - 1

    # ---- revenue
    _section(ws, row, "Revenue", ncols)
    row += 1
    for i, st in enumerate(a["revenue"]):
        ws.cell(row=row, column=1, value=f"{st['name']} - year 1 revenue")
        _inp(ws.cell(row=row, column=3), st["y1"], NUM)
        R["rev_y1"].append(f"$C${row}")
        y1row = row
        row += 1
        grow_row = yearly(f"{st['name']} - growth (%)", f"growth{i}", st["growth"], PCT_IN, first_blank=True)
        ws.cell(row=row, column=1, value=f"{st['name']} - annual revenue (calculated)")
        for k in range(n):
            c = ws.cell(row=row, column=4 + k)
            c.value = f"=$C${y1row}" if k == 0 else f"={CL(3 + k)}{row}*(1+{CL(4 + k)}{grow_row}/100)"
            _fmt(c, NUM)
        R["rev_annual"].append(row)
        row += 1
    # ---- costs
    _section(ws, row, "Costs", ncols)
    row += 1
    yearly("Cost of sales (% of revenue)", "cogs", a["cogs_pct"], PCT_IN)
    for j, o in enumerate(a["opex"]):
        ws.cell(row=row, column=1, value=f"{o['name']}  ({'fixed amount' if o['kind'] == 'fixed' else '% of revenue'})").font = Font(bold=True)
        ws.cell(row=row, column=2, value="On credit? (1/0)").font = Font(size=8, color=MUTED)
        _inp(ws.cell(row=row, column=3), 1 if o["creditors"] else 0, "0")
        R["creditor"].append(f"Assumptions!$C${row}")
        row += 1
        if o["kind"] == "fixed":
            y1row = row
            ws.cell(row=row, column=1, value="   Year 1 amount")
            _inp(ws.cell(row=row, column=3), o["y1"], NUM)
            row += 1
            grow_row = yearly("   Growth (%)", f"og{j}", o["growth"], PCT_IN, first_blank=True)
            ws.cell(row=row, column=1, value="   Annual amount (calculated)")
            for k in range(n):
                c = ws.cell(row=row, column=4 + k)
                c.value = f"=$C${y1row}" if k == 0 else f"={CL(3 + k)}{row}*(1+{CL(4 + k)}{grow_row}/100)"
                _fmt(c, NUM)
            R["opex_annual"].append(row)
            R["opex_pct"].append(None)
            row += 1
        else:
            pr = yearly("   % of revenue", f"op{j}", o["pct"], PCT_IN)
            R["opex_annual"].append(None)
            R["opex_pct"].append(pr)
    # ---- capex, WC, equity
    _section(ws, row, "Capital expenditure, working capital and equity", ncols)
    row += 1
    yearly("Capital expenditure", "capex", a["capex"], NUM)
    yearly("Receivable days (on revenue)", "rec_days", a["rec_days"], DAYS)
    yearly("Inventory days (on cost of sales)", "inv_days", a["inv_days"], DAYS)
    yearly("Payable days (on cost of sales + 'on credit' expenses)", "pay_days", a["pay_days"], DAYS)
    yearly("New share capital issued (start of year)", "equity_issue", a["equity_issue"], NUM)
    yearly("Dividend payout (% of profit after tax)", "payout", a["dividend_payout"], PCT_IN)
    row += 1

    # ---- opening balance sheet
    _section(ws, row, "Opening balance sheet (at the forecast start)", ncols)
    row += 1
    op = a["opening"]
    for label, key in (("Property, plant & equipment - cost", "ppe_cost"), ("Property, plant & equipment - accumulated depreciation", "ppe_accdep"),
                       ("Inventory", "inventory"), ("Trade receivables", "receivables"), ("Other assets", "other_assets"),
                       ("Cash and cash equivalents (negative = overdraft)", "cash"), ("Trade payables", "payables"),
                       ("Other liabilities", "other_liabilities"), ("Share capital", "share_capital"),
                       ("Tax losses brought forward (assessed loss)", "tax_loss_bf")):
        ws.cell(row=row, column=1, value=label)
        _inp(ws.cell(row=row, column=3), op[key], NUM)
        R["op"][key] = f"Assumptions!$C${row}"
        row += 1
    row += 1

    # ---- loans
    _section(ws, row, "Loans and borrowings", ncols)
    row += 1
    loan_open_cells = []
    for k, L in enumerate(a["loans"]):
        T = E.loan_timing(a, L)
        kind_txt = "bullet - repaid at the end date" if L["repay_type"] == "bullet" else "equal monthly principal instalments to the end date"
        dates_txt = ""
        if L.get("start_date") or L.get("end_date"):
            dates_txt = f"; dates {L.get('start_date') or '-'} to {L.get('end_date') or '-'}"
        ws.cell(row=row, column=1, value=f"{L['name']}  ({kind_txt}{dates_txt})").font = Font(bold=True)
        row += 1
        refs = {}
        for label, key, val, fmt in (("   Balance at start", "open", L["opening_balance"], NUM), ("   New drawdown", "draw", L["drawdown_amount"], NUM),
                                     ("   Interest rate (% p.a.)", "rate", L["rate"], PCT_IN),
                                     ("   Establishment fee (% of the amount drawn, on top of interest)", "fee_pct", L["fee_pct"], PCT_IN),
                                     ("   Establishment fee (fixed amount, on top of interest)", "fee_amt", L["fee_amount"], NUM),
                                     ("   Drawdown month no. (1 = first forecast month)", "draw_m", T["draw_idx"] + 1, "0"),
                                     ("   First principal instalment month no.", "first_m", T["first_idx"] + 1, "0"),
                                     ("   Last month no. of the loan (end date)", "end_m", T["end_idx"] + 1, "0")):
            ws.cell(row=row, column=1, value=label)
            _inp(ws.cell(row=row, column=3), val, fmt)
            refs[key] = f"Assumptions!$C${row}"
            row += 1
        ws.cell(row=row, column=1, value="   Instalment months (calculated)")
        ws.cell(row=row, column=3, value=f"=MAX(1,{refs['end_m']}-{refs['first_m']}+1)")
        refs["term"] = f"Assumptions!$C${row}"
        row += 1
        ws.cell(row=row, column=1, value="   Principal repaid over the term (calculated)")
        ws.cell(row=row, column=3, value=f"={refs['open']}+{refs['draw']}").number_format = NUM
        refs["principal0"] = f"Assumptions!$C${row}"
        row += 1
        loan_open_cells.append(refs["open"])
        R["loan"].append(refs)
    ws.cell(row=row, column=1, value="Total borrowings at the start (calculated)").font = Font(bold=True)
    c = ws.cell(row=row, column=3, value="=" + ("+".join(loan_open_cells) if loan_open_cells else "0"))
    _fmt(c, NUM, bold=True)
    R["loans_open"] = f"Assumptions!$C${row}"
    row += 1
    ws.cell(row=row, column=1, value="Opening retained earnings & reserves (balancing figure, calculated)").font = Font(bold=True)
    o = R["op"]
    formula = (f"=({o['ppe_cost']}-{o['ppe_accdep']})+{o['inventory']}+{o['receivables']}+{o['other_assets']}+MAX({o['cash']},0)"
               f"-({o['payables']}+{o['other_liabilities']}+{R['loans_open']}+MAX(-{o['cash']},0))-{o['share_capital']}")
    c = ws.cell(row=row, column=3, value=formula)
    _fmt(c, NUM, bold=True)
    R["reserves"] = f"Assumptions!$C${row}"
    ws.freeze_panes = "B8"
    return ws, R


# --------------------------------------------------------------------------- Calc sheet

def _calc_sheet(wb, a, R, meta, periods):
    ws = wb.create_sheet("Calc")
    P = len(periods)
    ncols = FIRST_PERIOD_COL + P
    brand.xlsx_title_band(ws, "Calculation engine", _meta_lines(meta) + ["Monthly in year 1, annual thereafter - do not overwrite"], ncols=min(ncols, 30))
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 44, "B": 4})
    for i in range(P):
        ws.column_dimensions[CL(FIRST_PERIOD_COL + i)].width = 12.5
    nstreams, nopex, nloans = len(a["revenue"]), len(a["opex"]), len(a["loans"])

    # ---- layout: assign a row to every key first
    layout = []  # (key, label, fmt) ; key None -> section header

    def sec(text):
        layout.append((None, text, None))

    def line(key, label, fmt=NUM):
        layout.append((key, label, fmt))

    line("label", "Period", None)
    line("year_no", "Forecast year", "0")
    line("months", "Months in period", "0")
    line("pstart", "Months from the model start to this period (0 = first month)", "0")
    line("year_start", "First period of the year (1/0)", "0")
    line("year_end", "Last period of the year (1/0)", "0")
    line("share", "Share of annual revenue / pattern", "0.0000")
    sec("Income statement")
    for i, st in enumerate(a["revenue"]):
        line(f"rev{i}", f"Revenue - {st['name']}")
    line("revenue", "Revenue")
    line("cogs", "Cost of sales")
    for j, o in enumerate(a["opex"]):
        line(f"opex{j}", f"{o['name']}")
    line("opex", "Total operating expenses")
    line("creditor_opex", "  of which bought on credit")
    line("ebitda", "EBITDA")
    line("capex", "Capital expenditure")
    line("gross_open", "PPE gross cost - opening")
    line("accdep_open", "Accumulated depreciation - opening")
    line("dep", "Depreciation")
    line("ebit", "EBIT")
    sec("Loans")
    for k, L in enumerate(a["loans"]):
        line(f"ln{k}_open", f"{L['name']} - opening")
        line(f"ln{k}_draw", f"{L['name']} - drawdown")
        line(f"ln{k}_int", f"{L['name']} - interest")
        line(f"ln{k}_fee", f"{L['name']} - establishment fee")
        line(f"ln{k}_repay", f"{L['name']} - repayment")
        line(f"ln{k}_close", f"{L['name']} - closing")
    line("int_loans", "Loan interest")
    line("loan_fees", "Loan establishment fees")
    line("loan_draw", "Loan drawdowns")
    line("loan_repay", "Loan repayments")
    line("borrowings", "Borrowings - closing")
    sec("Finance cost, tax and profit")
    line("cash_open", "Cash - opening")
    line("int_income", "Interest income")
    line("int_od", "Overdraft interest")
    line("fin_cost", "Net finance cost")
    line("pbt", "Profit before tax")
    line("loss_open", "Tax losses - opening")
    line("loss_used", "Tax losses used")
    line("loss_added", "Tax losses added")
    line("loss_close", "Tax losses - closing")
    line("tax", "Tax")
    line("pat", "Profit after tax")
    line("dividends", "Dividends")
    line("equity_issue", "Share capital issued")
    sec("Working capital")
    line("inventory", "Inventory")
    line("receivables", "Trade receivables")
    line("payables", "Trade payables")
    line("d_inventory", "(Increase)/decrease in inventory")
    line("d_receivables", "(Increase)/decrease in receivables")
    line("d_payables", "Increase/(decrease) in payables")
    sec("Cash flow")
    line("cash_ops", "Cash generated from operations")
    line("interest_paid", "Interest paid")
    line("interest_received", "Interest received")
    line("tax_paid", "Tax paid")
    line("cfo", "Net cash from operating activities")
    line("cfi", "Net cash from investing activities")
    line("cff", "Net cash from financing activities")
    line("net_cash", "Net change in cash")
    line("cash", "Cash - closing")
    sec("Balance sheet")
    line("gross_close", "PPE gross cost - closing")
    line("accdep_close", "Accumulated depreciation - closing")
    line("nbv", "PPE net book value")
    line("share_capital", "Share capital")
    line("retained", "Retained earnings & reserves")
    line("total_assets", "Total assets")
    line("total_liabilities", "Total liabilities")
    line("total_equity", "Total equity")
    line("check", "Balance check (must be nil)")

    rw, r, fmts = {}, 8, {}
    for key, label, fmt in layout:
        if key is None:
            r += 1
            _section(ws, r, label, min(ncols, 30))
            r += 1
            continue
        rw[key] = r
        ws.cell(row=r, column=1, value=label)
        fmts[key] = fmt
        r += 1
    if True:  # header row styling
        _head(ws, rw["label"], [None] * (ncols), 1)
        ws.cell(row=rw["label"], column=1, value="Period")

    def rng(key):
        return f"Calc!$C${rw[key]}:${CL(FIRST_PERIOD_COL + P - 1)}${rw[key]}"

    R["calc_rng"] = rng
    R["calc_rw"] = rw
    R["calc_P"] = P

    bold_keys = {"revenue", "ebitda", "ebit", "pbt", "pat", "cfo", "net_cash", "cash", "total_assets", "total_liabilities", "total_equity", "borrowings"}
    for i, p in enumerate(periods):
        col = FIRST_PERIOD_COL + i
        c = CL(col)
        pc = CL(col - 1) if i else None

        def X(key, cc=c):
            return f"{cc}{rw[key]}"

        def Pv(key):
            return f"{pc}{rw[key]}"

        def idx(name):
            return f"INDEX({R['yr'](name)},{X('year_no')})"

        def put(key, value):
            cell = ws.cell(row=rw[key], column=col, value=value)
            if fmts.get(key):
                cell.number_format = fmts[key]
            if key in bold_keys:
                cell.font = Font(bold=True)
            return cell

        put("label", p["label"]).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=rw["label"], column=col).alignment = Alignment(horizontal="center")
        put("year_no", p["year_no"])
        put("months", p["months"])
        put("pstart", p["start_idx"])
        put("year_start", p["is_year_start"])
        put("year_end", p["is_year_end"])
        put("share", f"={R['season'][i]}/100" if p["year_no"] == 1 else 1)

        for s in range(nstreams):
            ra = R["rev_annual"][s]
            put(f"rev{s}", f"=INDEX(Assumptions!$D${ra}:${CL(3 + a['years'])}${ra},{X('year_no')})*{X('share')}")
        put("revenue", "=" + "+".join(X(f"rev{s}") for s in range(nstreams)))
        put("cogs", f"={X('revenue')}*{idx('cogs')}/100")
        for j, o in enumerate(a["opex"]):
            if o["kind"] == "fixed":
                ra = R["opex_annual"][j]
                put(f"opex{j}", f"=INDEX(Assumptions!$D${ra}:${CL(3 + a['years'])}${ra},{X('year_no')})*{X('months')}/12")
            else:
                put(f"opex{j}", f"={X('revenue')}*{idx(f'op{j}')}/100")
        put("opex", "=" + ("+".join(X(f"opex{j}") for j in range(nopex)) if nopex else "0"))
        put("creditor_opex", "=" + ("+".join(f"{R['creditor'][j]}*{X(f'opex{j}')}" for j in range(nopex)) if nopex else "0"))
        put("ebitda", f"={X('revenue')}-{X('cogs')}-{X('opex')}")
        put("capex", f"={idx('capex')}*{X('months')}/12")
        put("gross_open", f"={R['op']['ppe_cost']}" if i == 0 else f"={Pv('gross_close')}")
        put("accdep_open", f"={R['op']['ppe_accdep']}" if i == 0 else f"={Pv('accdep_close')}")
        put("dep", f"=MIN({X('gross_open')}-{X('accdep_open')}+{X('capex')},{R['dep_rate']}/100*({X('gross_open')}+{X('capex')})*{X('months')}/12)")
        put("ebit", f"={X('ebitda')}-{X('dep')}")

        for k, L in enumerate(a["loans"]):
            ref = R["loan"][k]
            put(f"ln{k}_open", f"={ref['open']}" if i == 0 else f"={Pv(f'ln{k}_close')}")
            ps, mo = X("pstart"), X("months")
            put(f"ln{k}_draw", f"=IF(AND({ps}<={ref['draw_m']}-1,{ref['draw_m']}-1<{ps}+{mo}),{ref['draw']},0)")
            put(f"ln{k}_int", f"={ref['rate']}/100/12*({X(f'ln{k}_open')}*MIN({mo},MAX(0,{ref['end_m']}-{ps}))"
                              f"+{X(f'ln{k}_draw')}*MAX(0,MIN({ps}+{mo},{ref['end_m']})-({ref['draw_m']}-1)))")
            put(f"ln{k}_fee", f"=IF({X(f'ln{k}_draw')}>0,{X(f'ln{k}_draw')}*{ref['fee_pct']}/100+{ref['fee_amt']},0)")
            if L["repay_type"] == "bullet":
                put(f"ln{k}_repay", f"=IF(AND({ps}<={ref['end_m']}-1,{ref['end_m']}-1<{ps}+{mo}),{X(f'ln{k}_open')}+{X(f'ln{k}_draw')},0)")
            else:
                put(f"ln{k}_repay", f"=MIN({X(f'ln{k}_open')}+{X(f'ln{k}_draw')},{ref['principal0']}/{ref['term']}"
                                    f"*MAX(0,MIN({ps}+{mo}-1,{ref['end_m']}-1)-MAX({ps},{ref['first_m']}-1)+1))")
            put(f"ln{k}_close", f"={X(f'ln{k}_open')}+{X(f'ln{k}_draw')}-{X(f'ln{k}_repay')}")
        put("int_loans", "=" + ("+".join(X(f"ln{k}_int") for k in range(nloans)) if nloans else "0"))
        put("loan_fees", "=" + ("+".join(X(f"ln{k}_fee") for k in range(nloans)) if nloans else "0"))
        put("loan_draw", "=" + ("+".join(X(f"ln{k}_draw") for k in range(nloans)) if nloans else "0"))
        put("loan_repay", "=" + ("+".join(X(f"ln{k}_repay") for k in range(nloans)) if nloans else "0"))
        put("borrowings", "=" + ("+".join(X(f"ln{k}_close") for k in range(nloans)) if nloans else "0"))

        put("cash_open", f"={R['op']['cash']}" if i == 0 else f"={Pv('cash')}")
        put("int_income", f"=MAX({X('cash_open')},0)*{R['deposit']}/100*{X('months')}/12")
        put("int_od", f"=MAX(-{X('cash_open')},0)*{R['overdraft']}/100*{X('months')}/12")
        put("fin_cost", f"={X('int_loans')}+{X('int_od')}-{X('int_income')}+{X('loan_fees')}")
        put("pbt", f"={X('ebit')}-{X('fin_cost')}")
        put("loss_open", f"={R['op']['tax_loss_bf']}" if i == 0 else f"={Pv('loss_close')}")
        put("loss_used", f"=MIN({X('loss_open')},MAX({X('pbt')},0))")
        put("loss_added", f"=MAX(-{X('pbt')},0)")
        put("loss_close", f"={X('loss_open')}-{X('loss_used')}+{X('loss_added')}")
        put("tax", f"=(MAX({X('pbt')},0)-{X('loss_used')})*{R['tax']}/100")
        put("pat", f"={X('pbt')}-{X('tax')}")
        put("dividends", f"={idx('payout')}/100*MAX({X('pat')},0)")
        put("equity_issue", f"=IF({X('year_start')}=1,{idx('equity_issue')},0)")

        ann = f"12/{X('months')}"
        put("inventory", f"={X('cogs')}*{ann}*{idx('inv_days')}/365")
        put("receivables", f"={X('revenue')}*{ann}*{idx('rec_days')}/365")
        put("payables", f"=({X('cogs')}+{X('creditor_opex')})*{ann}*{idx('pay_days')}/365")
        put("d_inventory", f"={R['op']['inventory']}-{X('inventory')}" if i == 0 else f"={Pv('inventory')}-{X('inventory')}")
        put("d_receivables", f"={R['op']['receivables']}-{X('receivables')}" if i == 0 else f"={Pv('receivables')}-{X('receivables')}")
        put("d_payables", f"={X('payables')}-{R['op']['payables']}" if i == 0 else f"={X('payables')}-{Pv('payables')}")

        put("cash_ops", f"={X('ebitda')}+{X('d_inventory')}+{X('d_receivables')}+{X('d_payables')}")
        put("interest_paid", f"={X('int_loans')}+{X('int_od')}")
        put("interest_received", f"={X('int_income')}")
        put("tax_paid", f"={X('tax')}")
        put("cfo", f"={X('cash_ops')}-{X('interest_paid')}-{X('loan_fees')}+{X('interest_received')}-{X('tax_paid')}")
        put("cfi", f"=-{X('capex')}")
        put("cff", f"={X('loan_draw')}-{X('loan_repay')}+{X('equity_issue')}-{X('dividends')}")
        put("net_cash", f"={X('cfo')}+{X('cfi')}+{X('cff')}")
        put("cash", f"={X('cash_open')}+{X('net_cash')}")

        put("gross_close", f"={X('gross_open')}+{X('capex')}")
        put("accdep_close", f"={X('accdep_open')}+{X('dep')}")
        put("nbv", f"={X('gross_close')}-{X('accdep_close')}")
        put("share_capital", (f"={R['op']['share_capital']}" if i == 0 else f"={Pv('share_capital')}") + f"+{X('equity_issue')}")
        put("retained", (f"={R['reserves']}" if i == 0 else f"={Pv('retained')}") + f"+{X('pat')}-{X('dividends')}")
        put("total_assets", f"={X('nbv')}+{X('inventory')}+{X('receivables')}+{R['op']['other_assets']}+MAX({X('cash')},0)")
        put("total_liabilities", f"={X('payables')}+{R['op']['other_liabilities']}+{X('borrowings')}+MAX(-{X('cash')},0)")
        put("total_equity", f"={X('share_capital')}+{X('retained')}")
        put("check", f"={X('total_assets')}-{X('total_liabilities')}-{X('total_equity')}")
    ws.freeze_panes = ws.cell(row=rw["label"] + 1, column=FIRST_PERIOD_COL)
    return ws, rw


# --------------------------------------------------------------------------- Annual summary

class _Rows:
    """Tiny row allocator for the presentation sheets."""

    def __init__(self, ws, start):
        self.ws, self.row = ws, start
        self.keys = {}

    def skip(self, n=1):
        self.row += n


def _annual_sheet(wb, a, R, meta, result):
    ws = wb.create_sheet("Annual Summary")
    n = a["years"]
    ncols = 3 + n
    brand.xlsx_title_band(ws, "Annual summary - income statement, balance sheet, cash flow and ratios",
                          _meta_lines(meta), ncols=max(ncols, 8))
    brand.xlsx_page_setup(ws, landscape=False if n <= 4 else True)
    _setw(ws, {"A": 46, "B": 3})
    for i in range(n):
        ws.column_dimensions[CL(3 + i)].width = 14
    rw = R["calc_rw"]
    yrs = f"Calc!$C${rw['year_no']}:${CL(FIRST_PERIOD_COL + R['calc_P'] - 1)}${rw['year_no']}"
    yend = f"Calc!$C${rw['year_end']}:${CL(FIRST_PERIOD_COL + R['calc_P'] - 1)}${rw['year_end']}"
    r = 8
    A = {}  # key -> row on this sheet

    def header():
        nonlocal r
        _head(ws, r, ["", ""] + [E.year_label(a, y) for y in range(1, n + 1)])
        r += 1

    def flow(key):
        return lambda c: f"=SUMIF({yrs},{c}$9y,{R['calc_rng'](key)})"

    def put_row(key, label, builder, fmt=NUM, bold=False, total=False):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        for i in range(n):
            c = CL(3 + i)
            cell = ws.cell(row=r, column=3 + i, value=builder(c, i))
            _fmt(cell, fmt, bold=bold, total=total)
        if bold or total:
            ws.cell(row=r, column=1).font = Font(bold=True)
        if total:
            ws.cell(row=r, column=1).fill = FILL_TOT
            ws.cell(row=r, column=2).fill = FILL_TOT
        A[key] = r
        r += 1

    def F(key):  # annual flow from Calc
        return lambda c, i: f"=SUMIF({yrs},{c}$8,{R['calc_rng'](key)})"

    def B(key):  # year-end balance from Calc
        return lambda c, i: f"=SUMIFS({R['calc_rng'](key)},{yrs},{c}$8,{yend},1)"

    # header with numeric year numbers in row 8 (used by SUMIF)
    ws.cell(row=8, column=1, value="Forecast year no.").font = Font(color=MUTED, size=9)
    for i in range(n):
        c = ws.cell(row=8, column=3 + i, value=i + 1)
        c.font, c.alignment = Font(color=MUTED, size=9), Alignment(horizontal="center")
    r = 9
    header()
    _section(ws, r, "Income statement", ncols)
    r += 1
    for s, st in enumerate(a["revenue"]):
        put_row(f"rev{s}", f"  {st['name']}", F(f"rev{s}"))
    put_row("revenue", "Revenue", F("revenue"), bold=True)
    put_row("cogs", "Cost of sales", F("cogs"))
    put_row("gross_profit", "Gross profit", lambda c, i: f"={c}{A['revenue']}-{c}{A['cogs']}", bold=True, total=True)
    for j, o in enumerate(a["opex"]):
        put_row(f"opex{j}", f"  {o['name']}", F(f"opex{j}"))
    put_row("opex", "Total operating expenses", F("opex"))
    put_row("ebitda", "EBITDA", F("ebitda"), bold=True, total=True)
    put_row("dep", "Depreciation", F("dep"))
    put_row("ebit", "Operating profit (EBIT)", lambda c, i: f"={c}{A['ebitda']}-{c}{A['dep']}", bold=True)
    put_row("int_income", "Interest income", F("int_income"))
    put_row("fin_exp", "Finance costs (loan interest, fees and overdraft)", lambda c, i: f"=SUMIF({yrs},{c}$8,{R['calc_rng']('int_loans')})+SUMIF({yrs},{c}$8,{R['calc_rng']('int_od')})+SUMIF({yrs},{c}$8,{R['calc_rng']('loan_fees')})")
    put_row("pbt", "Profit before tax", lambda c, i: f"={c}{A['ebit']}+{c}{A['int_income']}-{c}{A['fin_exp']}", bold=True)
    put_row("tax", "Income tax", F("tax"))
    put_row("pat", "Profit after tax", lambda c, i: f"={c}{A['pbt']}-{c}{A['tax']}", bold=True, total=True)
    put_row("dividends", "Dividends", F("dividends"))
    r += 1
    _section(ws, r, "Cash flow statement", ncols)
    r += 1
    put_row("cf_ebitda", "EBITDA", F("ebitda"))
    put_row("d_inventory", "(Increase)/decrease in inventory", F("d_inventory"))
    put_row("d_receivables", "(Increase)/decrease in receivables", F("d_receivables"))
    put_row("d_payables", "Increase/(decrease) in payables", F("d_payables"))
    put_row("cash_ops", "Cash generated from operations", F("cash_ops"), bold=True)
    put_row("interest_paid", "Interest paid", lambda c, i: f"=-SUMIF({yrs},{c}$8,{R['calc_rng']('interest_paid')})")
    put_row("fees_paid", "Loan establishment fees paid", lambda c, i: f"=-SUMIF({yrs},{c}$8,{R['calc_rng']('loan_fees')})")
    put_row("interest_received", "Interest received", F("interest_received"))
    put_row("tax_paid", "Tax paid", lambda c, i: f"=-SUMIF({yrs},{c}$8,{R['calc_rng']('tax_paid')})")
    put_row("cfo", "Net cash from operating activities", lambda c, i: f"={c}{A['cash_ops']}+{c}{A['interest_paid']}+{c}{A['fees_paid']}+{c}{A['interest_received']}+{c}{A['tax_paid']}", bold=True, total=True)
    put_row("capex", "Purchase of property, plant and equipment", lambda c, i: f"=-SUMIF({yrs},{c}$8,{R['calc_rng']('capex')})")
    put_row("cfi", "Net cash from investing activities", lambda c, i: f"={c}{A['capex']}", bold=True, total=True)
    put_row("loan_draw", "Loan drawdowns", F("loan_draw"))
    put_row("loan_repay", "Loan repayments", lambda c, i: f"=-SUMIF({yrs},{c}$8,{R['calc_rng']('loan_repay')})")
    put_row("equity_issue", "Share capital issued", F("equity_issue"))
    put_row("div_paid", "Dividends paid", lambda c, i: f"=-{c}{A['dividends']}")
    put_row("cff", "Net cash from financing activities", lambda c, i: f"={c}{A['loan_draw']}+{c}{A['loan_repay']}+{c}{A['equity_issue']}+{c}{A['div_paid']}", bold=True, total=True)
    put_row("net_cash", "Net change in cash", lambda c, i: f"={c}{A['cfo']}+{c}{A['cfi']}+{c}{A['cff']}", bold=True)
    put_row("cash_open", "Cash at start of year", lambda c, i: f"={R['op']['cash']}" if i == 0 else f"={CL(2 + i)}{r + 1}")
    put_row("cash_close", "Cash at end of year", lambda c, i: f"={c}{A['cash_open']}+{c}{A['net_cash']}", bold=True, total=True)
    r += 1
    _section(ws, r, "Balance sheet (year end)", ncols)
    r += 1
    put_row("nbv", "Property, plant and equipment", B("nbv"))
    put_row("inventory", "Inventory", B("inventory"))
    put_row("receivables", "Trade receivables", B("receivables"))
    put_row("other_assets", "Other assets", lambda c, i: f"={R['op']['other_assets']}")
    put_row("cash_pos", "Cash and cash equivalents", lambda c, i: f"=MAX({c}{A['cash_close']},0)")
    put_row("total_assets", "Total assets", lambda c, i: f"=SUM({c}{A['nbv']}:{c}{A['cash_pos']})", bold=True, total=True)
    put_row("payables", "Trade payables", B("payables"))
    put_row("other_liab", "Other liabilities", lambda c, i: f"={R['op']['other_liabilities']}")
    put_row("overdraft", "Bank overdraft", lambda c, i: f"=MAX(-{c}{A['cash_close']},0)")
    put_row("borrowings", "Borrowings", B("borrowings"))
    put_row("total_liabilities", "Total liabilities", lambda c, i: f"=SUM({c}{A['payables']}:{c}{A['borrowings']})", bold=True, total=True)
    put_row("share_capital", "Share capital", B("share_capital"))
    put_row("retained", "Retained earnings & reserves", B("retained"))
    put_row("total_equity", "Total equity", lambda c, i: f"={c}{A['share_capital']}+{c}{A['retained']}", bold=True, total=True)
    put_row("total_el", "Total equity and liabilities", lambda c, i: f"={c}{A['total_liabilities']}+{c}{A['total_equity']}", bold=True)
    put_row("bs_check", "Balance check (must be nil)", lambda c, i: f"=ROUND({c}{A['total_assets']}-{c}{A['total_el']},2)")
    r += 1
    # ---- ratios
    _section(ws, r, "Key ratios", ncols)
    r += 1

    def nxt(c, i, key, last_alt):
        return f"{CL(4 + i)}{A[key]}" if i + 1 < n else last_alt

    put_row("cur_portion", "  Current portion of borrowings (next year's repayments)",
            lambda c, i: f"=MIN({c}{A['borrowings']},-{CL(4 + i)}{A['loan_repay']})" if i + 1 < n else f"=MIN({c}{A['borrowings']},-{c}{A['loan_repay']})")
    put_row("current_assets", "  Current assets", lambda c, i: f"={c}{A['inventory']}+{c}{A['receivables']}+{c}{A['other_assets']}+{c}{A['cash_pos']}")
    put_row("current_liab", "  Current liabilities", lambda c, i: f"={c}{A['payables']}+{c}{A['other_liab']}+{c}{A['overdraft']}+{c}{A['cur_portion']}")
    put_row("net_debt", "  Net debt", lambda c, i: f"={c}{A['borrowings']}+{c}{A['overdraft']}-{c}{A['cash_pos']}")
    put_row("debt_service", "  Debt service (interest paid + principal repaid)", lambda c, i: f"=-{c}{A['interest_paid']}-{c}{A['fees_paid']}-{c}{A['loan_repay']}")
    put_row("r_current", "Current ratio", lambda c, i: f'=IFERROR({c}{A["current_assets"]}/{c}{A["current_liab"]},"n/a")', MULT)
    put_row("r_quick", "Quick ratio", lambda c, i: f'=IFERROR(({c}{A["receivables"]}+{c}{A["cash_pos"]})/{c}{A["current_liab"]},"n/a")', MULT)
    put_row("r_gm", "Gross margin", lambda c, i: f'=IFERROR({c}{A["gross_profit"]}/{c}{A["revenue"]},"n/a")', PCT)
    put_row("r_ebitda", "EBITDA margin", lambda c, i: f'=IFERROR({c}{A["ebitda"]}/{c}{A["revenue"]},"n/a")', PCT)
    put_row("r_net", "Net profit margin", lambda c, i: f'=IFERROR({c}{A["pat"]}/{c}{A["revenue"]},"n/a")', PCT)
    put_row("r_roe", "Return on equity", lambda c, i: f'=IFERROR({c}{A["pat"]}/{c}{A["total_equity"]},"n/a")', PCT)
    put_row("r_roce", "Return on capital employed", lambda c, i: f'=IFERROR({c}{A["ebit"]}/({c}{A["total_equity"]}+{c}{A["borrowings"]}),"n/a")', PCT)
    put_row("r_de", "Gearing (debt / equity)", lambda c, i: f'=IFERROR(({c}{A["borrowings"]}+{c}{A["overdraft"]})/{c}{A["total_equity"]},"n/a")', MULT)
    put_row("r_nd", "Net debt / EBITDA", lambda c, i: f'=IF({c}{A["ebitda"]}>0,{c}{A["net_debt"]}/{c}{A["ebitda"]},"n/a")', MULT)
    put_row("r_ic", "Interest cover (EBIT / interest)", lambda c, i: f'=IFERROR({c}{A["ebit"]}/({c}{A["fin_exp"]}-SUMIF({yrs},{c}$8,{R["calc_rng"]("loan_fees")})),"n/a")', MULT)
    put_row("r_dscr", "Debt service cover ((EBITDA - tax) / debt service)", lambda c, i: f'=IFERROR(({c}{A["ebitda"]}-{c}{A["tax"]})/{c}{A["debt_service"]},"n/a")', MULT)
    put_row("r_rec", "Receivable days", lambda c, i: f'=IFERROR({c}{A["receivables"]}*365/{c}{A["revenue"]},"n/a")', DAYS)
    put_row("r_inv", "Inventory days", lambda c, i: f'=IFERROR({c}{A["inventory"]}*365/{c}{A["cogs"]},"n/a")', DAYS)
    put_row("r_turn", "Asset turnover", lambda c, i: f'=IFERROR({c}{A["revenue"]}/{c}{A["total_assets"]},"n/a")', MULT)
    r += 1
    # ---- covenants
    cov = a["covenants"]
    defs = [("min_dscr", "Debt service cover - minimum", "r_dscr", ">="), ("min_interest_cover", "Interest cover - minimum", "r_ic", ">="),
            ("max_debt_equity", "Gearing (debt / equity) - maximum", "r_de", "<="), ("min_current_ratio", "Current ratio - minimum", "r_current", ">="),
            ("max_net_debt_ebitda", "Net debt / EBITDA - maximum", "r_nd", "<=")]
    live = [d for d in defs if cov.get(d[0]) is not None]
    if live:
        _section(ws, r, "Covenant tests (limits are inputs)", ncols)
        r += 1
        for key, label, rkey, op in live:
            ws.cell(row=r, column=1, value=label)
            _inp(ws.cell(row=r, column=2), cov[key], '0.00')
            ws.column_dimensions["B"].width = 8
            for i in range(n):
                c = CL(3 + i)
                ref = f"{c}{A[rkey]}"
                cell = ws.cell(row=r, column=3 + i, value=f'=IF(ISNUMBER({ref}),IF({ref}{op}$B{r},"Pass","BREACH"),"Pass")')
                cell.alignment = Alignment(horizontal="center")
            r += 1
        r += 1
    ws.freeze_panes = "C10"
    R["annual_rows"] = A
    return ws, A


# --------------------------------------------------------------------------- Monthly year 1

def _monthly_sheet(wb, a, R, meta):
    ws = wb.create_sheet("Monthly Year 1")
    rw = R["calc_rw"]
    brand.xlsx_title_band(ws, "Year 1 by month", _meta_lines(meta), ncols=16)
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 42, "B": 3})
    for i in range(14):
        ws.column_dimensions[CL(3 + i)].width = 12
    r = 8
    _head(ws, r, ["", ""] + [f"=Calc!{CL(FIRST_PERIOD_COL + m)}{rw['label']}" for m in range(12)] + ["Year 1"])
    r += 1
    spec = [("Income statement", None, None, None),
            ("Revenue", "revenue", 1, "flow"), ("Cost of sales", "cogs", 1, "flow"), ("Operating expenses", "opex", 1, "flow"),
            ("EBITDA", "ebitda", 1, "flow"), ("Depreciation", "dep", 1, "flow"), ("Net finance cost", "fin_cost", 1, "flow"),
            ("Profit before tax", "pbt", 1, "flow"), ("Tax", "tax", 1, "flow"), ("Profit after tax", "pat", 1, "flow"),
            ("Cash flow", None, None, None),
            ("Net cash from operating activities", "cfo", 1, "flow"), ("Capital expenditure", "cfi", 1, "flow"),
            ("Net cash from financing activities", "cff", 1, "flow"), ("Net change in cash", "net_cash", 1, "flow"),
            ("Cash at end of month", "cash", 1, "bal"),
            ("Balance sheet", None, None, None),
            ("Property, plant and equipment", "nbv", 1, "bal"), ("Inventory", "inventory", 1, "bal"), ("Trade receivables", "receivables", 1, "bal"),
            ("Trade payables", "payables", 1, "bal"), ("Borrowings", "borrowings", 1, "bal"), ("Total equity", "total_equity", 1, "bal"),
            ("Balance check (nil)", "check", 1, "bal")]
    for label, key, _, kind in spec:
        if key is None:
            _section(ws, r, label, 15)
            r += 1
            continue
        ws.cell(row=r, column=1, value=label)
        for m in range(12):
            cell = ws.cell(row=r, column=3 + m, value=f"=Calc!{CL(FIRST_PERIOD_COL + m)}{rw[key]}")
            _fmt(cell, NUM)
        total = ws.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})" if kind == "flow" else f"=N{r}")
        _fmt(total, NUM, bold=True)
        r += 1
    ws.freeze_panes = "C9"
    return ws



# --------------------------------------------------------------------------- Cashflow forecast sheet

def _cashflow_forecast_sheet(wb, a, R, meta):
    """The firm's cash flow forecast layout with live formulas: year 1 by month, a Year 1 column and the later years. Every
    line reads the Calc sheet, and the closing cash is reconciled to the cash flow statement and the balance sheet cash."""
    ws = wb.create_sheet("Cashflow forecast")
    rw = R["calc_rw"]
    years = a["years"]
    ncol = 12 + 1 + (years - 1)
    brand.xlsx_title_band(ws, "Cash flow forecast", _meta_lines(meta) + ["Year 1 by month, then annual - reconciled to the cash flow statement"], ncols=min(ncol + 2, 30))
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 50, "B": 3})
    for i in range(ncol):
        ws.column_dimensions[CL(3 + i)].width = 12.5
    # column kinds: ('m', calc_col) for months, ('y1', None), ('y', calc_col) for later years
    cols = [("m", CL(FIRST_PERIOD_COL + m)) for m in range(12)] + [("y1", None)] + [("y", CL(FIRST_PERIOD_COL + 12 + (y - 2))) for y in range(2, years + 1)]
    r = 8
    _head(ws, r, ["", ""] + [f"=Calc!{c}{rw['label']}" if k != "y1" else "Year 1" for k, c in cols])
    r += 1
    first_col, last_col = CL(3), CL(14)
    rows = {}

    def put(label, key, build, kind="flow", style="line", fmt=NUM, bold=False):
        """build(calc_col_letter) -> formula text without '='. kind: flow (sum for Year 1) | open | close | none."""
        nonlocal r
        ws.cell(row=r, column=1, value=("    " if style == "indent" else "") + label)
        for i, (k, c) in enumerate(cols):
            cell = ws.cell(row=r, column=3 + i)
            if k == "y1":
                if kind == "flow":
                    cell.value = f"=SUM({first_col}{r}:{last_col}{r})"
                elif kind == "open":
                    cell.value = f"={first_col}{r}"
                elif kind == "close":
                    cell.value = f"={last_col}{r}"
                else:
                    continue
            else:
                f = build(c)
                if f is None:
                    continue
                cell.value = "=" + f
            _fmt(cell, fmt, bold=bold or style in ("total", "grand"), total=style in ("total", "grand"))
        if bold or style in ("total", "grand"):
            ws.cell(row=r, column=1).font = Font(bold=True)
        if key:
            rows[key] = r
        r += 1
        return r - 1

    def calc(k):
        return lambda c: f"Calc!{c}{rw[k]}"

    def neg(k):
        return lambda c: f"-Calc!{c}{rw[k]}"

    def gap():
        nonlocal r
        r += 1

    sums = []   # rows whose Year-1-or-period values add up to the net movement
    put("Opening cash balance", "open", calc("cash_open"), "open", "total")
    gap()
    if a["loans"]:
        sums.append(put("Loan finance", "loan", calc("loan_draw"), "flow", "total"))
        for k, L in enumerate(a["loans"]):
            put(L["name"], None, calc(f"ln{k}_draw"), "flow", "indent")
        gap()
    ws.cell(row=r, column=1, value="Revenue distribution during year 1").font = Font(italic=True, color=MUTED)
    for i, (k, c) in enumerate(cols):
        if k == "m":
            cell = ws.cell(row=r, column=3 + i, value=f"=Calc!{c}{rw['share']}")
            _fmt(cell, PCT)
        elif k == "y1":
            cell = ws.cell(row=r, column=3 + i, value=f"=SUM({first_col}{r}:{last_col}{r})")
            _fmt(cell, PCT)
    r += 1
    sums.append(put("Revenue", "rev", calc("revenue"), "flow", "total"))
    for i, st in enumerate(a["revenue"]):
        put(st["name"], None, calc(f"rev{i}"), "flow", "indent")
    gap()
    sums.append(put("Cost of sales", "cogs", neg("cogs"), "flow", "total"))
    gap()
    sums.append(put("Operating expenditure", "opex", neg("opex"), "flow", "total"))
    for j, o in enumerate(a["opex"]):
        put(o["name"], None, neg(f"opex{j}"), "flow", "indent")
    gap()
    sums.append(put("Working capital movements", "wc", lambda c: f"Calc!{c}{rw['d_inventory']}+Calc!{c}{rw['d_receivables']}+Calc!{c}{rw['d_payables']}", "flow", "total"))
    put("Change in inventory", None, calc("d_inventory"), "flow", "indent")
    put("Change in receivables", None, calc("d_receivables"), "flow", "indent")
    put("Change in payables", None, calc("d_payables"), "flow", "indent")
    gap()
    sums.append(put("Finance charges", "fin", lambda c: f"-Calc!{c}{rw['loan_fees']}-Calc!{c}{rw['loan_repay']}-Calc!{c}{rw['int_loans']}-Calc!{c}{rw['int_od']}+Calc!{c}{rw['int_income']}", "flow", "total"))
    for k, L in enumerate(a["loans"]):
        fee_txt = []
        if L["fee_pct"] > 0:
            fee_txt.append(f"{L['fee_pct']:g}%")
        if L["fee_amount"] > 0:
            fee_txt.append(f"{L['fee_amount']:,.0f} fixed")
        put(f"Establishment fee @ {' + '.join(fee_txt) or '0%'} - {L['name']}", None, neg(f"ln{k}_fee"), "flow", "indent")
    for k, L in enumerate(a["loans"]):
        put(f"Loan repayment - principal - {L['name']}", None, neg(f"ln{k}_repay"), "flow", "indent")
    for k, L in enumerate(a["loans"]):
        put(f"Loan repayment - interest @ {L['rate']:g}% - {L['name']}", None, neg(f"ln{k}_int"), "flow", "indent")
    put("Overdraft interest", None, neg("int_od"), "flow", "indent")
    put("Interest received", None, calc("int_income"), "flow", "indent")
    gap()
    sums.append(put("Taxation", "tax", neg("tax"), "flow", "total"))
    gap()
    sums.append(put("Capital expenditure", "capex", neg("capex"), "flow", "total"))
    gap()
    sums.append(put("Shareholders", "own", lambda c: f"Calc!{c}{rw['equity_issue']}-Calc!{c}{rw['dividends']}", "flow", "total"))
    put("Share capital issued", None, calc("equity_issue"), "flow", "indent")
    put("Dividends paid", None, neg("dividends"), "flow", "indent")
    gap()

    def net_formula(c):
        return None

    # net movement = sum of the section totals on this sheet
    net_r = r
    ws.cell(row=net_r, column=1, value="Net cash movement").font = Font(bold=True)
    for i, (k, c) in enumerate(cols):
        col = CL(3 + i)
        cell = ws.cell(row=net_r, column=3 + i, value="=" + "+".join(f"{col}{x}" for x in sums))
        _fmt(cell, NUM, bold=True, total=True)
    r += 1
    close_r = r
    ws.cell(row=close_r, column=1, value="Closing cash balance").font = Font(bold=True)
    for i, (k, c) in enumerate(cols):
        col = CL(3 + i)
        cell = ws.cell(row=close_r, column=3 + i, value=f"={col}{rows['open']}+{col}{net_r}")
        _fmt(cell, NUM, bold=True, total=True)
    r += 2
    _section(ws, r, "Reconciliation to the cash flow statement and balance sheet", ncol + 2)
    r += 1
    stmt = {}
    for label, key in (("Net cash from operating activities", "cfo"), ("Net cash from investing activities", "cfi"), ("Net cash from financing activities", "cff")):
        stmt[key] = put(label, None, calc(key), "flow", "indent")
    sn = put("Net change in cash per the cash flow statement", None, calc("net_cash"), "flow", "total")
    dn = put("Difference: forecast less cash flow statement (nil)", None, lambda c: None, "none", "line")
    for i, (k, c) in enumerate(cols):
        col = CL(3 + i)
        _fmt(ws.cell(row=dn, column=3 + i, value=f"=ROUND({col}{net_r}-{col}{sn},2)"), NUM)
    sc = put("Closing cash per the balance sheet (cash less overdraft)", None, calc("cash"), "close", "total")
    dc = put("Difference: closing cash less balance sheet (nil)", None, lambda c: None, "none", "line")
    for i, (k, c) in enumerate(cols):
        col = CL(3 + i)
        _fmt(ws.cell(row=dc, column=3 + i, value=f"=ROUND({col}{close_r}-{col}{sc},2)"), NUM)
    r += 1
    last = CL(2 + ncol)
    ws.cell(row=r, column=1, value="Reconciliation status").font = Font(bold=True)
    c = ws.cell(row=r, column=3, value=f'=IF(AND(SUMPRODUCT(ABS({first_col}{dn}:{last}{dn}))<0.01,SUMPRODUCT(ABS({first_col}{dc}:{last}{dc}))<0.01),"Reconciles","DOES NOT RECONCILE")')
    c.font = Font(bold=True)
    rows["status"] = r
    rows["net"], rows["close"], rows["diff_net"], rows["diff_close"] = net_r, close_r, dn, dc
    ws.freeze_panes = "C9"
    R["cff_rows"] = rows
    return ws


# --------------------------------------------------------------------------- Valuation sheet

def _valuation_sheet(wb, a, R, A, meta, result, val):
    ws = wb.create_sheet("Valuation")
    v = T.normalise_valuation(val)
    n = a["years"]
    ncols = max(3 + n, 9)
    brand.xlsx_title_band(ws, "Valuation - discounted cash flow and multiples", _meta_lines(meta), ncols=ncols)
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 46, "B": 14})
    for i in range(n + 2):
        ws.column_dimensions[CL(3 + i)].width = 14
    r = 8
    W = {}

    def inp(label, key, value, fmt=PCT_IN, note=None):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        _inp(ws.cell(row=r, column=2), value, fmt)
        if note:
            ws.cell(row=r, column=3, value=note).font = Font(size=8, color=MUTED)
        W[key] = f"$B${r}"
        r += 1

    def calc(label, key, formula, fmt=NUM, bold=False):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=formula)
        _fmt(c, fmt, bold=bold)
        if bold:
            ws.cell(row=r, column=1).font = Font(bold=True)
        W[key] = f"$B${r}"
        r += 1

    w = v["wacc"]
    _section(ws, r, "Discount rate (WACC)", ncols)
    r += 1
    inp("Risk-free rate (%)", "rf", w["rf"])
    inp("Equity risk premium (%)", "erp", w["erp"])
    inp("Beta", "beta", w["beta"], '0.00')
    inp("Country risk premium (%)", "crp", w["crp"])
    inp("Size premium (%)", "size", w["size"])
    inp("Company-specific risk premium (%)", "specific", w["specific"])
    inp("Pre-tax cost of debt (%)", "kd", w["kd"])
    inp("Debt / (debt + equity) (%)", "dw", w["debt_weight"])
    ws.cell(row=r, column=1, value="Tax rate (%) - from Assumptions")
    ws.cell(row=r, column=2, value=f"={R['tax']}").number_format = PCT_IN
    W["tax"] = f"$B${r}"
    r += 1
    calc("Cost of equity (%)", "ke", f"={W['rf']}+{W['beta']}*{W['erp']}+{W['crp']}+{W['size']}+{W['specific']}", '0.00')
    calc("After-tax cost of debt (%)", "kd_at", f"={W['kd']}*(1-{W['tax']}/100)", '0.00')
    calc("WACC calculated (%)", "wacc_calc", f"={W['ke']}*(1-{W['dw']}/100)+{W['kd_at']}*{W['dw']}/100", '0.00')
    ws.cell(row=r, column=1, value="WACC override (%) - leave blank to use calculated")
    _inp(ws.cell(row=r, column=2), w["override"], PCT_IN)
    W["override"] = f"$B${r}"
    r += 1
    calc("WACC used (%)", "wacc", f'=IF(ISNUMBER({W["override"]}),{W["override"]},{W["wacc_calc"]})', '0.00', bold=True)
    r += 1
    _section(ws, r, "Terminal value and bridge", ncols)
    r += 1
    inp("Terminal growth rate (%)", "g", v["terminal"]["growth"])
    inp("Exit multiple (EV / EBITDA, x)", "mult", v["terminal"]["exit_multiple"], '0.00')
    inp("Mid-year discounting (1 = yes, 0 = no)", "mid", 1 if v["mid_year"] else 0, "0")
    ws.cell(row=r, column=1, value="Terminal value method used")
    ws.cell(row=r, column=2, value="Gordon growth" if v["terminal"]["method"] == "gordon" else "Exit multiple")
    W["method"] = v["terminal"]["method"]
    r += 1
    debt_f = f"={R['loans_open']}+MAX(-{R['op']['cash']},0)"
    cash_f = f"=MAX({R['op']['cash']},0)"
    ws.cell(row=r, column=1, value="Debt at valuation date")
    if v["bridge"]["debt"] is None:
        c = ws.cell(row=r, column=2, value=debt_f)
        _fmt(c, NUM)
    else:
        _inp(ws.cell(row=r, column=2), v["bridge"]["debt"], NUM)
    W["debt"] = f"$B${r}"
    r += 1
    ws.cell(row=r, column=1, value="Cash at valuation date")
    if v["bridge"]["cash"] is None:
        c = ws.cell(row=r, column=2, value=cash_f)
        _fmt(c, NUM)
    else:
        _inp(ws.cell(row=r, column=2), v["bridge"]["cash"], NUM)
    W["cash"] = f"$B${r}"
    r += 1
    inp("Other adjustments (non-operating assets, +)", "other", v["bridge"]["other"], NUM)
    inp("Minority interests (-)", "minorities", v["bridge"]["minorities"], NUM)
    inp("Stake being valued (%)", "stake", v["stake_pct"])
    inp("Discount for lack of marketability (%)", "dlom", v["dlom"])
    r += 1

    # ---- FCFF table
    _section(ws, r, "Unlevered free cash flow", ncols)
    r += 1
    hdr = r
    _head(ws, r, ["", ""] + [E.year_label(a, y) for y in range(1, n + 1)])
    r += 1
    rows = {}

    def line(key, label, builder, fmt=NUM, bold=False):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        for i in range(n):
            col = CL(3 + i)
            c = ws.cell(row=r, column=3 + i, value=builder(col, i))
            _fmt(c, fmt, bold=bold)
        if bold:
            ws.cell(row=r, column=1).font = Font(bold=True)
        rows[key] = r
        r += 1

    def AS(key, col):
        return f"'Annual Summary'!{col}{A[key]}"

    line("ebit", "EBIT", lambda c, i: f"={AS('ebit', c)}")
    line("tax", "Tax on EBIT", lambda c, i: f"=MAX({c}{rows['ebit']},0)*{W['tax']}/100")
    line("nopat", "NOPAT", lambda c, i: f"={c}{rows['ebit']}-{c}{rows['tax']}")
    line("dep", "Add back: depreciation", lambda c, i: f"={AS('dep', c)}")
    line("capex", "Less: capital expenditure", lambda c, i: f"={AS('capex', c)}")
    line("nwc", "Change in working capital", lambda c, i: f"={AS('d_inventory', c)}+{AS('d_receivables', c)}+{AS('d_payables', c)}")
    line("fcff", "Free cash flow to the firm", lambda c, i: f"={c}{rows['nopat']}+{c}{rows['dep']}+{c}{rows['capex']}+{c}{rows['nwc']}", bold=True)
    line("t", "Discount period (years)", lambda c, i: f"={i + 1}-0.5*{W['mid']}", '0.0')
    line("df", "Discount factor", lambda c, i: f"=1/(1+{W['wacc']}/100)^{c}{rows['t']}", '0.0000')
    line("pv", "Present value", lambda c, i: f"={c}{rows['fcff']}*{c}{rows['df']}", bold=True)
    line("ebitda", "EBITDA (for the exit multiple)", lambda c, i: f"={AS('ebitda', c)}")
    r += 1
    first, last = CL(3), CL(2 + n)
    _section(ws, r, "Enterprise and equity value", ncols)
    r += 1
    calc("Sum of present values of cash flows", "sum_pv", f"=SUM({first}{rows['pv']}:{last}{rows['pv']})")
    calc("Terminal value - Gordon growth", "tv_g", f'=IF({W["wacc"]}>{W["g"]},{last}{rows["fcff"]}*(1+{W["g"]}/100)/(({W["wacc"]}-{W["g"]})/100),"n/a")')
    calc("PV of terminal value - Gordon growth", "pv_tv_g", f'=IF(ISNUMBER({W["tv_g"]}),{W["tv_g"]}/(1+{W["wacc"]}/100)^({n}-0.5*{W["mid"]}),"n/a")')
    calc("Terminal value - exit multiple", "tv_x", f"={last}{rows['ebitda']}*{W['mult']}")
    calc("PV of terminal value - exit multiple", "pv_tv_x", f"={W['tv_x']}/(1+{W['wacc']}/100)^{n}")
    if W["method"] == "exit_multiple":
        calc("PV of terminal value used", "pv_tv", f"={W['pv_tv_x']}")
    else:
        calc("PV of terminal value used", "pv_tv", f"=IF(ISNUMBER({W['pv_tv_g']}),{W['pv_tv_g']},{W['pv_tv_x']})")
    calc("Enterprise value", "ev", f"={W['sum_pv']}+{W['pv_tv']}", NUM, bold=True)
    calc("Less: debt", "l_debt", f"=-{W['debt']}")
    calc("Add: cash", "l_cash", f"={W['cash']}")
    calc("Add: other adjustments", "l_other", f"={W['other']}")
    calc("Less: minority interests", "l_min", f"=-{W['minorities']}")
    calc("Equity value (100%)", "equity", f"={W['ev']}+{W['l_debt']}+{W['l_cash']}+{W['l_other']}+{W['l_min']}", NUM, bold=True)
    calc("Equity value of the stake", "stake_val", f"={W['equity']}*{W['stake']}/100")
    calc("Equity value of the stake after marketability discount", "after_dlom", f"={W['stake_val']}*(1-{W['dlom']}/100)", NUM, bold=True)
    calc("Terminal value as % of enterprise value", "tv_share", f"=IFERROR({W['pv_tv']}/{W['ev']},0)", PCT.replace("0.0%", "0.0%"))
    ws[W["tv_share"].replace("$", "")].number_format = '0.0"%"' if False else PCT
    r += 1

    # ---- sensitivity (live)
    _section(ws, r, "Sensitivity - enterprise value", ncols)
    r += 1
    exit_mode = W["method"] == "exit_multiple"
    ws.cell(row=r, column=1, value="WACC (down) vs " + ("exit multiple (across)" if exit_mode else "terminal growth (across)")).font = Font(bold=True)
    cols = [(v["terminal"]["exit_multiple"] + s) if exit_mode else (v["terminal"]["growth"] + s)
            for s in ((-1, -0.5, 0, 0.5, 1) if exit_mode else (-1, -0.5, 0, 0.5, 1))]
    for j, cv in enumerate(cols):
        c = ws.cell(row=r, column=3 + j, value=f"={W['mult' if exit_mode else 'g']}+({cv}-{v['terminal']['exit_multiple' if exit_mode else 'growth']})")
        c.font, c.fill, c.number_format = Font(bold=True, color="FFFFFF"), FILL_HEAD, '0.00'
    top = r
    r += 1
    fc_rng = f"${first}${rows['fcff']}:${last}${rows['fcff']}"
    t_rng = f"${first}${rows['t']}:${last}${rows['t']}"
    for s in (-2, -1, 0, 1, 2):
        wc = ws.cell(row=r, column=2, value=f"={W['wacc']}+({s})")
        wc.number_format, wc.font = '0.00', Font(bold=True)
        for j in range(5):
            col = CL(3 + j)
            wv, gv = f"$B{r}", f"{col}${top}"
            if exit_mode:
                f_ = (f"=SUMPRODUCT({fc_rng}/(1+{wv}/100)^{t_rng})+${last}${rows['ebitda']}*{gv}/(1+{wv}/100)^{n}")
            else:
                f_ = (f'=IF({wv}>{gv},SUMPRODUCT({fc_rng}/(1+{wv}/100)^{t_rng})+${last}${rows["fcff"]}*(1+{gv}/100)/(({wv}-{gv})/100)/(1+{wv}/100)^({n}-0.5*{W["mid"]}),"n/a")')
            _fmt(ws.cell(row=r, column=3 + j, value=f_), NUM)
        r += 1
    r += 1

    # ---- multiples
    _section(ws, r, "Trading multiples (enter comparable companies)", ncols)
    r += 1
    _head(ws, r, ["Comparable", "", "EV / EBITDA", "EV / Revenue", "P / E"])
    r += 1
    comps = v["comparables"] or [{"name": "", "ev_ebitda": None, "ev_revenue": None, "pe": None}] * 3
    c0 = r
    for c in comps:
        _inp(ws.cell(row=r, column=1), c["name"])
        for j, k in enumerate(("ev_ebitda", "ev_revenue", "pe")):
            _inp(ws.cell(row=r, column=3 + j), c[k], '0.00')
        r += 1
    c1 = r - 1
    stats = {}
    for label, fn in (("Low", "MIN"), ("Median", "MEDIAN"), ("Mean", "AVERAGE"), ("High", "MAX")):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        for j in range(3):
            col = CL(3 + j)
            _fmt(ws.cell(row=r, column=3 + j, value=f'=IFERROR({fn}({col}{c0}:{col}{c1}),"n/a")'), '0.00', bold=True)
        stats[label] = r
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Earnings base (year 1 forecast)").font = Font(bold=True)
    for j, key in enumerate(("ebitda", "revenue", "pat")):
        _fmt(ws.cell(row=r, column=3 + j, value=f"='Annual Summary'!{first}{A[key]}"), NUM)
    base_row = r
    r += 1
    _head(ws, r, ["Equity value from multiples", "", "EV / EBITDA", "EV / Revenue", "P / E"])
    r += 1
    for label in ("Low", "Median", "High"):
        ws.cell(row=r, column=1, value=label)
        for j in range(3):
            col = CL(3 + j)
            core = f"{col}{stats[label]}*{col}{base_row}"
            f_ = (f'=IFERROR({core}+{W["other"]}-{W["minorities"]},"n/a")' if j == 2
                  else f'=IFERROR({core}-{W["debt"]}+{W["cash"]}+{W["other"]}-{W["minorities"]},"n/a")')
            _fmt(ws.cell(row=r, column=3 + j, value=f_), NUM)
        r += 1
    R["valuation_cells"] = W
    return ws


# --------------------------------------------------------------------------- snapshot sheets

def _snapshot_sheet(wb, a, meta, scenarios, sens):
    """Scenario comparison and sensitivity grids (values at the export date - the engine runs them)."""
    ws = wb.create_sheet("Scenarios")
    brand.xlsx_title_band(ws, "Scenario comparison and sensitivity", _meta_lines(meta) + [f"Values as at {date.today():%d %b %Y} (re-run in the app to refresh)"], ncols=10)
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 44})
    for i in range(2, 12):
        ws.column_dimensions[CL(i)].width = 16
    r = 8
    if scenarios:
        _section(ws, r, "Scenarios", 10)
        r += 1
        _head(ws, r, ["Measure"] + [s["name"] for s in scenarios])
        r += 1
        rows = [("Revenue - year 1", "revenue_y1"), ("Revenue - final year", "revenue_last"), ("EBITDA - year 1", "ebitda_y1"),
                ("EBITDA - final year", "ebitda_last"), ("Total profit after tax", "pat_total"), ("Closing cash - final year", "cash_last"),
                ("Lowest cash balance", "cash_min"), ("Borrowings - final year", "borrowings_last"), ("Equity - final year", "equity_last")]
        for label, key in rows:
            ws.cell(row=r, column=1, value=label)
            for j, s in enumerate(scenarios):
                _fmt(ws.cell(row=r, column=2 + j, value=s["summary"].get(key)), NUM)
            r += 1
        ws.cell(row=r, column=1, value="Lowest debt service cover")
        for j, s in enumerate(scenarios):
            _fmt(ws.cell(row=r, column=2 + j, value=s["summary"].get("dscr_min")), MULT)
        r += 2
    for grid in sens or []:
        _section(ws, r, grid["title"], 10)
        r += 1
        _head(ws, r, [grid["y_label"] + " (down) / " + grid["x_label"] + " (across)"] + [str(x) for x in grid["x_values"]])
        r += 1
        for yv, line in zip(grid["y_values"], grid["rows"]):
            ws.cell(row=r, column=1, value=yv)
            for j, val in enumerate(line):
                _fmt(ws.cell(row=r, column=2 + j, value=val), NUM)
            r += 1
        r += 1
    return ws


def _bva_sheet(wb, a, meta, result, actuals):
    bva = T.budget_vs_actual(result, actuals)
    if not bva["n"]:
        return None
    ws = wb.create_sheet("Budget vs Actual")
    brand.xlsx_title_band(ws, "Budget vs actual", _meta_lines(meta), ncols=9)
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 30, "B": 22})
    for i in range(3, 9):
        ws.column_dimensions[CL(i)].width = 15
    r = 8
    for p in bva["periods"] + [{"label": "Cumulative to date", "lines": bva["cumulative"]}]:
        _section(ws, r, p["label"], 8)
        r += 1
        _head(ws, r, ["Line", "", "Budget", "Actual", "Variance", "Variance %", "Result"])
        r += 1
        for ln in p["lines"]:
            ws.cell(row=r, column=1, value=ln["label"])
            if ln["actual"] is None:
                r += 1
                continue
            _fmt(ws.cell(row=r, column=3, value=ln["budget"]), NUM)
            _inp(ws.cell(row=r, column=4), ln["actual"], NUM)
            _fmt(ws.cell(row=r, column=5, value=f"=D{r}-C{r}"), NUM)
            _fmt(ws.cell(row=r, column=6, value=f'=IF(ABS(C{r})>0.000001,E{r}/ABS(C{r}),"n/a")'), PCT)
            sign = ">=" if ln["kind"] == "income" else "<="
            ws.cell(row=r, column=7, value=f'=IF(ABS(E{r})<0.005,"On budget",IF(E{r}{sign}0,"Favourable","Adverse"))')
            r += 1
        r += 1
    return ws


def _cover_sheet(wb, a, meta, R, A):
    ws = wb.active
    ws.title = "Cover"
    brand.xlsx_title_band(ws, meta.get("name") or "Financial forecast model", _meta_lines(meta), ncols=8)
    brand.xlsx_page_setup(ws, landscape=False)
    _setw(ws, {"A": 34, "B": 60})
    r = 8
    for k, v in (("Model reference", meta.get("number")), ("Client", meta.get("client")), ("Prepared by", meta.get("prepared_by")),
                 ("Status", meta.get("status")), ("Currency", a["currency"]), ("Forecast", f"{a['years']} year(s) from {a['start_month']}"),
                 ("Generated", f"{date.today():%d %B %Y}")):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, color=NAVY)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    rw = R["calc_rw"]
    chk = R["calc_rng"]("check")
    ws.cell(row=r, column=1, value="Model integrity").font = Font(bold=True, color=NAVY)
    ws.cell(row=r, column=2, value=f'=IF(SUMPRODUCT(ABS({chk}))<0.01,"OK - balance sheet balances in every period","ERROR - balance sheet does not balance")')
    r += 2
    ws.cell(row=r, column=1, value="Sheets").font = Font(bold=True, color=NAVY)
    r += 1
    for name, text in (("Assumptions", "All inputs (blue on yellow)"), ("Annual Summary", "Income statement, cash flow, balance sheet, ratios, covenants"),
                       ("Monthly Year 1", "Year 1 month by month"), ("Cashflow forecast", "Cash flow forecast in the firm layout, reconciled to the statement"), ("Calc", "The calculation engine (monthly year 1, annual after)"),
                       ("Valuation", "WACC, DCF, sensitivity and trading multiples"), ("Scenarios", "Scenario comparison and sensitivity snapshots")):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=text)
        r += 1
    r += 1
    for line in ("Convention: blue figures on a pale-yellow fill are inputs; everything else is a formula.",
                 "The model is VAT-exclusive. Interest is charged on opening balances; tax and dividends are paid in the period they arise.",
                 "Working capital balances use annualised run-rates and days / 365. Opening reserves are the balancing figure of the opening balance sheet.",
                 "Number of years, revenue streams, expense lines, loans and repayment types are fixed in this export; re-export from the app to change them.",
                 "Forecasts are estimates, not predictions, and depend entirely on the assumptions entered."):
        c = ws.cell(row=r, column=1, value=line)
        c.font = Font(size=9, color=MUTED)
        r += 1
    return ws


# --------------------------------------------------------------------------- public entry points

def build_model_workbook(meta, a, valuation=None, scenarios=None, sensitivity=None, actuals=None):
    """Return bytes of the formula workbook for one model. `scenarios` = [{'name','summary'}...] (optional)."""
    a = E.normalise(a)
    result = E.run(a)
    wb = Workbook()
    _cover_sheet_placeholder = wb.active
    periods = E.build_periods(a)
    ws_a, R = _assumptions_sheet(wb, a, meta)
    ws_c, rw = _calc_sheet(wb, a, R, meta, periods)
    ws_s, A = _annual_sheet(wb, a, R, meta, result)
    _monthly_sheet(wb, a, R, meta)
    _cashflow_forecast_sheet(wb, a, R, meta)
    _valuation_sheet(wb, a, R, A, meta, result, valuation)
    if scenarios or sensitivity:
        _snapshot_sheet(wb, a, meta, scenarios, sensitivity)
    if actuals:
        _bva_sheet(wb, a, meta, result, actuals)
    _cover_sheet(wb, a, meta, R, A)
    order = ["Cover", "Assumptions", "Annual Summary", "Monthly Year 1", "Cashflow forecast", "Valuation", "Scenarios", "Budget vs Actual", "Calc"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title in ("Cover", "Assumptions") else GOLD
    wb["Assumptions"].sheet_properties.tabColor = "0000FF"
    wb.properties.creator = "Neverlank Chartered Accountants"
    wb.properties.title = meta.get("name") or "Financial forecast model"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {"calc_rw": rw, "annual_rows": A, "valuation_cells": R.get("valuation_cells"), "cff_rows": R.get("cff_rows"), "calc_P": R["calc_P"]}


def build_statements_workbook(meta, result):
    """Plain-values workbook of the statements (annual, then year 1 by month) and the cash flow forecast, for circulating or
    pasting into other documents. The live-formula model workbook is a separate download."""
    import cf_report as RP
    wb = Workbook()
    wb.remove(wb.active)

    def write(ws, r, columns, lines, ncols):
        _head(ws, r, [""] + list(columns))
        r += 1
        for ln in lines:
            kind = ln["kind"]
            if kind == "gap":
                r += 1
                continue
            if kind == "head":
                _section(ws, r, ln["label"], ncols)
                r += 1
                continue
            ws.cell(row=r, column=1, value=("    " if kind == "indent" else "") + ln["label"])
            fmt = {"pct": PCT, "x": '0.00"x"', "days": "0"}.get(ln.get("fmt", "money"), NUM)
            for i, v in enumerate(ln["values"]):
                if v is None:
                    continue
                _fmt(ws.cell(row=r, column=2 + i, value=v), fmt, bold=kind in ("total", "grand"), total=kind in ("total", "grand"))
            if kind in ("total", "grand"):
                ws.cell(row=r, column=1).font = Font(bold=True)
            if kind == "memo":
                ws.cell(row=r, column=1).font = Font(italic=True, color=MUTED)
            r += 1
        return r + 1

    def sheet(title, heading, blocks, ncols):
        ws = wb.create_sheet(title)
        r = brand.xlsx_title_band(ws, heading, _meta_lines(meta) + ["Values as at export - the live-formula model is a separate download"], ncols=max(ncols + 1, 8))
        brand.xlsx_page_setup(ws)
        _setw(ws, {"A": 50})
        for i in range(ncols):
            ws.column_dimensions[CL(2 + i)].width = 13
        r = max(r, 8)
        for name, columns, lines in blocks:
            if name:
                _section(ws, r, name, ncols + 1)
                r += 1
            r = write(ws, r, columns, lines, ncols + 1)
        ws.freeze_panes = "B9"
        ws.sheet_properties.tabColor = GOLD
        return ws

    ann = RP.statement_tables(result, "annual")
    mon = RP.statement_tables(result, "monthly")
    cff = RP.cashflow_forecast(result)
    sheet("Annual statements", "Financial statements - annual", [(t, ann["columns"], ln) for t, ln in ann["statements"]], len(ann["columns"]))
    sheet("Monthly year 1", "Financial statements - year 1 by month", [(t, mon["columns"], ln) for t, ln in mon["statements"]], len(mon["columns"]))
    ws = sheet("Cashflow forecast", "Cash flow forecast", [(None, cff["columns"], cff["lines"])], len(cff["columns"]))
    ws.cell(row=ws.max_row + 1, column=1, value="Reconciles to the cash flow statement" if cff["ok"] else "DOES NOT RECONCILE to the cash flow statement").font = Font(bold=True)
    wb.properties.creator = "Neverlank Chartered Accountants"
    wb.properties.title = meta.get("name") or "Financial statements"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()



def build_project_workbook(meta, p, narrative=None):
    """Capital budgeting workbook with live NPV / IRR / MIRR / payback formulas."""
    p = T.normalise_project(p)
    flows = p["cash_flows"]
    n = len(flows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Project"
    brand.xlsx_title_band(ws, meta.get("name") or "Capital budgeting analysis", _meta_lines(meta), ncols=9)
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 38, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16})
    r = 8
    _section(ws, r, "Inputs", 8)
    r += 1
    refs = {}
    for label, key, val, fmt in (("Discount rate (%)", "rate", p["discount_rate"], PCT_IN), ("Reinvestment rate for MIRR (%)", "reinvest", p["reinvest_rate"], PCT_IN),
                                 ("Initial investment", "initial", p["initial"], NUM), ("Residual / terminal value (received in the final year)", "residual", p["residual"], NUM)):
        ws.cell(row=r, column=1, value=label)
        _inp(ws.cell(row=r, column=2), val, fmt)
        refs[key] = f"$B${r}"
        r += 1
    r += 1
    _section(ws, r, "Cash flows", 8)
    r += 1
    _head(ws, r, ["Year", "Net cash flow", "Cash flow incl. investment & residual", "Cumulative", "Discount factor", "Present value", "Cumulative PV"])
    r += 1
    first = r
    for t in range(n + 1):
        ws.cell(row=r, column=1, value=t)
        if t == 0:
            ws.cell(row=r, column=2, value=0).number_format = NUM
            ws.cell(row=r, column=3, value=f"=-{refs['initial']}")
        else:
            _inp(ws.cell(row=r, column=2), flows[t - 1], NUM)
            ws.cell(row=r, column=3, value=f"=B{r}" + (f"+{refs['residual']}" if t == n else ""))
        ws.cell(row=r, column=4, value=f"=C{r}" if t == 0 else f"=D{r - 1}+C{r}")
        ws.cell(row=r, column=5, value=f"=1/(1+{refs['rate']}/100)^A{r}")
        ws.cell(row=r, column=6, value=f"=C{r}*E{r}")
        ws.cell(row=r, column=7, value=f"=F{r}" if t == 0 else f"=G{r - 1}+F{r}")
        for c, fm in ((3, NUM), (4, NUM), (5, '0.0000'), (6, NUM), (7, NUM)):
            ws.cell(row=r, column=c).number_format = fm
        r += 1
    last = r - 1
    r += 1
    _section(ws, r, "Results", 8)
    r += 1
    res = {}
    for label, key, formula, fmt in (
            ("Net present value (NPV)", "npv", f"=SUM(F{first}:F{last})", NUM),
            ("Internal rate of return (IRR, %)", "irr", f'=IFERROR(IRR(C{first}:C{last})*100,"n/a")', '0.00'),
            ("Modified IRR (MIRR, %)", "mirr", f'=IFERROR(MIRR(C{first}:C{last},{refs["rate"]}/100,{refs["reinvest"]}/100)*100,"n/a")', '0.00'),
            ("Profitability index (PV of inflows / investment)", "pi", f'=IFERROR(SUM(F{first + 1}:F{last})/{refs["initial"]},"n/a")', '0.00'),
            ("Equivalent annual annuity", "eaa", f'=IFERROR(F{last + 2}*({refs["rate"]}/100)/(1-(1+{refs["rate"]}/100)^-{n}),"n/a")', NUM),
            ("Decision", "decision", None, None)):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        res[key] = f"B{r}"
        if key == "decision":
            ws.cell(row=r, column=2, value=f'=IF({res["npv"]}>0,"ACCEPT - positive NPV","REJECT - NPV not positive")').font = Font(bold=True)
        else:
            _fmt(ws.cell(row=r, column=2, value=formula), fmt, bold=True)
        r += 1
    # EAA formula refers to the NPV cell
    eaa_row = int(res["eaa"][1:])
    ws[res["eaa"]].value = f'=IFERROR({res["npv"]}*({refs["rate"]}/100)/(1-(1+{refs["rate"]}/100)^-{n}),"n/a")'
    # payback (simple)
    ws.cell(row=r, column=1, value="Payback period (years)").font = Font(bold=True)
    pb = (f'=IFERROR(INDEX(A{first}:A{last},MATCH(TRUE,INDEX(D{first}:D{last}>=0,0),0))-1'
          f'+(-INDEX(D{first}:D{last},MATCH(TRUE,INDEX(D{first}:D{last}>=0,0),0)-1))/INDEX(C{first}:C{last},MATCH(TRUE,INDEX(D{first}:D{last}>=0,0),0)),"not recovered")')
    _fmt(ws.cell(row=r, column=2, value=pb), '0.00', bold=True)
    r += 1
    ws.cell(row=r, column=1, value="Discounted payback (years)").font = Font(bold=True)
    dpb = (f'=IFERROR(INDEX(A{first}:A{last},MATCH(TRUE,INDEX(G{first}:G{last}>=0,0),0))-1'
           f'+(-INDEX(G{first}:G{last},MATCH(TRUE,INDEX(G{first}:G{last}>=0,0),0)-1))/INDEX(F{first}:F{last},MATCH(TRUE,INDEX(G{first}:G{last}>=0,0),0)),"not recovered")')
    _fmt(ws.cell(row=r, column=2, value=dpb), '0.00', bold=True)
    r += 2
    if narrative:
        _section(ws, r, "Feasibility notes", 8)
        r += 1
        for key, label in T.FEASIBILITY_SECTIONS:
            txt = (narrative or {}).get(key)
            if txt:
                ws.cell(row=r, column=1, value=label).font = Font(bold=True, color=NAVY)
                c = ws.cell(row=r, column=2, value=txt)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
                ws.row_dimensions[r].height = max(30, 15 * (len(txt) // 90 + 1))
                r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {"res": res, "first": first, "last": last}


def build_loan_workbook(meta, principal, rate, months, repay_type="annuity", grace_months=0, balloon_pct=0.0, start=None):
    """Monthly loan schedule with live formulas (change principal / rate / term at the top)."""
    months = max(1, min(int(months), 360))
    wb = Workbook()
    ws = wb.active
    ws.title = "Loan schedule"
    brand.xlsx_title_band(ws, meta.get("name") or "Loan amortisation schedule", _meta_lines(meta), ncols=8)
    brand.xlsx_page_setup(ws)
    _setw(ws, {"A": 34, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16})
    r = 8
    refs = {}
    for label, key, val, fmt in (("Loan amount", "p", principal, NUM), ("Interest rate (% p.a.)", "rate", rate, PCT_IN), ("Term (months)", "n", months, "0"),
                                 ("Grace period (months, interest only)", "grace", grace_months, "0"), ("Balloon at the end (% of loan)", "balloon", balloon_pct, PCT_IN)):
        ws.cell(row=r, column=1, value=label)
        _inp(ws.cell(row=r, column=2), val, fmt)
        refs[key] = f"$B${r}"
        r += 1
    ws.cell(row=r, column=1, value="Repayment type")
    ws.cell(row=r, column=2, value={"annuity": "Equal instalments", "straight": "Equal principal", "bullet": "Interest only, bullet"}[repay_type])
    r += 1
    ws.cell(row=r, column=1, value="Level instalment (equal-instalment loans)")
    amort = f"({refs['n']}-{refs['grace']})"
    mr = f"({refs['rate']}/100/12)"
    bal = f"({refs['p']}*{refs['balloon']}/100)"
    inst = f"=IF({mr}=0,({refs['p']}-{bal})/{amort},({refs['p']}-{bal}/(1+{mr})^{amort})*{mr}/(1-(1+{mr})^-{amort}))"
    _fmt(ws.cell(row=r, column=2, value=inst), NUM2, bold=True)
    inst_ref = f"$B${r}"
    r += 2
    _head(ws, r, ["Month", "Opening balance", "Interest", "Principal", "Payment", "Closing balance"])
    r += 1
    first = r
    for m in range(1, months + 1):
        ws.cell(row=r, column=1, value=m)
        ws.cell(row=r, column=2, value=f"={refs['p']}" if m == 1 else f"=F{r - 1}")
        ws.cell(row=r, column=3, value=f"=B{r}*{mr}")
        if m == months:
            pr = f"=B{r}"
        elif repay_type == "bullet":
            pr = "=0"
        elif repay_type == "straight":
            pr = f"=IF(A{r}<={refs['grace']},0,({refs['p']}-{bal})/{amort})"
        else:
            pr = f"=IF(A{r}<={refs['grace']},0,MIN(B{r},{inst_ref}-C{r}))"
        ws.cell(row=r, column=4, value=pr)
        ws.cell(row=r, column=5, value=f"=C{r}+D{r}")
        ws.cell(row=r, column=6, value=f"=B{r}-D{r}")
        for c in range(2, 7):
            ws.cell(row=r, column=c).number_format = NUM2
        r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(bold=True)
    for c in (3, 4, 5):
        _fmt(ws.cell(row=r, column=c, value=f"=SUM({CL(c)}{first}:{CL(c)}{r - 1})"), NUM2, total=True)
    ws.freeze_panes = ws.cell(row=first, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
