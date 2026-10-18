"""Finance > Firm Finance: the firm's Asset Register and its own Income Statement, Balance Sheet, Cash Flow Statement,
Statement of Changes in Equity and Notes. The maths is in firm_statements.py. Gated by "view_firm_finances"
(Partner/Admin by default) like the rest of the firm's finances."""
import io
from datetime import date, datetime, timedelta

from flask import Blueprint, render_template, redirect, url_for, request, flash, abort, send_file
from flask_login import current_user

from extensions import db
from models import (
    FirmAsset, FirmAssetClass, FirmYearEnd, FirmBalanceLine, FIRM_ASSET_METHODS, FIRM_ENTRY_CATEGORIES,
    FIRM_ENTRY_CATEGORY_CODES, FIRM_FRAMEWORKS, INVOICE_CURRENCIES,
)
from finance import finance_required, _parse_date, _parse_float
import cashbook
import firm_statements as fs
import financials

firmfs_bp = Blueprint("firmfs", __name__, url_prefix="/firm-finance")

# balance-sheet style categories that are normally carried from one year-end to the next
CARRY_CATEGORIES = {
    "intangible_assets", "investment_property", "long_term_investments", "deferred_tax_asset", "other_noncurrent_assets",
    "long_term_borrowings", "deferred_tax_liability", "long_term_provisions", "other_noncurrent_liabilities",
    "share_capital", "share_premium", "other_reserves",
}


def _year_ends():
    return FirmYearEnd.query.order_by(FirmYearEnd.period_end).all()


def _pick_year_end():
    """The year-end named by ?ye=, else the latest one."""
    yes = _year_ends()
    if not yes:
        return None, yes
    want = request.args.get("ye", type=int)
    for y in yes:
        if y.id == want:
            return y, yes
    return yes[-1], yes


def _statements_or_none():
    ye, yes = _pick_year_end()
    if ye is None:
        return None, None, yes
    return fs.build_firm_statements(ye.id, year_ends=yes), ye, yes


def _next_asset_code():
    n = FirmAsset.query.count() + 1
    while FirmAsset.query.filter_by(asset_code=f"FA-{n:04d}").first():
        n += 1
    return f"FA-{n:04d}"


# ------------------------------------------------------------------ overview

@firmfs_bp.route("/")
@finance_required
def index():
    yes = _year_ends()
    latest = fs.build_firm_statements(yes[-1].id, year_ends=yes) if yes else None
    rows = []
    for y in reversed(yes):
        rows.append({"ye": y, "lines": len(y.lines)})
    return render_template("firmfs/index.html", year_ends=rows, latest=latest, active="index",
                           asset_count=FirmAsset.query.count(), today=date.today())


# ------------------------------------------------------------------ year-ends

@firmfs_bp.route("/year-ends/new", methods=["POST"])
@finance_required
def create_year_end():
    pe = _parse_date(request.form.get("period_end"))
    if not pe:
        flash("Enter the year-end date.", "danger")
        return redirect(url_for("firmfs.index"))
    if FirmYearEnd.query.filter_by(period_end=pe).first():
        flash("That year-end already exists.", "danger")
        return redirect(url_for("firmfs.index"))
    typed_rate = _parse_float(request.form.get("zwg_rate"))
    from_register = None
    if not typed_rate:
        import firm_fx
        from_register = firm_fx.register_rate("ZWG", pe)
        typed_rate = from_register.rate if from_register else None
    ye = FirmYearEnd(
        period_end=pe, zwg_rate=typed_rate,
        opening_nil=bool(request.form.get("opening_nil")), created_by_id=current_user.id,
        reporting_framework=request.form.get("reporting_framework") if request.form.get("reporting_framework") in dict(FIRM_FRAMEWORKS) else "ifrs_for_smes",
    )
    db.session.add(ye)
    db.session.commit()
    flash(f"Year-end {pe.strftime('%d %B %Y')} created. Enter the balances the app cannot know."
          + (f" Its ZWG rate ({from_register.rate:g}) was taken from the Exchange rates register (dated {from_register.rate_date.strftime('%d %b %Y')})." if from_register else ""), "success")
    return redirect(url_for("firmfs.year_end", ye_id=ye.id))


@firmfs_bp.route("/year-ends/<int:ye_id>")
@finance_required
def year_end(ye_id):
    ye = FirmYearEnd.query.get_or_404(ye_id)
    stm = fs.build_firm_statements(ye.id)
    cur = stm["cur"]
    d = cur["derived"]
    sched = cur["schedule"]["total"]
    pbt = stm["pl"]["profit_before_tax"]["current"]
    yes = _year_ends()
    earlier = [y for y in yes if y.period_end < ye.period_end]
    grouped = {}
    for ln in ye.lines:
        grouped.setdefault(ln.fs_category, []).append(ln)
    ordered = [(code, label, grouped[code]) for code, label in FIRM_ENTRY_CATEGORIES if code in grouped]
    import firm_fx
    fx = firm_fx.year_end_info(ye)
    cash_info = cashbook.statement_balances(ye.period_end, ye.zwg_rate)
    cash_auto = cashbook.uses_cash_book(ye, ye.period_end)
    manual_cash = cashbook.manual_summary(cur["period_start"], ye.period_end)
    return render_template(
        "firmfs/year_end.html", ye=ye, stm=stm, d=d, sched=sched, pbt=pbt, ordered=ordered, active="index",
        cash_info=cash_info, manual_cash=manual_cash, cash_auto=cash_auto, fx=fx, fx_stale_days=firm_fx.STALE_DAYS,
        categories=FIRM_ENTRY_CATEGORIES, currencies=INVOICE_CURRENCIES, frameworks=FIRM_FRAMEWORKS,
        previous=earlier[-1] if earlier else None,
        indicative_tax=(pbt * (ye.tax_rate_percent or 0) / 100.0) if pbt > 0 else 0.0,
        expense_total=sum(d["expenses"].values()),
    )


@firmfs_bp.route("/year-ends/<int:ye_id>/settings", methods=["POST"])
@finance_required
def save_year_end(ye_id):
    ye = FirmYearEnd.query.get_or_404(ye_id)
    pe = _parse_date(request.form.get("period_end"))
    if pe and pe != ye.period_end:
        if FirmYearEnd.query.filter(FirmYearEnd.period_end == pe, FirmYearEnd.id != ye.id).first():
            flash("Another year-end already has that date.", "danger")
            return redirect(url_for("firmfs.year_end", ye_id=ye.id))
        ye.period_end = pe
    ye.zwg_rate = _parse_float(request.form.get("zwg_rate"))
    if not ye.zwg_rate:
        import firm_fx
        row = firm_fx.register_rate("ZWG", ye.period_end)
        if row:
            ye.zwg_rate = row.rate
            flash(f"No rate was typed, so the register rate {row.rate:g} (dated {row.rate_date.strftime('%d %b %Y')}) is used.", "info")
    rate = _parse_float(request.form.get("tax_rate_percent"))
    if rate is not None:
        ye.tax_rate_percent = rate
    ye.opening_nil = bool(request.form.get("opening_nil"))
    ye.cash_manual = not request.form.get("cash_from_book")
    if request.form.get("reporting_framework") in dict(FIRM_FRAMEWORKS):
        ye.reporting_framework = request.form.get("reporting_framework")
    ye.extra_note = (request.form.get("extra_note") or "").strip() or None
    db.session.commit()
    flash("Year-end settings saved.", "success")
    return redirect(url_for("firmfs.year_end", ye_id=ye.id))


@firmfs_bp.route("/year-ends/<int:ye_id>/delete", methods=["POST"])
@finance_required
def delete_year_end(ye_id):
    ye = FirmYearEnd.query.get_or_404(ye_id)
    label = ye.period_end.strftime("%d %B %Y")
    db.session.delete(ye)
    db.session.commit()
    flash(f"Year-end {label} and its entered balances were deleted. Invoices, expenses and the Asset Register are untouched.", "info")
    return redirect(url_for("firmfs.index"))


def _line_from_form(line):
    code = request.form.get("fs_category")
    desc = (request.form.get("description") or "").strip()
    amount = _parse_float(request.form.get("amount"))
    currency = request.form.get("currency") if request.form.get("currency") in INVOICE_CURRENCIES else "USD"
    if code not in FIRM_ENTRY_CATEGORY_CODES:
        return "Choose what the figure is."
    if not desc:
        return "Describe the figure (for example 'CBZ current account')."
    if amount is None:
        return "Enter the amount as a number."
    line.fs_category, line.description, line.amount, line.currency = code, desc[:200], amount, currency
    return None


@firmfs_bp.route("/year-ends/<int:ye_id>/lines/add", methods=["POST"])
@finance_required
def add_line(ye_id):
    ye = FirmYearEnd.query.get_or_404(ye_id)
    line = FirmBalanceLine(year_end_id=ye.id)
    err = _line_from_form(line)
    if err:
        flash(err, "danger")
    else:
        db.session.add(line)
        db.session.commit()
        flash("Figure added.", "success")
    return redirect(url_for("firmfs.year_end", ye_id=ye.id) + "#lines")


@firmfs_bp.route("/year-ends/<int:ye_id>/lines/<int:line_id>/edit", methods=["POST"])
@finance_required
def edit_line(ye_id, line_id):
    line = FirmBalanceLine.query.filter_by(id=line_id, year_end_id=ye_id).first_or_404()
    err = _line_from_form(line)
    if err:
        flash(err, "danger")
    else:
        db.session.commit()
        flash("Figure updated.", "success")
    return redirect(url_for("firmfs.year_end", ye_id=ye_id) + "#lines")


@firmfs_bp.route("/year-ends/<int:ye_id>/lines/<int:line_id>/delete", methods=["POST"])
@finance_required
def delete_line(ye_id, line_id):
    line = FirmBalanceLine.query.filter_by(id=line_id, year_end_id=ye_id).first_or_404()
    db.session.delete(line)
    db.session.commit()
    flash("Figure removed.", "info")
    return redirect(url_for("firmfs.year_end", ye_id=ye_id) + "#lines")


@firmfs_bp.route("/year-ends/<int:ye_id>/copy-previous", methods=["POST"])
@finance_required
def copy_previous(ye_id):
    """Start a year-end from the previous one: loans, capital, reserves and other long-term balances usually carry over
    (they are copied as a starting point to be corrected, never as the answer)."""
    ye = FirmYearEnd.query.get_or_404(ye_id)
    earlier = FirmYearEnd.query.filter(FirmYearEnd.period_end < ye.period_end).order_by(FirmYearEnd.period_end.desc()).first()
    if not earlier:
        flash("There is no earlier year-end to copy from.", "danger")
        return redirect(url_for("firmfs.year_end", ye_id=ye.id))
    have = {(l.fs_category, l.description) for l in ye.lines}
    n = 0
    for l in earlier.lines:
        if l.fs_category in CARRY_CATEGORIES and (l.fs_category, l.description) not in have:
            db.session.add(FirmBalanceLine(year_end_id=ye.id, fs_category=l.fs_category, description=l.description,
                                           amount=l.amount, currency=l.currency))
            n += 1
    db.session.commit()
    flash(f"Copied {n} long-term balance(s) from {earlier.period_end.strftime('%d %B %Y')}. Check each amount.", "success" if n else "info")
    return redirect(url_for("firmfs.year_end", ye_id=ye.id) + "#lines")


# ------------------------------------------------------------------ asset register

def _period_for(as_at):
    prev = FirmYearEnd.query.filter(FirmYearEnd.period_end < as_at).order_by(FirmYearEnd.period_end.desc()).first()
    return fs.period_start_for(as_at, prev.period_end if prev else None)


@firmfs_bp.route("/assets")
@finance_required
def assets():
    yes = _year_ends()
    default_at = yes[-1].period_end if yes else date.today()
    as_at = _parse_date(request.args.get("as_at")) or default_at
    classes, assets_ = fs.load_register()
    start = _period_for(as_at)
    sched = fs.register_schedule(assets_, classes, start, as_at)
    show_disposed = request.args.get("disposed") == "1"
    rows = [r for r in sched["assets"] if show_disposed or not r["disposed_in_period"]]
    # assets the period does not cover (bought later, or sold before it) are listed under a note
    in_scope_ids = {r["id"] for r in sched["assets"]}
    outside = [a for a in assets_ if a["id"] not in in_scope_ids]
    return render_template(
        "firmfs/assets.html", sched=sched, rows=rows, outside=outside, as_at=as_at, start=start, active="assets",
        show_disposed=show_disposed, class_count=len(classes), year_ends=yes, today=date.today(),
    )


def _classes_for_form(also=None):
    q = FirmAssetClass.query.order_by(FirmAssetClass.order, FirmAssetClass.name).all()
    return [c for c in q if c.is_active or (also is not None and c.id == also)]


def _asset_from_form(asset):
    desc = (request.form.get("description") or "").strip()
    acquired = _parse_date(request.form.get("date_acquired"))
    cost = _parse_float(request.form.get("cost"))
    residual = 0.0 if not (request.form.get("residual_value") or "").strip() else _parse_float(request.form.get("residual_value"))
    cid = request.form.get("asset_class_id", type=int)
    bf = _parse_float(request.form.get("acc_dep_bf"))
    bf_date = _parse_date(request.form.get("acc_dep_bf_date"))
    if not desc:
        return "Describe the asset."
    if not acquired:
        return "Enter the date the asset was acquired."
    if cost is None or cost < 0:
        return "Enter the cost as a number (USD)."
    if residual is None or residual < 0 or residual > cost:
        return "The residual value cannot be negative or more than the cost."
    if cid and not FirmAssetClass.query.get(cid):
        return "Choose an asset class from the list."
    if (bf is None) != (bf_date is None):
        return "Enter both the accumulated depreciation brought forward and the date it is stated at, or neither."
    if bf is not None and (bf < 0 or bf > cost):
        return "Accumulated depreciation brought forward cannot be negative or more than the cost."
    code = (request.form.get("asset_code") or "").strip()
    if code:
        clash = FirmAsset.query.filter(FirmAsset.asset_code == code, FirmAsset.id != (asset.id or 0)).first()
        if clash:
            return f"Asset code {code} is already used."
    asset.asset_code = code or asset.asset_code or _next_asset_code()
    asset.description = desc[:250]
    asset.asset_class_id = cid or None
    asset.date_acquired, asset.cost, asset.residual_value = acquired, cost, residual
    asset.supplier = (request.form.get("supplier") or "").strip()[:150] or None
    asset.location = (request.form.get("location") or "").strip()[:120] or None
    asset.serial_number = (request.form.get("serial_number") or "").strip()[:120] or None
    asset.notes = (request.form.get("notes") or "").strip() or None
    asset.acc_dep_bf, asset.acc_dep_bf_date = bf, bf_date
    return None


@firmfs_bp.route("/assets/new", methods=["GET", "POST"])
@finance_required
def new_asset():
    asset = FirmAsset()
    if request.method == "POST":
        err = _asset_from_form(asset)
        if err:
            flash(err, "danger")
        else:
            asset.created_by_id = current_user.id
            db.session.add(asset)
            db.session.commit()
            flash(f"{asset.asset_code} added to the register.", "success")
            return redirect(url_for("firmfs.assets"))
    return render_template("firmfs/asset_form.html", asset=asset, classes=_classes_for_form(), active="assets", editing=False,
                           form=request.form if request.method == "POST" else None)


@firmfs_bp.route("/assets/<int:asset_id>/edit", methods=["GET", "POST"])
@finance_required
def edit_asset(asset_id):
    asset = FirmAsset.query.get_or_404(asset_id)
    if request.method == "POST":
        err = _asset_from_form(asset)
        if err:
            flash(err, "danger")
            db.session.rollback()
        else:
            db.session.commit()
            flash(f"{asset.asset_code} updated.", "success")
            return redirect(url_for("firmfs.assets"))
        asset = FirmAsset.query.get_or_404(asset_id)
    return render_template("firmfs/asset_form.html", asset=asset, classes=_classes_for_form(asset.asset_class_id), active="assets",
                           editing=True, form=request.form if request.method == "POST" else None)


@firmfs_bp.route("/assets/<int:asset_id>/dispose", methods=["POST"])
@finance_required
def dispose_asset(asset_id):
    asset = FirmAsset.query.get_or_404(asset_id)
    on = _parse_date(request.form.get("disposed_on"))
    proceeds = 0.0 if not (request.form.get("disposal_proceeds") or "").strip() else _parse_float(request.form.get("disposal_proceeds"))
    if not on or on < asset.date_acquired:
        flash("Enter a disposal date that is not before the acquisition date.", "danger")
    elif proceeds is None or proceeds < 0:
        flash("Enter the sale proceeds as a number (0 if scrapped).", "danger")
    else:
        asset.disposed_on, asset.disposal_proceeds = on, proceeds
        asset.disposal_note = (request.form.get("disposal_note") or "").strip()[:250] or None
        db.session.commit()
        flash(f"{asset.asset_code} marked as disposed on {on.strftime('%d %b %Y')}.", "success")
    return redirect(url_for("firmfs.assets"))


@firmfs_bp.route("/assets/<int:asset_id>/undo-disposal", methods=["POST"])
@finance_required
def undo_disposal(asset_id):
    asset = FirmAsset.query.get_or_404(asset_id)
    asset.disposed_on, asset.disposal_proceeds, asset.disposal_note = None, 0.0, None
    db.session.commit()
    flash(f"Disposal of {asset.asset_code} reversed.", "info")
    return redirect(url_for("firmfs.assets"))


@firmfs_bp.route("/assets/<int:asset_id>/delete", methods=["POST"])
@finance_required
def delete_asset(asset_id):
    asset = FirmAsset.query.get_or_404(asset_id)
    code = asset.asset_code
    db.session.delete(asset)
    db.session.commit()
    flash(f"{code} removed from the register. (Use 'Dispose' for an asset that was sold or scrapped.)", "info")
    return redirect(url_for("firmfs.assets"))


# ------------------------------------------------------------------ asset classes

@firmfs_bp.route("/assets/classes")
@finance_required
def asset_classes():
    classes = FirmAssetClass.query.order_by(FirmAssetClass.order, FirmAssetClass.name).all()
    counts = {c.id: len(c.assets) for c in classes}
    return render_template("firmfs/classes.html", classes=classes, counts=counts, methods=FIRM_ASSET_METHODS, active="assets")


def _class_from_form(cls):
    name = " ".join((request.form.get("name") or "").split())
    method = request.form.get("depreciation_method")
    rate = _parse_float(request.form.get("rate_percent"))
    life = _parse_float(request.form.get("useful_life_years"))
    if not name:
        return "Enter a class name."
    clash = FirmAssetClass.query.filter(db.func.lower(FirmAssetClass.name) == name.lower(), FirmAssetClass.id != (cls.id or 0)).first()
    if clash:
        return f"There is already a class called {clash.name}."
    if method not in FIRM_ASSET_METHODS:
        return "Choose a depreciation method."
    if rate is not None and not (0 <= rate <= 100):
        return "The rate must be between 0 and 100 per cent a year."
    if life is not None and life <= 0:
        return "The useful life must be more than zero years."
    if method != "Not depreciated" and rate is None and life is None:
        return "Enter a rate a year or a useful life in years (or choose 'Not depreciated')."
    cls.name, cls.depreciation_method = name[:80], method
    cls.rate_percent = rate
    cls.useful_life_years = life
    return None


@firmfs_bp.route("/assets/classes/add", methods=["POST"])
@finance_required
def add_class():
    cls = FirmAssetClass(order=FirmAssetClass.query.count())
    err = _class_from_form(cls)
    if err:
        flash(err, "danger")
    else:
        db.session.add(cls)
        db.session.commit()
        flash(f"Class '{cls.name}' added.", "success")
    return redirect(url_for("firmfs.asset_classes"))


@firmfs_bp.route("/assets/classes/<int:class_id>/edit", methods=["POST"])
@finance_required
def edit_class(class_id):
    cls = FirmAssetClass.query.get_or_404(class_id)
    err = _class_from_form(cls)
    if err:
        flash(err, "danger")
        db.session.rollback()
    else:
        db.session.commit()
        flash(f"Class '{cls.name}' updated. Depreciation in every statement is recalculated from the new policy.", "success")
    return redirect(url_for("firmfs.asset_classes"))


@firmfs_bp.route("/assets/classes/<int:class_id>/toggle", methods=["POST"])
@finance_required
def toggle_class(class_id):
    cls = FirmAssetClass.query.get_or_404(class_id)
    cls.is_active = not cls.is_active
    db.session.commit()
    flash(f"Class '{cls.name}' {'switched on' if cls.is_active else 'switched off (existing assets keep it)'}.", "info")
    return redirect(url_for("firmfs.asset_classes"))


@firmfs_bp.route("/assets/classes/<int:class_id>/delete", methods=["POST"])
@finance_required
def delete_class(class_id):
    cls = FirmAssetClass.query.get_or_404(class_id)
    if cls.assets:
        flash("A class with assets cannot be deleted; switch it off instead.", "danger")
    else:
        db.session.delete(cls)
        db.session.commit()
        flash("Class deleted.", "info")
    return redirect(url_for("firmfs.asset_classes"))


@firmfs_bp.route("/assets/classes/seed", methods=["POST"])
@finance_required
def seed_classes():
    """Offer a typical starter set once; the firm edits it to its own policy."""
    if FirmAssetClass.query.count():
        flash("Classes already exist.", "info")
    else:
        starter = [
            ("Land and buildings", "Straight-line", 2.5), ("Motor vehicles", "Straight-line", 20.0),
            ("Furniture and fittings", "Straight-line", 10.0), ("Office equipment", "Straight-line", 20.0),
            ("Computer equipment", "Straight-line", 33.33),
        ]
        for i, (n, m, r) in enumerate(starter):
            db.session.add(FirmAssetClass(name=n, depreciation_method=m, rate_percent=r, order=i))
        db.session.commit()
        flash("A starter set of classes was added. Adjust the rates to the firm's own depreciation policy.", "success")
    return redirect(url_for("firmfs.asset_classes"))


# ------------------------------------------------------------------ statements

def _statement_page(template, active):
    stm, ye, yes = _statements_or_none()
    if stm is None:
        flash("Add a year-end first: it tells the statements which date to report at.", "info")
        return redirect(url_for("firmfs.index"))
    return render_template(template, stm=stm, ye=ye, year_ends=yes, active=active, firm_title=fs.FIRM_NAME)


@firmfs_bp.route("/income-statement")
@finance_required
def income_statement():
    return _statement_page("firmfs/income_statement.html", "is")


@firmfs_bp.route("/balance-sheet")
@finance_required
def balance_sheet():
    return _statement_page("firmfs/balance_sheet.html", "bs")


@firmfs_bp.route("/cash-flow")
@finance_required
def cash_flow():
    return _statement_page("firmfs/cash_flow.html", "cf")


@firmfs_bp.route("/changes-in-equity")
@finance_required
def changes_in_equity():
    return _statement_page("firmfs/equity.html", "eq")


@firmfs_bp.route("/notes")
@finance_required
def notes():
    return _statement_page("firmfs/notes.html", "notes")


@firmfs_bp.route("/print")
@finance_required
def print_all():
    stm, ye, yes = _statements_or_none()
    if stm is None:
        return redirect(url_for("firmfs.index"))
    return render_template("firmfs/print.html", stm=stm, ye=ye, firm_title=fs.FIRM_NAME, today=date.today())


@firmfs_bp.route("/export.xlsx")
@finance_required
def export_xlsx():
    stm, ye, yes = _statements_or_none()
    if stm is None:
        return redirect(url_for("firmfs.index"))
    data = build_xlsx(stm)
    name = f"Neverlank_financial_statements_{ye.period_end.isoformat()}.xlsx"
    return send_file(io.BytesIO(data), as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ------------------------------------------------------------------ Excel

def build_xlsx(stm):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E8E2D0")
    num = '#,##0.00;(#,##0.00);"-"'
    wb = Workbook()
    cur_label = stm["labels"]["current"]
    pri_label = stm["labels"]["prior"]
    has_prior = stm["has_prior"]

    def sheet(title, subtitle, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append([fs.FIRM_NAME]); ws["A1"].font = Font(bold=True, size=13)
        ws.append([title]); ws["A2"].font = Font(bold=True, size=12)
        ws.append([subtitle])
        ws.append([])
        ws.column_dimensions["A"].width = 62
        for col in "BCDEFGHIJKL":
            ws.column_dimensions[col].width = 17
        return ws

    def two_col_header(ws, with_note=True):
        row = ["", "Note" if with_note else "", "USD " + cur_label.replace("Year ended ", "")]
        if has_prior:
            row.append("USD " + pri_label.replace("Year ended ", ""))
        ws.append(row)
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = head; c.alignment = Alignment(horizontal="right" if c.column > 1 else "left")

    def put(ws, label, cur, pri, note="", is_bold=False, indent=0):
        row = [("   " * indent) + label, note, cur]
        if has_prior:
            row.append(pri)
        ws.append(row)
        r = ws.max_row
        for c in ws[r][2:]:
            c.number_format = num
        if is_bold:
            for c in ws[r]:
                c.font = bold

    # income statement
    ws = sheet("Income Statement", cur_label, first=True)
    two_col_header(ws)
    for r in stm["pl"]["rows"]:
        put(ws, r["label"], r["current"], r.get("prior"), r.get("note", ""), r.get("bold"))

    # financial position
    ws = sheet("Financial Position", "As at " + stm["cur"]["period_end"].strftime("%d %B %Y"))
    two_col_header(ws)
    s = stm["sfp"]

    def section(title, rows, total, total_label):
        ws.append([title]); ws.cell(ws.max_row, 1).font = bold
        for r in rows:
            put(ws, r["label"], r["current"], r.get("prior"), r.get("note", ""), indent=1)
        put(ws, total_label, total["current"], total["prior"], is_bold=True)

    section("Non-current assets", s["non_current_assets"], s["nca_total"], "Total non-current assets")
    section("Current assets", s["current_assets"], s["ca_total"], "Total current assets")
    put(ws, "Total assets", s["total_assets"]["current"], s["total_assets"]["prior"], is_bold=True)
    section("Equity", s["equity"], s["total_equity"], "Total equity")
    section("Non-current liabilities", s["non_current_liabilities"], s["ncl_total"], "Total non-current liabilities")
    section("Current liabilities", s["current_liabilities"], s["cl_total"], "Total current liabilities")
    put(ws, "Total equity and liabilities", s["total_equity_and_liabilities"]["current"], s["total_equity_and_liabilities"]["prior"], is_bold=True)

    # cash flow
    ws = sheet("Cash Flow", cur_label)
    cf = stm["cf"]
    if cf is None:
        ws.append(["Not available: the previous year-end's balances are needed as the opening position."])
    else:
        ws.append(["", "Note", "USD"])
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = head
        for title, rows, total, tl in (
            ("Cash flows from operating activities", cf["operating_rows"], cf["net_operating"], "Net cash from operating activities"),
            ("Cash flows from investing activities", cf["investing_rows"], cf["net_investing"], "Net cash from investing activities"),
            ("Cash flows from financing activities", cf["financing_rows"], cf["net_financing"], "Net cash from financing activities"),
        ):
            ws.append([title]); ws.cell(ws.max_row, 1).font = bold
            for r in rows:
                ws.append(["   " + r["label"], r.get("note", ""), r["current"]])
                ws.cell(ws.max_row, 3).number_format = num
                if r.get("bold"):
                    for c in ws[ws.max_row]:
                        c.font = bold
            ws.append([tl, "", total]); ws.cell(ws.max_row, 3).number_format = num
            for c in ws[ws.max_row]:
                c.font = bold
        for label, v, b in (("Net increase/(decrease) in cash", cf["net_movement"], True), ("Cash at the start of the period", cf["cash_open"], False),
                            ("Cash at the end of the period", cf["cash_close_computed"], True), ("Cash per the balance sheet", cf["cash_close_actual"], False),
                            ("Unreconciled difference", cf["variance"], False)):
            ws.append([label, "", v]); ws.cell(ws.max_row, 3).number_format = num
            if b:
                for c in ws[ws.max_row]:
                    c.font = bold

    # changes in equity
    ws = sheet("Changes in Equity", cur_label)
    for block in stm["soce"]:
        ws.append([block["heading"]]); ws.cell(ws.max_row, 1).font = bold
        ws.append([""] + [lbl for _c, lbl in block["columns"]] + ["Total"])
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = head
        for r in block["rows"]:
            ws.append([r["label"]] + [r["values"][c] for c, _l in block["columns"]] + [r["total"]])
            for c in ws[ws.max_row][1:]:
                c.number_format = num
            if r["bold"]:
                for c in ws[ws.max_row]:
                    c.font = bold
        ws.append([])

    # notes
    ws = sheet("Notes", cur_label)
    for n in stm["notes"]:
        ws.append([f"Note {n['number']}  {n['title']}"]); ws.cell(ws.max_row, 1).font = bold
        for p in n.get("paragraphs", []):
            ws.append([p]); ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        for title, text in n.get("policies", []):
            ws.append([f"{title}: {text}"]); ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        if n.get("movement"):
            ws.append(["Class", "Opening cost", "Additions", "Disposals", "Closing cost", "Opening depn", "Charge", "Disposals", "Closing depn", "NBV opening", "NBV closing"])
            for c in ws[ws.max_row]:
                c.font = bold; c.fill = head
            for m in n["movement"] + [dict(n["movement_total"], class_name="Total")]:
                ws.append([m["class_name"], m["opening_cost"], m["additions"], -m["disposals_cost"], m["closing_cost"], m["opening_acc_dep"],
                           m["charge"], -m["disposals_acc_dep"], m["closing_acc_dep"], m["opening_nbv"], m["closing_nbv"]])
                for c in ws[ws.max_row][1:]:
                    c.number_format = num
        for a in n.get("accounts", []):
            ws.append(["   " + a["name"], a["current"]] + ([a["prior"]] if has_prior else []))
            for c in ws[ws.max_row][1:]:
                c.number_format = num
        if n.get("extra_table"):
            ws.append([n["extra_table"]["title"]])
            for label, v in n["extra_table"]["rows"]:
                ws.append(["   " + label, v]); ws.cell(ws.max_row, 2).number_format = num
        if n.get("total") and (n.get("accounts") or n.get("movement")):
            ws.append(["Total", n["total"]["current"]] + ([n["total"]["prior"]] if has_prior else []))
            for c in ws[ws.max_row]:
                c.font = bold
            for c in ws[ws.max_row][1:]:
                c.number_format = num
        ws.append([])

    # asset register
    ws = sheet("Asset Register", "Movement for " + cur_label)
    ws.append(["Code", "Description", "Class", "Acquired", "Cost", "Opening depn", "Charge", "Closing depn", "NBV at end", "Disposed on", "Proceeds"])
    for c in ws[ws.max_row]:
        c.font = bold; c.fill = head
    for a in stm["cur"]["schedule"]["assets"]:
        ws.append([a["asset_code"], a["description"], a["class_name"], a["date_acquired"], a["cost"], a["opening_acc_dep"], a["charge"],
                   a["closing_acc_dep"], a["closing_nbv"], a["disposed_on"], a["proceeds"] or None])
        for c in ws[ws.max_row][4:9]:
            c.number_format = num
        ws.cell(ws.max_row, 4).number_format = "dd mmm yyyy"
        if a["disposed_on"]:
            ws.cell(ws.max_row, 10).number_format = "dd mmm yyyy"

    ws = wb.create_sheet("Checks")
    ws.append(["Check", "Result", "Detail"])
    for c in ws[1]:
        c.font = bold; c.fill = head
    for c in stm["checks"]:
        ws.append([c["title"], "OK" if c["ok"] else "CHECK", c["detail"]])
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 10; ws.column_dimensions["C"].width = 110

    for w in wb.worksheets:
        w.oddFooter.left.text = "&8" + fs.FIRM_NAME + " - management accounts, unaudited"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# the Cash Book's routes live in firm_cashbook.py and register on this same blueprint
import firm_cashbook  # noqa: E402,F401
import firm_taxpages  # noqa: E402,F401   (Tax: allowances, computation, deferred tax, QPD)
import firm_ledgerpages  # noqa: E402,F401   (Ledger: journals, chart of accounts, trial balance)
import firm_fx  # noqa: E402,F401   (Exchange rates register)
