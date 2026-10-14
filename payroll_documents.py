"""Generates a Payslip as both a Word (.docx) and a PDF document - the firm
asked for both formats to always be available for the same payslip. Reuses
the firm's letterhead constants/helpers from workpapers.py rather than
duplicating them, and (like workpapers.py) never recalculates a figure -
every number here is the SAME figure already stored on the Payslip record
and shown on screen, so a downloaded document can never disagree with the
app.
"""
import io
import os

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image

from xml.sax.saxutils import escape

import brand_docs as brand
from workpapers import FIRM_NAME, FIRM_ADDRESS, _LOGO_PATH, _GOLD, _GOLD_HEX, _fmt_date, _fmt_num


def _employer_name(period):
    """Neverlank's own name for an internal payslip; the client's name for
    a client-payroll-service payslip (Neverlank is processing payroll ON
    BEHALF OF that client, so the client is the employer of record)."""
    if period.scope == "client" and period.client:
        return period.client.name
    return FIRM_NAME


def _payslip_rows(payslip):
    """The (label, amount) rows shared by both the Word and PDF layout, so
    the two formats can never show different figures for the same
    payslip."""
    earnings = [("Basic Salary", payslip.basic_salary or 0.0)]
    for item in payslip.items:
        if item.category == "Allowance":
            earnings.append((item.label, item.amount or 0.0))
    for item in payslip.items:
        if item.category == "Exempt Income":
            earnings.append((f"{item.label} (Exempt)", item.amount or 0.0))

    deductions = []
    if payslip.tax_credits_total:
        # Show the workings (PAYE before credits, then the credit itself)
        # purely for transparency - these two informational rows are never
        # separately summed into Total Deductions, which uses the stored
        # paye_tax (already net of credits) via Payslip.total_deductions.
        deductions.append(("PAYE before Tax Credits", payslip.paye_before_credits or 0.0))
        for item in payslip.items:
            if item.category == "Tax Credit":
                deductions.append((f"Less: {item.label}", -(item.amount or 0.0)))
    deductions.append(("PAYE (Income Tax)", payslip.paye_tax or 0.0))
    deductions.append(("AIDS Levy", payslip.aids_levy or 0.0))
    if payslip.nssa_employee:
        deductions.append(("NSSA (Employee)", payslip.nssa_employee))
    for item in payslip.items:
        if item.category == "Deduction":
            deductions.append((item.label, item.amount or 0.0))

    return earnings, deductions


def _leave_rows(leave_balances):
    """(leave_type_name, accrued, taken, available) rows for the payslip's
    Leave section - ACCUMULATIVE leave types only (e.g. Annual/Vacation
    Leave), the one(s) the firm tracks a running per-employee balance for.
    Sick/Special/Maternity are statutory provisions, not a balance, so they
    never appear on the payslip - see leave_calc.sync_all_balances_for_employee,
    which is what `leave_balances` already comes from."""
    return [
        (lt.name, bal.accrued_days or 0.0, bal.taken_days or 0.0, bal.current_balance)
        for lt, bal in (leave_balances or [])
    ]


# ---------------------------------------------------------------- shared look

def _letterhead_firm(period):
    """Letterhead lines: the firm's own for an internal payslip; the client's
    name (as employer of record) with 'Payroll processed by <firm>' for a
    client-payroll-service payslip."""
    if period.scope == "client" and period.client:
        return brand.letter_firm(name=period.client.name, strapline=f"Payroll processed by {FIRM_NAME}", address="")
    return brand.letter_firm()


def _meta_items(employee, period):
    return [
        ("Employee no.", employee.employee_number or "—"),
        ("Pay date", _fmt_date(period.pay_date)),
        ("Period", f"{_fmt_date(period.period_start)} - {_fmt_date(period.period_end)}"),
        ("Currency", period.currency or "—"),
    ]


def _employee_lines(employee):
    return [x for x in (
        employee.job_title or None,
        "  ·  ".join(p for p in (f"NSSA No: {employee.nssa_number}" if employee.nssa_number else "",
                                f"National ID: {employee.national_id}" if employee.national_id else "") if p) or None,
    ) if x]


CAVEAT = ("Tax figures calculated per this firm's own editable Payroll Tax Settings. "
          "Verify current PAYE bands, AIDS levy and NSSA rates against ZIMRA/NSSA before relying on this payslip.")


# ---------------------------------------------------------------- Word (.docx)

def generate_payslip_docx(payslip, leave_balances=None):
    employee = payslip.employee
    period = payslip.period

    doc = Document()
    brand.setup_document(doc)
    brand.branded_footer(doc, _letterhead_firm(period))
    brand.letterhead_band(doc, "Payslip", sub=period.name, number=employee.employee_number or None,
                          firm=_letterhead_firm(period))
    brand.meta_strip(doc, _meta_items(employee, period))
    brand.party_block(doc, "Employee", employee.full_name, _employee_lines(employee))

    earnings, deductions = _payslip_rows(payslip)
    max_rows = max(len(earnings), len(deductions))
    body_table = doc.add_table(rows=max_rows + 2, cols=4)
    for idx, text in enumerate(["Earnings", "Amount", "Deductions", "Amount"]):
        body_table.rows[0].cells[idx].paragraphs[0].add_run(text)
    for i in range(max_rows):
        row = body_table.rows[i + 1].cells
        e_label, e_amt = earnings[i] if i < len(earnings) else ("", None)
        d_label, d_amt = deductions[i] if i < len(deductions) else ("", None)
        for idx, text in enumerate((e_label, _fmt_num(e_amt) if e_amt is not None else "",
                                    d_label, _fmt_num(d_amt) if d_amt is not None else "")):
            row[idx].paragraphs[0].add_run(text)
    total_row = body_table.rows[max_rows + 1].cells
    for idx, text in enumerate(["Gross Pay", _fmt_num(payslip.gross_pay), "Total Deductions", _fmt_num(payslip.total_deductions)]):
        total_row[idx].paragraphs[0].add_run(text).bold = True
        brand.shade(total_row[idx], brand.TINT)
    brand.style_table(body_table)
    for row in body_table.rows:
        for idx in (1, 3):
            row.cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    leave_rows = _leave_rows(leave_balances)
    if leave_rows:
        brand.heading(doc, "Leave", level=2)
        leave_table = doc.add_table(rows=len(leave_rows) + 1, cols=4)
        for idx, text in enumerate(["Leave Type", "Accrued (days)", "Taken (days)", "Available (days)"]):
            leave_table.rows[0].cells[idx].paragraphs[0].add_run(text)
        for i, (name, accrued, taken, available) in enumerate(leave_rows):
            row = leave_table.rows[i + 1].cells
            for idx, text in enumerate([name, _fmt_num(accrued), _fmt_num(taken), _fmt_num(available)]):
                row[idx].paragraphs[0].add_run(text)
        brand.style_table(leave_table, numeric_from=1)

    doc.add_paragraph()
    brand.total_panel(doc, "NET PAY", f"{period.currency or ''} {_fmt_num(payslip.net_pay)}", width_in=3.4)
    doc.add_paragraph()
    caveat_p = doc.add_paragraph()
    caveat_run = caveat_p.add_run(CAVEAT)
    caveat_run.font.size = Pt(7.5)
    caveat_run.italic = True
    caveat_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    brand.fix_document(doc)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------- PDF

def generate_payslip_pdf(payslip, leave_balances=None):
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.units import mm
    employee = payslip.employee
    period = payslip.period
    firm = _letterhead_firm(period)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=12 * mm, bottomMargin=16 * mm,
        title=f"Payslip - {employee.full_name} - {period.name}", author=firm.get("name", ""),
    )
    width = A4[0] - 32 * mm
    styles = getSampleStyleSheet()
    base = ParagraphStyle("pb", parent=styles["Normal"], fontSize=9.5, leading=13)
    navy = colors.HexColor(f"#{brand.NAVY}")
    gold = colors.HexColor(f"#{brand.GOLD}")
    tint = colors.HexColor(f"#{brand.TINT}")

    story = [brand.pdf_letterhead(width, "Payslip", number=employee.employee_number or None, sub=period.name, firm=firm),
             Spacer(1, 8 * mm), brand.pdf_meta_strip(width, _meta_items(employee, period)), Spacer(1, 5 * mm)]

    party = [Paragraph("<font size='7' color='#b8934a'><b>EMPLOYEE</b></font>", base),
             Paragraph(f"<font name='Helvetica-Bold' size='11.5' color='#{brand.NAVY}'>{escape(employee.full_name)}</font>", ParagraphStyle("pn", parent=base, leading=15))]
    for line in _employee_lines(employee):
        party.append(Paragraph(f"<font size='8.5' color='#555555'>{escape(line)}</font>", base))
    pt = Table([[party]], colWidths=[width])
    pt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), tint), ("LINEBEFORE", (0, 0), (0, -1), 3.5, gold),
                            ("LEFTPADDING", (0, 0), (-1, -1), 12), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [pt, Spacer(1, 6 * mm)]

    earnings, deductions = _payslip_rows(payslip)
    max_rows = max(len(earnings), len(deductions))
    body = [["EARNINGS", "AMOUNT", "DEDUCTIONS", "AMOUNT"]]
    for i in range(max_rows):
        e_label, e_amt = earnings[i] if i < len(earnings) else ("", None)
        d_label, d_amt = deductions[i] if i < len(deductions) else ("", None)
        body.append([e_label, _fmt_num(e_amt) if e_amt is not None else "", d_label, _fmt_num(d_amt) if d_amt is not None else ""])
    body.append(["Gross Pay", _fmt_num(payslip.gross_pay), "Total Deductions", _fmt_num(payslip.total_deductions)])
    body_table = Table(body, colWidths=[width * 0.32, width * 0.18, width * 0.32, width * 0.18], repeatRows=1)
    style = brand.pdf_table_style()
    body_table.setStyle(style)
    body_table.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("ALIGN", (3, 0), (3, -1), "RIGHT"),
        ("FONTSIZE", (0, 1), (-1, -1), 9.5),
        ("BACKGROUND", (0, -1), (-1, -1), tint), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LINEABOVE", (0, -1), (-1, -1), 0.75, navy),
    ]))
    story += [body_table, Spacer(1, 6 * mm)]

    leave_rows = _leave_rows(leave_balances)
    if leave_rows:
        story += [Paragraph("<font size='7.5' color='#b8934a'><b>LEAVE</b></font>", base), Spacer(1, 2 * mm)]
        leave_body = [["LEAVE TYPE", "ACCRUED (DAYS)", "TAKEN (DAYS)", "AVAILABLE (DAYS)"]]
        for name, accrued, taken, available in leave_rows:
            leave_body.append([name, _fmt_num(accrued), _fmt_num(taken), _fmt_num(available)])
        leave_table = Table(leave_body, colWidths=[width * 0.4, width * 0.2, width * 0.2, width * 0.2], repeatRows=1)
        leave_table.setStyle(brand.pdf_table_style())
        leave_table.setStyle(TableStyle([("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("FONTSIZE", (0, 1), (-1, -1), 9.5)]))
        story += [leave_table, Spacer(1, 6 * mm)]

    right = ParagraphStyle("nr", parent=base, alignment=TA_RIGHT)
    net = Table([[Paragraph("<font color='white'><b>NET PAY</b></font>", base),
                  Paragraph(f"<font color='#{brand.GOLD_LIGHT}' size='12'><b>{escape(period.currency or '')} {escape(_fmt_num(payslip.net_pay))}</b></font>", right)]],
                colWidths=[40 * mm, 52 * mm], hAlign="RIGHT")
    net.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), navy), ("TOPPADDING", (0, 0), (-1, -1), 8),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10)]))
    story += [net, Spacer(1, 8 * mm),
              Paragraph(f"<font size='7.5' color='#6b7684'><i>{escape(CAVEAT)}</i></font>", base)]

    footer = brand.pdf_footer(firm)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf
