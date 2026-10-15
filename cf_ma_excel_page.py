"""One-page Excel exports for the Management Accounts pack: the page you are looking at, as a workbook (values as at export).

Each tab of the pack (Summary, Income statement, Balance sheet, Cash flow, KPIs, Variances, Variance model, Reforecast, Actuals,
Budget & forecast) has a matching builder here. The three statements use the same column layout as the screen and the PDF.
The full pack workbook with live formulas stays in cf_ma_excel.build_pack_workbook.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as CL

import brand_docs as brand
import cf_ma as M
import cf_ma_excel as X
import cf_ma_pdf as P

NUM, PCT = X.NUM, X.PCT
NAVY, GREEN, RED = X.NAVY, X.GREEN, X.RED
PTS = '0.0;(0.0);"-"'

PAGES = {
    "summary": "Summary",
    "income": "Income statement",
    "balance": "Balance sheet",
    "cashflow": "Cash flow",
    "kpis": "KPIs",
    "variances": "Variances",
    "varmodel": "Variance model",
    "reforecast": "Reforecast",
    "actuals": "Actuals",
    "budget": "Budget and forecast",
}


def _labels(d):
    months = M.fy_months(d)
    ridx = M.report_index(d)
    return {"month": M.month_label(months[ridx], False), "fy": M.fy_label(d), "prior": M.prior_fy_label(d),
            "last_month": M.month_label(months[ridx - 1], False) if ridx > 0 else "n/a", "n": ridx + 1}


def _head(ws, title, meta, subtitle, ncols):
    r = brand.xlsx_title_band(ws, title, [meta.get("client", ""), meta.get("name", ""), subtitle], ncols=max(ncols, 8))
    brand.xlsx_page_setup(ws)
    return r


def _hdr(ws, r, heads, left=(1,)):
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font, c.fill = Font(bold=True), X.FILL_SUB
        c.alignment = Alignment(horizontal="left" if i in left else "right", wrap_text=True, vertical="center")


def _val(ws, r, c, v, fmt=NUM, bold=False, color=None):
    cell = ws.cell(row=r, column=c, value=None if v is None else (round(v, 4) if isinstance(v, float) else v))
    cell.number_format = fmt
    if bold or color:
        cell.font = Font(bold=bold, color=color)
    return cell


def _put_note(ws, r, text, ncols=8):
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(italic=True, color=X.MUTED)
    return r + 1


# =========================================================================== the three statements (firm layout)

def _statement(wb, title, meta, d, cols, rows, subtitle, sections=False, note=None):
    ws = wb.active
    ws.title = title[:31]
    ncols = len(cols)
    r = _head(ws, title, meta, subtitle, ncols)
    band, sub = r, r + 1
    li = next(k for k, c in enumerate(cols) if c["kind"] == "label")
    # group band with merged titles
    i = 0
    while i < ncols:
        g = cols[i].get("group", "")
        j = i
        while j + 1 < ncols and cols[j + 1].get("group", "") == g and cols[j + 1]["kind"] == "val" and cols[i]["kind"] == "val":
            j += 1
        ws.cell(row=band, column=i + 1, value=g)
        if j > i:
            ws.merge_cells(start_row=band, start_column=i + 1, end_row=band, end_column=j + 1)
        i = j + 1
    for k, c in enumerate(cols, start=1):
        b = ws.cell(row=band, column=k)
        b.fill = X.FILL_BAND
        b.font = Font(bold=True, color="1D2733")
        b.alignment = Alignment(horizontal="center")
        h = ws.cell(row=sub, column=k, value=c["head"])
        h.font, h.fill = Font(bold=True), X.FILL_SUB
        h.alignment = Alignment(horizontal="left" if c["kind"] == "label" else "right", wrap_text=True)
        ws.column_dimensions[CL(k)].width = 42 if c["kind"] == "label" else 13.5
    r = sub
    group = None
    for row in rows:
        kind = row["kind"]
        g = (("Revenue" if row.get("income") else "Operating expenses") if kind == "indent" else None) if sections else None
        if g and g != group:
            r += 1
            ws.cell(row=r, column=li + 1, value=g).font = Font(bold=True)
            for k in range(1, ncols + 1):
                ws.cell(row=r, column=k).fill = X.FILL_SEC
        group = g
        r += 1
        for k, c in enumerate(cols, start=1):
            if c["kind"] == "label":
                lab = ws.cell(row=r, column=k, value=("    " if kind == "indent" else "") + row["label"])
                lab.font = Font(bold=kind in ("total", "grand"), italic=kind in ("memo", "pct"))
                continue
            v = c["get"](row)
            if v is None:
                continue
            if kind == "pct":
                if c["fmt"] in ("money", "plain"):
                    cell = _val(ws, r, k, v, PCT)
                elif c["fmt"] == "var":
                    cell = _val(ws, r, k, v * 100, PTS)
                else:
                    continue
                cell.font = Font(italic=True, color=X.MUTED)
                continue
            fmt = PCT if c["fmt"] == "varpct" else NUM
            col = None
            if c["fmt"] in ("var", "varpct"):
                fv = (c["fav"](row) if c.get("fav") else v) if c["fmt"] == "varpct" else v
                col = GREEN if (fv or 0) > 0.0005 else (RED if (fv or 0) < -0.0005 else None)
            cell = _val(ws, r, k, v, fmt, bold=kind in ("total", "grand"), color=col)
            if kind in ("total", "grand"):
                cell.border = X.TOP
    ws.freeze_panes = ws.cell(row=sub + 1, column=li + 2 if li == 1 else 3)
    if note:
        _put_note(ws, r + 2, note)
    return ws


def _income(wb, meta, d, labels, args):
    ytd = args.get("view") == "ytd"
    var = args.get("var") != "0"
    inc = M.income_rows(d)
    cols = P._income_cols(labels, ytd, var)
    sub = f"{'Year to date' if ytd else 'Month'}: {labels['month']}. Variances favourable positive, adverse in brackets."
    return _statement(wb, "Income statement", meta, d, cols, inc["rows"], sub, sections=True,
                      note=f"{labels['prior']} is the prior-year accounts spread evenly across the months. Values as at export.")


def _balance(wb, meta, d, labels, args):
    bal = M.balance_rows(d)
    cols = P._balance_cols(labels, args.get("var") != "0")
    return _statement(wb, "Balance sheet", meta, d, cols, bal["rows"], f"As at {labels['month']}. Retained earnings include the current year's profit.",
                      note="Values as at export.")


def _cash(wb, meta, d, labels, args):
    cfl = M.cashflow_rows(d)
    cols = P._cash_cols(labels, args.get("var") != "0")
    return _statement(wb, "Cash flow", meta, d, cols, cfl["rows"], f"{labels['month']}. Indirect method, derived from the movement between balance sheets; cash is net of overdraft.",
                      note="Values as at export.")


# =========================================================================== summary, KPIs, variances

def _kpi_fmt(k, v):
    if v is None:
        return None, NUM
    f = k["fmt"]
    if f == "pct":
        return v, PCT
    if f == "x":
        return (None if v == float("inf") else v), '0.00"x"'
    if f in ("days", "months"):
        return v, '#,##0.0' if f == "months" else '#,##0'
    return v, NUM


def _kpi_rows(ws, r, d, labels):
    _hdr(ws, r, ["Indicator", "Actual", "Budget", labels["prior"], "Status"])
    for k in M.kpi_table(d):
        r += 1
        ws.cell(row=r, column=1, value=k["label"])
        for j, key in enumerate(("actual", "budget", "prior"), start=2):
            v, fmt = _kpi_fmt(k, k[key])
            if v is not None:
                _val(ws, r, j, v, fmt)
        st = {"green": "On track", "amber": "Watch", "red": "Concern"}.get(k["rag"], "")
        c = ws.cell(row=r, column=5, value=st)
        c.font = Font(color={"green": GREEN, "amber": "B26A00", "red": RED}.get(k["rag"]))
        c.alignment = Alignment(horizontal="right")
    return r + 1


def _kpis(wb, meta, d, labels, args):
    ws = wb.active
    ws.title = "KPIs"
    r = _head(ws, "Key performance indicators", meta, f"Year to date, {labels['month']}", 6)
    for col, w in zip("ABCDE", (46, 14, 14, 14, 12)):
        ws.column_dimensions[col].width = w
    if M.vectors(d)["actual_ytd"] is None:
        ws.cell(row=r, column=1, value="No actual figures have been captured for this reporting month.")
        return ws
    _kpi_rows(ws, r, d, labels)
    return ws


def _summary(wb, meta, d, labels, args):
    ws = wb.active
    ws.title = "Summary"
    r = _head(ws, "Summary", meta, f"{labels['month']}  ·  {labels['fy']}", 10)
    ws.column_dimensions["A"].width = 46
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 14
    V = M.vectors(d)
    if V["actual_ytd"] is None:
        ws.cell(row=r, column=1, value="No actual figures have been captured for this reporting month.")
        return ws
    com = d.get("commentary") or {}
    ws.cell(row=r, column=1, value="Key messages").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    msgs = [com["summary"]] if com.get("summary") else M.key_messages(d, "")
    for t in msgs:
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        ws.row_dimensions[r].height = max(15, 15 * (len(t) // 120 + 1))
        r += 1
    r += 1
    inc = {x["key"]: x for x in M.income_rows(d)["rows"]}
    _hdr(ws, r, ["", "Month actual", "Month budget", "Var", "YTD actual", "YTD budget", "Var", "FY forecast", "FY budget", "Var"])
    for key, label in (("revenue", "Revenue"), ("gross_profit", "Gross profit"), ("ebitda", "EBITDA"), ("pat", "Profit after tax")):
        x = inc.get(key)
        if not x:
            continue
        r += 1
        ws.cell(row=r, column=1, value=label).font = Font(bold=key in ("ebitda", "pat"))
        for j, v in enumerate((x["month"]["a"], x["month"]["b"], x["month"]["var"], x["ytd"]["a"], x["ytd"]["b"], x["ytd"]["var"], x["fy"]["f"], x["fy"]["b"], x["fy"]["var"]), start=2):
            _val(ws, r, j, v, NUM, color=(GREEN if v and v > 0.5 else RED if v and v < -0.5 else None) if j in (4, 7, 10) else None)
    r += 2
    ws.cell(row=r, column=1, value="Key performance indicators (year to date)").font = Font(bold=True, size=12, color=NAVY)
    _kpi_rows(ws, r + 1, d, labels)
    return ws


def _variances(wb, meta, d, labels, args):
    ws = X._variance_sheet(wb, meta, d, labels)
    ws.title = "Variances"
    return ws


# =========================================================================== variance model and reforecast

def _varmodel(wb, meta, d, ctx):
    import cf_ma_var as VM
    b = ctx.get("br")
    if not b:
        ws = wb.active
        ws.title = "Variance model"
        r = _head(ws, "Variance model", meta, "", 8)
        ws.cell(row=r, column=1, value="There are no figures to compare for that combination of period and comparison.")
        return ws
    ex = {"bridge": b, "drivers": ctx.get("dr"), "narr": ctx.get("narr") or []}
    ws = X._bridge_sheet(wb, meta, d, ex)
    wb.remove(wb["Sheet"]) if "Sheet" in wb.sheetnames else None
    return ws


def _reforecast(wb, meta, d, ctx, labels):
    res, sc = ctx["res"], ctx["sc"]
    ws = wb.active
    ws.title = "Reforecast"
    r = _head(ws, f"Reforecast - {sc.get('name') or 'this scenario'}", meta,
              f"Actuals to {labels['month']} plus the remaining months rebuilt from {res.get('base_desc', '')}" if not res.get("error") else "", 9)
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 14
    if res.get("error"):
        ws.cell(row=r, column=1, value=res["error"])
        return ws
    for n in res.get("notes") or []:
        r = _put_note(ws, r, n)
    r += 1
    ws.cell(row=r, column=1, value=f"Full year {labels['fy']}").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    _hdr(ws, r, ["", "This scenario", "Pack outlook now", "Var", "Budget", "Var", labels["prior"], "Var"])
    for t in res["table"]:
        r += 1
        bold = t["key"] in ("ebitda", "pat")
        ws.cell(row=r, column=1, value=t["label"]).font = Font(bold=bold)
        for j, key in enumerate(("scenario", "current", "var_current", "budget", "var_budget", "prior", "var_prior"), start=2):
            v = t.get(key)
            _val(ws, r, j, v, NUM, bold=bold and j == 2, color=(GREEN if v and v > 0.5 else RED if v and v < -0.5 else None) if key.startswith("var") else None)
    req = res.get("required")
    if req:
        r += 2
        ws.cell(row=r, column=1, value="What it takes to hit the budget").font = Font(bold=True, size=12, color=NAVY)
        r += 1
        _hdr(ws, r, ["", "Needed in total", "Needed per month", "This scenario per month", "Year to date per month"])
        for lab, a, b_, c_, d_ in (("Revenue", req["revenue_needed"], req["revenue_per_month"], req["scenario_revenue_per_month"], req["ytd_revenue_per_month"]),
                                   ("Profit after tax", req["pat_needed"], req["pat_per_month"], req["scenario_pat_per_month"], req["ytd_pat_per_month"])):
            r += 1
            ws.cell(row=r, column=1, value=lab)
            for j, v in enumerate((a, b_, c_, d_), start=2):
                _val(ws, r, j, v)
    r += 2
    ws.cell(row=r, column=1, value="Month by month").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    _hdr(ws, r, ["Month", "Basis", "Revenue", "Budget", "EBITDA", "Budget", "Profit after tax", "Budget", "Cumulative PAT", "Cumulative budget"])
    for m in res["months"]:
        r += 1
        ws.cell(row=r, column=1, value=m["label"])
        ws.cell(row=r, column=2, value="Actual" if m["kind"] == "actual" else "Reforecast").alignment = Alignment(horizontal="right")
        for j, key in enumerate(("revenue", "b_revenue", "ebitda", "b_ebitda", "pat", "b_pat", "cum_pat", "cum_b_pat"), start=3):
            _val(ws, r, j, m.get(key))
    cash = res.get("cash")
    if cash:
        r += 2
        ws.cell(row=r, column=1, value="Cash and funding (net of overdraft, indicative)").font = Font(bold=True, size=12, color=NAVY)
        r += 1
        _hdr(ws, r, ["Month end", "This scenario", "Pack outlook now", "Budget"])
        for m in cash["months"]:
            r += 1
            ws.cell(row=r, column=1, value=m["label"])
            _val(ws, r, 2, m["net_cash"])
            _val(ws, r, 3, m["current_net_cash"])
            _val(ws, r, 4, m["budget_net_cash"])
        r += 1
        _put_note(ws, r, "Indicative: working capital keeps its relationship to revenue and costs at the reporting month; capital spending, loan repayments and dividends beyond the P&L are not modelled.")
    ws.freeze_panes = "B5"
    return ws


# =========================================================================== actuals and budget grids

def _month_grid(ws, r, d, title, rows_by_month, line_defs, months):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12, color=NAVY)
    r += 1
    _hdr(ws, r, ["Line", *[M.month_label(m) for m in months], "Total"])
    first = r + 1
    for x in line_defs:
        r += 1
        ws.cell(row=r, column=1, value=x["label"])
        for j, m in enumerate(months, start=2):
            v = (rows_by_month.get(m) or {}).get(x["key"])
            if v is not None:
                _val(ws, r, j, M.f(v))
        c = ws.cell(row=r, column=len(months) + 2, value=f"=SUM(B{r}:{CL(len(months) + 1)}{r})")
        c.number_format = NUM
        c.font = Font(bold=True)
    return r + 2


def _actuals(wb, meta, d, labels, args):
    ws = wb.active
    ws.title = "Actuals"
    months = M.fy_months(d)
    r = _head(ws, "Actuals", meta, f"Each month's own income statement, {labels['fy']} (cumulative trial balances are converted to monthly figures)", len(months) + 2)
    ws.column_dimensions["A"].width = 38
    for k in range(2, len(months) + 3):
        ws.column_dimensions[CL(k)].width = 12.5
    defs = [x for x in M.pl_line_defs(d["structure"]) if x["key"] != "div"]
    mon = M.monthly_actuals(d)
    by_month = {m: (mon[i]["values"] or {}) for i, m in enumerate(months)}
    r = _month_grid(ws, r, d, "Income statement by month", by_month, defs, months)
    status = M.data_status(d)
    ws.cell(row=r, column=1, value="Source of each month").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    _hdr(ws, r, ["Month", "Source", "Note"], left=(1, 2, 3))
    for s in status:
        r += 1
        ws.cell(row=r, column=1, value=s["label"])
        src = {"tb": "Trial balance", "manual": "Captured manually", "modelled": "MODELLED from prior year"}.get(s["source"], s["source"] or "")
        ws.cell(row=r, column=2, value=src + (" (estimated)" if s.get("est") else ""))
    ws.freeze_panes = "B5"
    return ws


def _budget(wb, meta, d, labels, args):
    ws = wb.active
    ws.title = "Budget and forecast"
    months = M.fy_months(d)
    r = _head(ws, "Budget and forecast", meta, f"Monthly income statement, {labels['fy']}", len(months) + 2)
    ws.column_dimensions["A"].width = 38
    for k in range(2, len(months) + 3):
        ws.column_dimensions[CL(k)].width = 12.5
    defs = [x for x in M.pl_line_defs(d["structure"]) if x["key"] != "div"]
    any_plan = False
    for name, title in (("budget", "Budget"), ("forecast", "Forecast")):
        plan = M.budget_months(d, name)
        if plan:
            any_plan = True
            r = _month_grid(ws, r, d, title, plan, defs, months)
    if not any_plan:
        ws.cell(row=r, column=1, value="No budget or forecast has been loaded for this pack.")
    ws.freeze_panes = "B5"
    return ws


# =========================================================================== entry point

def build_page_workbook(meta, d, tab, args=None, ctx=None):
    """Workbook for one tab. args: request args (view, var); ctx: the tab's context for varmodel / reforecast."""
    args = args or {}
    labels = _labels(d)
    wb = Workbook()
    if tab == "summary":
        _summary(wb, meta, d, labels, args)
    elif tab == "income":
        _income(wb, meta, d, labels, args)
    elif tab == "balance":
        _balance(wb, meta, d, labels, args)
    elif tab == "cashflow":
        _cash(wb, meta, d, labels, args)
    elif tab == "kpis":
        _kpis(wb, meta, d, labels, args)
    elif tab == "variances":
        wb.remove(wb.active)
        _variances(wb, meta, d, labels, args)
    elif tab == "varmodel":
        _varmodel(wb, meta, d, ctx or {})
    elif tab == "reforecast":
        _reforecast(wb, meta, d, ctx, labels)
    elif tab == "actuals":
        _actuals(wb, meta, d, labels, args)
    elif tab == "budget":
        _budget(wb, meta, d, labels, args)
    else:
        raise KeyError(tab)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = X.GOLD
    wb.properties.creator = "Neverlank Chartered Accountants"
    wb.properties.title = f"{meta.get('name') or 'Management accounts'} - {PAGES.get(tab, tab)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
