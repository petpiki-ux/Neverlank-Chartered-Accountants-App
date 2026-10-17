"""Routes for the Cash Book (Finance > Firm Finance > Cash book). The maths is in cashbook.py; the routes are
registered on the same blueprint as the rest of Firm Finance (imported at the bottom of firm_finance.py)."""
import io
import uuid
from datetime import date

from flask import render_template, redirect, url_for, request, flash, send_file
from flask_login import current_user

from extensions import db
from models import (
    CashAccount, CashEntry, CashAssign, CashClear, CASH_ACCOUNT_KINDS, CASH_ENTRY_CATEGORIES, CASH_TRANSFER_CATEGORY,
    INVOICE_CURRENCIES, PAYMENT_METHODS, FirmYearEnd,
)
from finance import finance_required, _parse_date, _parse_float
from firm_finance import firmfs_bp
import cashbook
import firm_statements as fs

BACK_PREFIX = "/firm-finance/cash-book"


def _back(default):
    b = request.form.get("back") or ""
    return b if b.startswith(BACK_PREFIX) else default


def _periods():
    """Each year-end's reporting period, for quick filters."""
    out, prev = [], None
    for y in FirmYearEnd.query.order_by(FirmYearEnd.period_end).all():
        out.append({"label": y.period_end.strftime("%d %b %Y"), "start": fs.period_start_for(y.period_end, prev), "end": y.period_end})
        prev = y.period_end
    return list(reversed(out))


def _line_label(t):
    return f"{t['ref'] + ' - ' if t['ref'] else ''}{t['description']}"


# ------------------------------------------------------------------ the book

@firmfs_bp.route("/cash-book")
@finance_required
def cash_book():
    accounts, txns, clears = cashbook.load_all()
    today = date.today()
    end = _parse_date(request.args.get("to")) or today
    start = _parse_date(request.args.get("from")) or date(end.year, 1, 1)
    if start > end:
        start, end = end, start
    sel = request.args.get("account", "")
    selected = None
    if sel == "all":
        selected = None
    elif sel.isdigit():
        selected = next((a for a in accounts if a.id == int(sel)), None)
    else:
        selected = next((a for a in accounts if a.is_active), accounts[0] if accounts else None)
    lg = rec = None
    stmt_date = _parse_date(request.args.get("stmt_date"))
    stmt_balance = _parse_float(request.args.get("stmt_balance"))
    if selected is not None:
        lg = cashbook.ledger(selected, txns, clears, start, end)
        if stmt_date and stmt_balance is not None:
            rec = cashbook.reconciliation(selected, txns, clears, stmt_date, stmt_balance)
    rows, totals = cashbook.summary(accounts, txns, start, end)
    un = cashbook.unassigned(txns)
    moves = {a.source_key for a in CashAssign.query.all()}
    return render_template(
        "firmfs/cash_book.html", ye=None, active="cash", accounts=accounts, selected=selected, lg=lg, rec=rec,
        start=start, end=end, rows=rows, totals=totals, unassigned=un, periods=_periods(), moves=moves,
        stmt_date=stmt_date, stmt_balance=stmt_balance, today=today, label=_line_label,
        back=request.full_path.rstrip("?"),
    )


@firmfs_bp.route("/cash-book/assign", methods=["POST"])
@finance_required
def cash_assign():
    key = request.form.get("key", "")
    acc_id = request.form.get("account_id", type=int)
    accounts, txns, _c = cashbook.load_all()
    t = next((x for x in txns if x["key"] == key and x["source"] != "Entry"), None)
    back = _back(url_for("firmfs.cash_book"))
    if not t:
        flash("That transaction was not found.", "danger")
        return redirect(back)
    existing = CashAssign.query.filter_by(source_key=key).first()
    if not acc_id:
        if existing:
            db.session.delete(existing)
            db.session.commit()
        flash("Moved back to the account the rules choose.", "info")
        return redirect(back)
    acct = next((a for a in accounts if a.id == acc_id), None)
    if not acct or acct.currency != t["currency"]:
        flash(f"Choose an account in {t['currency']}: that is the currency of this transaction.", "danger")
        return redirect(back)
    if existing:
        existing.account_id = acct.id
    else:
        db.session.add(CashAssign(source_key=key, account_id=acct.id))
    # a line already marked as on the statement of the old account no longer applies to the new one
    CashClear.query.filter_by(source_key=key).delete()
    db.session.commit()
    flash(f"Moved to {acct.name}.", "success")
    return redirect(back)


@firmfs_bp.route("/cash-book/clear", methods=["POST"])
@finance_required
def cash_clear():
    acc_id = request.form.get("account_id", type=int)
    keys = request.form.getlist("keys")
    action = request.form.get("action")
    back = _back(url_for("firmfs.cash_book", account=acc_id))
    accounts, txns, _c = cashbook.load_all()
    acct = next((a for a in accounts if a.id == acc_id), None)
    if not acct:
        flash("Choose an account.", "danger")
        return redirect(back)
    valid = {t["key"] for t in cashbook.account_rows(acct, txns)}
    keys = [k for k in keys if k in valid]
    if not keys:
        flash("Tick the lines first.", "danger")
        return redirect(back)
    if action == "unmark":
        CashClear.query.filter(CashClear.account_id == acct.id, CashClear.source_key.in_(keys)).delete(synchronize_session=False)
        db.session.commit()
        flash(f"{len(keys)} line(s) marked as not yet on the statement.", "info")
        return redirect(back)
    on = _parse_date(request.form.get("cleared_on"))
    if not on:
        flash("Enter the statement date the lines cleared on.", "danger")
        return redirect(back)
    have = {c.source_key: c for c in CashClear.query.filter(CashClear.account_id == acct.id, CashClear.source_key.in_(keys)).all()}
    for k in keys:
        if k in have:
            have[k].cleared_on = on
        else:
            db.session.add(CashClear(account_id=acct.id, source_key=k, cleared_on=on))
    db.session.commit()
    flash(f"{len(keys)} line(s) marked as on the bank statement of {on.strftime('%d %b %Y')}.", "success")
    return redirect(back)


# ------------------------------------------------------------------ manual entries

def _entry_form_context(entry=None, form=None):
    return dict(ye=None, active="cash", accounts=cashbook.accounts_in_order(active_only=True), entry=entry, form=form,
                categories=CASH_ENTRY_CATEGORIES, today=date.today())


def _read_entry_form(editing=None):
    kind = request.form.get("kind")
    acc_id = request.form.get("account_id", type=int)
    amount = _parse_float(request.form.get("amount"))
    on = _parse_date(request.form.get("entry_date"))
    desc = (request.form.get("description") or "").strip()[:250]
    ref = (request.form.get("reference") or "").strip()[:120]
    cat = request.form.get("category")
    acct = CashAccount.query.get(acc_id) if acc_id else None
    if kind not in ("in", "out", "transfer"):
        return None, "Choose money in, money out or a transfer."
    if not acct:
        return None, "Choose the account."
    if not on:
        return None, "Enter the date."
    if amount is None or amount <= 0:
        return None, "Enter the amount as a number above zero."
    if on < acct.opening_date:
        return None, f"{acct.name} starts on {acct.opening_date.strftime('%d %b %Y')}; an entry cannot be dated before the book starts. Change the account's opening date and balance instead."
    out = {"kind": kind, "acct": acct, "amount": amount, "on": on, "desc": desc, "ref": ref}
    if kind == "transfer":
        to_id = request.form.get("to_account_id", type=int)
        to = CashAccount.query.get(to_id) if to_id else None
        if not to or to.id == acct.id:
            return None, "Choose a different account to transfer to."
        if to.currency != acct.currency:
            return None, "A transfer must be between accounts in the same currency."
        if on < to.opening_date:
            return None, f"{to.name} starts on {to.opening_date.strftime('%d %b %Y')}; the transfer cannot be dated before then."
        out["to"] = to
    else:
        if cat not in CASH_ENTRY_CATEGORIES:
            return None, "Choose what the money was for."
        out["category"] = cat
    return out, None


@firmfs_bp.route("/cash-book/entries/new", methods=["GET", "POST"])
@finance_required
def cash_entry_new():
    if request.method == "POST":
        d, err = _read_entry_form()
        if err:
            flash(err, "danger")
        else:
            if d["kind"] == "transfer":
                tid = str(uuid.uuid4())
                text = d["desc"] or f"Transfer {d['acct'].name} to {d['to'].name}"
                db.session.add(CashEntry(account_id=d["acct"].id, entry_date=d["on"], direction="out", amount=d["amount"], category=CASH_TRANSFER_CATEGORY,
                                         description=text, reference=d["ref"], transfer_id=tid, created_by_id=current_user.id))
                db.session.add(CashEntry(account_id=d["to"].id, entry_date=d["on"], direction="in", amount=d["amount"], category=CASH_TRANSFER_CATEGORY,
                                         description=text, reference=d["ref"], transfer_id=tid, created_by_id=current_user.id))
                flash(f"Transferred {d['amount']:,.2f} from {d['acct'].name} to {d['to'].name}.", "success")
            else:
                db.session.add(CashEntry(account_id=d["acct"].id, entry_date=d["on"], direction=d["kind"], amount=d["amount"], category=d["category"],
                                         description=d["desc"], reference=d["ref"], created_by_id=current_user.id))
                flash("Entry added to the cash book.", "success")
            db.session.commit()
            return redirect(url_for("firmfs.cash_book", account=d["acct"].id))
    return render_template("firmfs/cash_entry.html", **_entry_form_context(None, request.form if request.method == "POST" else None))


@firmfs_bp.route("/cash-book/entries/<int:entry_id>/edit", methods=["GET", "POST"])
@finance_required
def cash_entry_edit(entry_id):
    en = CashEntry.query.get_or_404(entry_id)
    if en.transfer_id:
        flash("A transfer cannot be edited: delete it and enter it again.", "info")
        return redirect(url_for("firmfs.cash_book", account=en.account_id))
    if request.method == "POST":
        d, err = _read_entry_form()
        if err:
            flash(err, "danger")
        elif d["kind"] == "transfer":
            flash("To make this a transfer, delete it and enter the transfer.", "danger")
        else:
            moved = CashClear.query.filter_by(source_key=f"entry:{en.id}").first()
            if moved and moved.account_id != d["acct"].id:
                db.session.delete(moved)
            en.account_id, en.entry_date, en.direction, en.amount = d["acct"].id, d["on"], d["kind"], d["amount"]
            en.category, en.description, en.reference = d["category"], d["desc"], d["ref"]
            db.session.commit()
            flash("Entry updated.", "success")
            return redirect(url_for("firmfs.cash_book", account=en.account_id))
    return render_template("firmfs/cash_entry.html", **_entry_form_context(en, request.form if request.method == "POST" else None))


@firmfs_bp.route("/cash-book/entries/<int:entry_id>/delete", methods=["POST"])
@finance_required
def cash_entry_delete(entry_id):
    en = CashEntry.query.get_or_404(entry_id)
    acc = en.account_id
    group = CashEntry.query.filter_by(transfer_id=en.transfer_id).all() if en.transfer_id else [en]
    for e in group:
        CashClear.query.filter_by(source_key=f"entry:{e.id}").delete()
        db.session.delete(e)
    db.session.commit()
    flash("Transfer deleted from both accounts." if len(group) > 1 else "Entry deleted.", "info")
    return redirect(url_for("firmfs.cash_book", account=acc))


# ------------------------------------------------------------------ accounts

@firmfs_bp.route("/cash-book/accounts")
@finance_required
def cash_accounts():
    accounts, txns, _c = cashbook.load_all()
    today = date.today()
    info = {a.id: {"balance": cashbook.balance_at(a, txns, today),
                   "lines": len(cashbook.account_rows(a, txns))} for a in accounts}
    return render_template("firmfs/cash_accounts.html", ye=None, active="cash", accounts=accounts, info=info, kinds=CASH_ACCOUNT_KINDS,
                           currencies=INVOICE_CURRENCIES, methods=PAYMENT_METHODS, today=today, edit=None, form=None)


def _read_account_form(acct):
    name = " ".join((request.form.get("name") or "").split())
    kind = request.form.get("kind")
    cur = request.form.get("currency")
    raw_opening = (request.form.get("opening_balance") or "").strip()
    opening = 0.0 if raw_opening == "" else _parse_float(raw_opening)
    on = _parse_date(request.form.get("opening_date"))
    methods = [m for m in request.form.getlist("methods") if m in PAYMENT_METHODS]
    if not name:
        return "Enter the account name."
    clash = CashAccount.query.filter(db.func.lower(CashAccount.name) == name.lower(), CashAccount.id != (acct.id or 0)).first()
    if clash:
        return f"There is already an account called {clash.name}."
    if kind not in CASH_ACCOUNT_KINDS:
        return "Choose the kind of account."
    if cur not in INVOICE_CURRENCIES:
        return "Choose the currency."
    if opening is None:
        return "Enter the opening balance as a number (0 if none)."
    if not on:
        return "Enter the date the book starts for this account."
    acct.name, acct.kind, acct.currency = name[:80], kind, cur
    acct.opening_balance, acct.opening_date = opening, on
    acct.methods = ", ".join(methods) or None
    acct.default_in = bool(request.form.get("default_in"))
    acct.default_out = bool(request.form.get("default_out"))
    acct.reference = (request.form.get("reference") or "").strip()[:120] or None
    return None


def _single_defaults(acct):
    """At most one 'default in' and one 'default out' account per currency: the one just saved wins."""
    others = CashAccount.query.filter(CashAccount.currency == acct.currency, CashAccount.id != (acct.id or 0)).all()
    for o in others:
        if acct.default_in:
            o.default_in = False
        if acct.default_out:
            o.default_out = False


@firmfs_bp.route("/cash-book/accounts/add", methods=["POST"])
@finance_required
def cash_account_add():
    acct = CashAccount(order=CashAccount.query.count())
    err = _read_account_form(acct)
    if err:
        flash(err, "danger")
        return redirect(url_for("firmfs.cash_accounts"))
    db.session.add(acct)
    db.session.flush()
    _single_defaults(acct)
    db.session.commit()
    flash(f"Account '{acct.name}' added. Recorded receipts and payments are matched to it from now on.", "success")
    return redirect(url_for("firmfs.cash_accounts"))


@firmfs_bp.route("/cash-book/accounts/<int:account_id>/edit", methods=["GET", "POST"])
@finance_required
def cash_account_edit(account_id):
    acct = CashAccount.query.get_or_404(account_id)
    if request.method == "POST":
        err = _read_account_form(acct)
        if err:
            flash(err, "danger")
            db.session.rollback()
            acct = CashAccount.query.get_or_404(account_id)
        else:
            _single_defaults(acct)
            db.session.commit()
            flash(f"Account '{acct.name}' updated.", "success")
            return redirect(url_for("firmfs.cash_accounts"))
    accounts, txns, _c = cashbook.load_all()
    return render_template("firmfs/cash_accounts.html", ye=None, active="cash", accounts=accounts, info={}, kinds=CASH_ACCOUNT_KINDS,
                           currencies=INVOICE_CURRENCIES, methods=PAYMENT_METHODS, today=date.today(), edit=acct,
                           form=request.form if request.method == "POST" else None)


@firmfs_bp.route("/cash-book/accounts/<int:account_id>/toggle", methods=["POST"])
@finance_required
def cash_account_toggle(account_id):
    acct = CashAccount.query.get_or_404(account_id)
    acct.is_active = not acct.is_active
    if not acct.is_active:
        acct.default_in = acct.default_out = False
    db.session.commit()
    flash(f"'{acct.name}' {'switched on' if acct.is_active else 'switched off. Its history stays, and new transactions will not be matched to it'}.", "info")
    return redirect(url_for("firmfs.cash_accounts"))


@firmfs_bp.route("/cash-book/accounts/<int:account_id>/delete", methods=["POST"])
@finance_required
def cash_account_delete(account_id):
    acct = CashAccount.query.get_or_404(account_id)
    accounts, txns, _c = cashbook.load_all()
    used = bool(cashbook.account_rows(acct, txns)) or CashEntry.query.filter_by(account_id=acct.id).count() \
        or CashAssign.query.filter_by(account_id=acct.id).count() or CashClear.query.filter_by(account_id=acct.id).count()
    if used:
        flash("An account with transactions cannot be deleted; switch it off instead.", "danger")
    else:
        db.session.delete(acct)
        db.session.commit()
        flash("Account deleted.", "info")
    return redirect(url_for("firmfs.cash_accounts"))


@firmfs_bp.route("/cash-book/accounts/seed", methods=["POST"])
@finance_required
def cash_accounts_seed():
    """A typical starter set, each opening at nil on the date of the earliest recorded transaction so that everything
    already recorded is counted; the Partner then sets the true opening balances."""
    if CashAccount.query.count():
        flash("Accounts already exist.", "info")
        return redirect(url_for("firmfs.cash_accounts"))
    txns = cashbook.collect_transactions()
    first = min([t["date"] for t in txns], default=date.today())
    spec = [
        ("Bank (USD)", "Bank", "USD", "Bank transfer, Card, Cheque, Other", True, True),
        ("Cash on hand (USD)", "Cash", "USD", "Cash", False, False),
        ("EcoCash (USD)", "Mobile money", "USD", "EcoCash", False, False),
    ]
    if any(t["currency"] == "ZWG" for t in txns):
        spec.append(("Bank (ZWG)", "Bank", "ZWG", "Bank transfer, Card, Cheque, EcoCash, Cash, Other", True, True))
    for i, (n, k, c, m, di, do) in enumerate(spec):
        db.session.add(CashAccount(name=n, kind=k, currency=c, methods=m, default_in=di, default_out=do, opening_balance=0.0,
                                   opening_date=first, order=i))
    db.session.commit()
    flash(f"{len(spec)} accounts created, each starting at nil on {first.strftime('%d %b %Y')} (the earliest recorded transaction). "
          "Edit each to enter the real opening balance and start date.", "success")
    return redirect(url_for("firmfs.cash_accounts"))


# ------------------------------------------------------------------ Excel

@firmfs_bp.route("/cash-book/export.xlsx")
@finance_required
def cash_book_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    accounts, txns, clears = cashbook.load_all()
    today = date.today()
    end = _parse_date(request.args.get("to")) or today
    start = _parse_date(request.args.get("from")) or date(end.year, 1, 1)
    sel = request.args.get("account", "all")
    chosen = [a for a in accounts if sel == "all" or str(a.id) == sel]
    bold, head = Font(bold=True), PatternFill("solid", fgColor="E8E2D0")
    num = '#,##0.00;(#,##0.00);"-"'
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([fs.FIRM_NAME]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Cash book: {start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')}"])
    ws.append([])
    ws.append(["Account", "Currency", "Opening", "Money in", "Money out", "Closing"])
    for c in ws[ws.max_row]:
        c.font = bold; c.fill = head
    rows, totals = cashbook.summary(chosen, txns, start, end)
    for r in rows:
        ws.append([r["account"].name, r["account"].currency, r["opening"], r["in"], r["out"], r["closing"]])
        for c in ws[ws.max_row][2:]:
            c.number_format = num
    for cur, t in totals.items():
        ws.append([f"Total {cur}", cur, t["opening"], t["in"], t["out"], t["closing"]])
        for c in ws[ws.max_row]:
            c.font = bold
        for c in ws[ws.max_row][2:]:
            c.number_format = num
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    for a in chosen:
        lg = cashbook.ledger(a, txns, clears, start, end)
        w = wb.create_sheet(a.name[:31].replace("/", "-").replace("\\", "-").replace("?", "").replace("*", "").replace("[", "(").replace("]", ")").replace(":", "-"))
        w.append([f"{a.name} ({a.currency})"]); w["A1"].font = Font(bold=True, size=12)
        w.append([f"{start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')}"])
        w.append([])
        w.append(["Date", "Reference", "Details", "Source", "Money in", "Money out", "Balance", "On statement"])
        for c in w[w.max_row]:
            c.font = bold; c.fill = head
        w.append(["", "", "Opening balance", "", "", "", lg["opening"], ""])
        w.cell(w.max_row, 7).number_format = num
        for t in lg["lines"]:
            w.append([t["date"], t["ref"], t["description"], t["source"], t["amount"] if t["amount"] > 0 else None,
                      -t["amount"] if t["amount"] < 0 else None, t["balance"], t["cleared_on"]])
            w.cell(w.max_row, 1).number_format = "dd mmm yyyy"
            for ci in (5, 6, 7):
                w.cell(w.max_row, ci).number_format = num
            if t["cleared_on"]:
                w.cell(w.max_row, 8).number_format = "dd mmm yyyy"
        w.append(["", "", "Totals / closing balance", "", lg["total_in"], lg["total_out"], lg["closing"], ""])
        for c in w[w.max_row]:
            c.font = bold
        for ci in (5, 6, 7):
            w.cell(w.max_row, ci).number_format = num
        w.column_dimensions["A"].width = 13; w.column_dimensions["B"].width = 18; w.column_dimensions["C"].width = 56
        w.column_dimensions["D"].width = 10
        for col in "EFGH":
            w.column_dimensions[col].width = 15
    un = cashbook.unassigned(txns)
    if un:
        w = wb.create_sheet("Unassigned")
        w.append(["Date", "Reference", "Details", "Source", "Currency", "Amount"])
        for c in w[1]:
            c.font = bold; c.fill = head
        for t in un:
            w.append([t["date"], t["ref"], t["description"], t["source"], t["currency"], t["amount"]])
            w.cell(w.max_row, 1).number_format = "dd mmm yyyy"
            w.cell(w.max_row, 6).number_format = num
    for w in wb.worksheets:
        w.oddFooter.left.text = "&8" + fs.FIRM_NAME + " - cash book"
    out = io.BytesIO()
    wb.save(out)
    return send_file(io.BytesIO(out.getvalue()), as_attachment=True, download_name=f"Neverlank_cash_book_{end.isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
