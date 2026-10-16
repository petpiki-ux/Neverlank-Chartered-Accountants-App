"""Reply letters to ZIMRA: the template library, the merge of case facts into a draft, and the Word / PDF output on the
firm letterhead.

A template is a starting point, never a finished letter. Whatever the app does not know (or must not guess - the section
of the Act relied on, the grounds of an objection, the reason for a late filing) is left as a visible
``[CONFIRM: ...]`` placeholder. A letter that still contains a placeholder cannot be approved or marked as sent, so a
half-finished template can never go to ZIMRA by accident.
"""
import io
import re
from datetime import date
from xml.sax.saxutils import escape

PLACEHOLDER_RE = re.compile(r"\[CONFIRM:[^\]]*\]")
DEFAULT_ADDRESSEE = "The Commissioner General\nZimbabwe Revenue Authority\nHarare"
DEFAULT_SALUTATION = "Dear Sir/Madam"
DEFAULT_SIGNATORY_TITLE = "Neverlank Chartered Accountants"

# kind = the ZimraCorrespondence.kind the letter is logged under once sent
TEMPLATES = [
    {
        "key": "ack_extension", "label": "Acknowledgement and request for more time", "kind": "Response / submission",
        "for_kinds": ["Notice of audit"],
        "about": "Acknowledges ZIMRA's letter and asks for an extension of the deadline, giving the reason.",
        "subject": "{client_name}{tin_part}{matter_part} - your letter{their_ref_part}{their_date_part}: request for extension of time",
        "body": (
            "We acknowledge receipt of your letter{their_date_phrase}{their_ref_phrase} concerning {matter}. We act for {client_name} "
            "(TIN {client_tin}) and write on its behalf. [CONFIRM: reference to our authority to act on file with ZIMRA, if any]\n\n"
            "Our client's response is being prepared. We respectfully request an extension of time until [CONFIRM: requested date] "
            "to respond, because [CONFIRM: reason for the extension, for example records for the periods requested are being retrieved].\n\n"
            "We will revert with a full response by that date. Thank you for your co-operation."),
    },
    {
        "key": "info_response", "label": "Response to an information request", "kind": "Response / submission",
        "for_kinds": ["Information request"],
        "about": "Lists the documents and information supplied, in the order ZIMRA asked for them.",
        "subject": "{client_name}{tin_part}{matter_part} - response to your request for information{their_ref_part}",
        "body": (
            "We refer to your letter{their_date_phrase}{their_ref_phrase} requesting information from {client_name} (TIN {client_tin}) "
            "for {tax_head_period}.\n\n"
            "In response, we enclose the following documents and information, numbered to match your request:\n\n"
            "- [CONFIRM: item 1 requested, and the document enclosed]\n"
            "- [CONFIRM: item 2 requested, and the document enclosed]\n\n"
            "[CONFIRM: any item that is not available, with the reason and the date it will be supplied, or delete this line]\n\n"
            "The information is provided from the records held by our client. Please contact us if you need anything further."),
    },
    {
        "key": "response_findings", "label": "Response to audit findings", "kind": "Response / submission",
        "for_kinds": ["Assessment / findings"],
        "about": "Takes each finding in turn: accepted, or disputed with the reason and the law relied on. Reserves the right to object.",
        "subject": "{client_name}{tin_part}{matter_part} - response to audit findings{their_ref_part}",
        "body": (
            "We refer to your letter{their_date_phrase}{their_ref_phrase} setting out the audit findings for {client_name} (TIN {client_tin}) "
            "for {tax_head_period}, which indicate additional tax of {currency} {amount_assessed} and penalties and interest of "
            "{currency} {penalties_interest}.\n\n"
            "We have reviewed the findings with our client and respond as follows:\n\n"
            "1. [CONFIRM: finding 1 - accepted or disputed, the reason, the supporting documents and the law relied on]\n"
            "2. [CONFIRM: finding 2 - as above]\n\n"
            "On this basis our client [CONFIRM: accepts the adjustments to ... / asks that the findings be revised to ...]. "
            "Our client reserves all its rights, including the right to object to any assessment issued, within the time allowed by law.\n\n"
            "We would welcome a meeting to discuss the findings before the assessment is finalised."),
    },
    {
        "key": "objection", "label": "Notice of objection to an assessment", "kind": "Objection",
        "for_kinds": [],
        "about": "Objects to an assessment, states the grounds and the relief sought. Statutory basis and time limit must be confirmed.",
        "subject": "{client_name}{tin_part}{matter_part} - notice of objection to assessment{their_ref_part}",
        "body": (
            "We act for {client_name} (TIN {client_tin}). We refer to the notice of assessment{their_date_phrase}{their_ref_phrase} for "
            "{tax_head_period}, which raised additional tax of {currency} {amount_assessed} and penalties and interest of "
            "{currency} {penalties_interest}.\n\n"
            "On behalf of our client we object to the assessment under [CONFIRM: section of the Act under which the objection is lodged]. "
            "[CONFIRM: date the assessment was served, and that this objection is within the statutory time limit - if it is late, "
            "an application for condonation must go with it]\n\n"
            "The grounds of objection are:\n\n"
            "1. [CONFIRM: ground 1 - the part of the assessment objected to, the amount, and why]\n"
            "2. [CONFIRM: ground 2 - as above]\n\n"
            "We ask that the assessment be [CONFIRM: reduced to ... / set aside]. [CONFIRM: whether to ask that collection of the disputed "
            "amount be suspended while the objection is decided, and whether any payment on account is required by law]\n\n"
            "Documents in support are enclosed: [CONFIRM: list of documents enclosed]."),
    },
    {
        "key": "condonation", "label": "Application for waiver of penalties / condonation", "kind": "Response / submission",
        "for_kinds": [],
        "about": "Asks ZIMRA to waive or reduce penalties and interest, or to condone a late filing or objection, with the reasons.",
        "subject": "{client_name}{tin_part}{matter_part} - application for waiver of penalties / condonation{their_ref_part}",
        "body": (
            "We act for {client_name} (TIN {client_tin}). We refer to {their_doc_phrase} for {tax_head_period}, which raised penalties and "
            "interest of {currency} {penalties_interest}.\n\n"
            "On behalf of our client we apply for [CONFIRM: waiver / reduction of penalties and interest, or condonation of the late "
            "filing or objection] under [CONFIRM: provision of the Act relied on].\n\n"
            "The circumstances are as follows: [CONFIRM: the facts that caused the default, with dates, and why it was outside our "
            "client's reasonable control].\n\n"
            "Our client has since [CONFIRM: filed the returns / paid the tax / put controls in place, with dates]. Its compliance "
            "history is [CONFIRM: compliance record].\n\n"
            "We therefore respectfully ask that the application be granted, and we are available to provide any further information."),
    },
    {
        "key": "payment_plan", "label": "Proposal to pay by instalments", "kind": "Payment plan",
        "for_kinds": [],
        "about": "Proposes an initial payment and monthly instalments for an outstanding balance, and asks for the plan in writing.",
        "subject": "{client_name}{tin_part}{matter_part} - proposal for payment by instalments{their_ref_part}",
        "body": (
            "We act for {client_name} (TIN {client_tin}). We refer to the amount outstanding of {currency} {balance} in respect of "
            "{tax_head_period}.\n\n"
            "Our client wishes to settle the amount by instalments and proposes the following:\n\n"
            "- an initial payment of {currency} [CONFIRM: amount] on [CONFIRM: date]; and\n"
            "- the balance of {currency} [CONFIRM: amount] in [CONFIRM: number] equal monthly instalments of {currency} [CONFIRM: amount], "
            "the first on [CONFIRM: date].\n\n"
            "Full payment now is not possible because [CONFIRM: short, honest explanation of the cash-flow position]. "
            "[CONFIRM: confirm that all current returns are filed and current taxes are being paid, or delete this line]\n\n"
            "We ask that you confirm the agreed plan in writing and, [CONFIRM: what is requested in return - for example that collection "
            "action is suspended while instalments are paid on time]."),
    },
    {
        "key": "tcc", "label": "Request for a Tax Clearance Certificate", "kind": "Tax Clearance Certificate (TCC)",
        "for_kinds": ["Tax Clearance Certificate (TCC)"],
        "about": "Asks for a new or renewed Tax Clearance Certificate and states the compliance position.",
        "subject": "{client_name}{tin_part} - request for Tax Clearance Certificate{their_ref_part}",
        "body": (
            "We act for {client_name} (TIN {client_tin}) and apply for [CONFIRM: the issue / renewal] of a Tax Clearance Certificate "
            "for the purpose of [CONFIRM: purpose, for example a tender or a payment that requires it].{tcc_phrase}\n\n"
            "The compliance position is as follows: [CONFIRM: returns filed and taxes paid up to which date, and the reference of any "
            "payment plan agreement in force].\n\n"
            "We would be grateful if the certificate could be issued, and we will provide anything further that you need."),
    },
    {
        "key": "general", "label": "General letter to ZIMRA", "kind": "Other",
        "for_kinds": [],
        "about": "A blank letter with the client, case and ZIMRA details filled in.",
        "subject": "{client_name}{tin_part}{matter_part}{their_ref_part}",
        "body": (
            "We act for {client_name} (TIN {client_tin}) and write about {matter}.\n\n"
            "[CONFIRM: the content of the letter]\n\n"
            "We look forward to hearing from you."),
    },
]
TEMPLATE_BY_KEY = {t["key"]: t for t in TEMPLATES}


def suggested_template(incoming, case=None):
    """The template that fits the letter being answered (or the case) best."""
    if incoming is not None:
        for t in TEMPLATES:
            if incoming.kind in t["for_kinds"]:
                return t["key"]
    if case is not None and incoming is None:
        if case.case_type in ("Objection", "Appeal"):
            return "objection"
        if case.case_type == "Condonation / waiver of penalties":
            return "condonation"
        if case.case_type == "Payment plan / TCC":
            return "payment_plan"
    return "general"


def _date_text(d):
    return d.strftime("%d %B %Y").lstrip("0") if d else ""


def _money(v):
    return f"{v:,.2f}" if v is not None else ""


class _Fill(dict):
    """format_map dictionary: a field with no value becomes a visible [CONFIRM: ...] placeholder."""
    LABELS = {
        "client_tin": "client's TIN / BP number", "amount_assessed": "amount of tax assessed",
        "penalties_interest": "penalties and interest raised", "balance": "balance outstanding",
        "matter": "the matter this letter is about", "tax_head_period": "tax head and period",
        "currency": "currency",
    }

    def __missing__(self, key):
        return f"[CONFIRM: {self.LABELS.get(key, key.replace('_', ' '))}]"


def build_fields(case, incoming=None):
    """The merge fields for a case (and the letter being answered). Returns (body_fields, subject_fields): the body
    fields show a [CONFIRM] placeholder for anything unknown; the subject fields simply leave it out."""
    client = case.client
    tin = (client.tin_number or "").strip()
    their_date = _date_text(incoming.corr_date) if incoming is not None else ""
    their_ref = (incoming.reference or "").strip() if incoming is not None else ""
    if not their_ref:
        their_ref = (case.zimra_reference or "").strip()
    tax_head = (case.tax_head or "").strip()
    period = case.period_label or ""
    tax_head_period = " ".join(x for x in (tax_head, f"({period})" if period else "") if x)
    matter = f"{case.title}" + (f" ({tax_head_period})" if tax_head_period else "")
    cur = case.currency or "USD"
    tcc_bits = []
    if case.tcc_number:
        tcc_bits.append(f" Our client's previous certificate is No. {case.tcc_number}")
        if case.tcc_expiry:
            tcc_bits.append(f", which {'expired' if case.tcc_expiry < date.today() else 'expires'} on {_date_text(case.tcc_expiry)}")
        tcc_bits.append(".")
    their_doc = "your letter" + (f" dated {their_date}" if their_date else "") + (f" (ref {their_ref})" if their_ref else "")
    body = {
        "client_name": client.name, "client_tin": tin,
        "matter": f"{case.title} ({tax_head_period})" if tax_head_period else case.title,
        "tax_head_period": tax_head_period, "currency": cur,
        "amount_assessed": _money(case.amount_assessed), "penalties_interest": _money(case.penalties_interest),
        "balance": _money(case.balance),
        "their_date_phrase": f" dated {their_date}" if their_date else "",
        "their_ref_phrase": f" (your reference {their_ref})" if their_ref else "",
        "their_doc_phrase": their_doc, "tcc_phrase": "".join(tcc_bits),
    }
    subject = dict(body)
    subject.update({
        "tin_part": f" (TIN {tin})" if tin else "", "matter_part": f" - {case.title}" if case.title else "",
        "their_ref_part": f" (your ref {their_ref})" if their_ref else "",
        "their_date_part": f" dated {their_date}" if their_date else "",
    })
    required = {"client_tin", "amount_assessed", "penalties_interest", "balance", "tax_head_period"}
    return _Fill({k: v for k, v in body.items() if k not in required or v not in ("", None)}), _SubjectFill(subject)


class _SubjectFill(dict):
    def __missing__(self, key):
        return ""


def clean_subject(text):
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+([,.:;])", r"\1", text)
    return text[:300]


def new_draft(template_key, case, incoming, user, seq=1):
    """The field values for a new draft reply."""
    t = TEMPLATE_BY_KEY.get(template_key) or TEMPLATE_BY_KEY["general"]
    body_f, subj_f = build_fields(case, incoming)
    # body templates only use body fields; give the subject-only keys an empty value so they never show as placeholders
    for k in ("tin_part", "matter_part", "their_ref_part", "their_date_part"):
        body_f[k] = ""
    attention = f"\nAttention: {case.zimra_officer}" if case.zimra_officer else ""
    partner = case.engagement.partner.name if (case.engagement is not None and getattr(case.engagement, "partner", None)) else None
    return {
        "template_key": t["key"],
        "subject": clean_subject(t["subject"].format_map(subj_f)),
        "body": t["body"].format_map(body_f),
        "addressee": DEFAULT_ADDRESSEE + attention,
        "salutation": DEFAULT_SALUTATION,
        "our_reference": f"{case.case_number}/L{seq}",
        "signatory": partner or getattr(user, "name", "") or "",
        "signatory_title": DEFAULT_SIGNATORY_TITLE,
        "letter_date": date.today(),
    }


def unresolved(*texts):
    """Every [CONFIRM: ...] placeholder still in the letter, in order, without repeats."""
    seen, out = set(), []
    for text in texts:
        for m in PLACEHOLDER_RE.findall(text or ""):
            if m not in seen:
                seen.add(m)
                out.append(m)
    return out


def paragraphs(body):
    """Split letter text into blocks: ('p', text) paragraphs, ('b', text) bullets and ('n', text) numbered points."""
    blocks = []
    for chunk in re.split(r"\n\s*\n", (body or "").strip()):
        lines = [ln.rstrip() for ln in chunk.split("\n") if ln.strip()]
        if not lines:
            continue
        if all(ln.lstrip().startswith("- ") for ln in lines):
            blocks += [("b", ln.lstrip()[2:].strip()) for ln in lines]
        elif all(re.match(r"^\d+[.)]\s", ln.lstrip()) for ln in lines):
            blocks += [("n", ln.lstrip()) for ln in lines]
        else:
            blocks.append(("p", " ".join(ln.strip() for ln in lines)))
    return blocks


def _banner(reply):
    if reply.status in ("Approved", "Sent"):
        return None
    return "DRAFT - NOT APPROVED FOR SENDING"


def letter_filename(reply, ext):
    from re import sub
    ref = sub(r"[^A-Za-z0-9_-]+", "_", reply.our_reference or f"reply{reply.id}").strip("_")
    return f"{ref}.{ext}"


# --------------------------------------------------------------------------------------------- Word

def build_docx(reply, firm=None):
    """The letter as a Word document on the firm letterhead. Unresolved [CONFIRM] placeholders are highlighted."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Inches, Pt
    import brand_docs as brand

    firm = firm or brand.firm_info()
    doc = Document()
    brand.setup_document(doc, margins_in=0.95, base_size=11)
    brand.branded_footer(doc, firm, center=reply.our_reference or "")
    banner = _banner(reply)
    brand.letterhead_band(doc, "Letter", sub=banner or None, number=None, firm=firm)

    def add_runs(par, text, bold=False, size=11):
        pos = 0
        for m in PLACEHOLDER_RE.finditer(text):
            if m.start() > pos:
                brand._run(par, text[pos:m.start()], size=size, bold=bold)
            r = brand._run(par, m.group(0), size=size, bold=True)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            pos = m.end()
        if pos < len(text):
            brand._run(par, text[pos:], size=size, bold=bold)

    def para(text="", bold=False, after=8, align=None, indent=None, hanging=None):
        p = doc.add_paragraph()
        brand._tight(p, before=0, after=after)
        if align:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.left_indent = Inches(indent)
        if hanging is not None:
            p.paragraph_format.first_line_indent = Inches(-hanging)
        if text:
            add_runs(p, text, bold=bold)
        return p

    refs = []
    if reply.our_reference:
        refs.append(f"Our ref: {reply.our_reference}")
    if reply.in_reply_to is not None and reply.in_reply_to.reference:
        refs.append(f"Your ref: {reply.in_reply_to.reference}")
    elif reply.case.zimra_reference:
        refs.append(f"Your ref: {reply.case.zimra_reference}")
    para("  |  ".join(refs), after=2)
    para(_date_text(reply.letter_date), after=10)
    lines = [ln for ln in (reply.addressee or "").split("\n") if ln.strip()]
    for i, ln in enumerate(lines):
        para(ln, after=0 if i < len(lines) - 1 else 10)
    para(reply.salutation or DEFAULT_SALUTATION, after=8)
    para((reply.subject or "").upper(), bold=True, after=8)
    blocks = paragraphs(reply.body)
    for i, (kind, text) in enumerate(blocks):
        last_of_run = i + 1 >= len(blocks) or blocks[i + 1][0] != kind
        if kind == "b":
            p = doc.add_paragraph(style="List Bullet")
            brand._tight(p, before=0, after=8 if last_of_run else 3)
            add_runs(p, text)
        elif kind == "n":
            para(text, after=8 if last_of_run else 3, indent=0.3, hanging=0.3)
        else:
            para(text, after=8)
    para("Yours faithfully", after=30)
    if reply.signatory:
        para(reply.signatory, bold=True, after=0)
    if reply.signatory_title:
        para(reply.signatory_title, after=0)
    brand.fix_document(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------------------------- PDF

def build_pdf(reply, firm=None):
    """The letter as a PDF on the firm letterhead (unresolved placeholders in red)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    import brand_docs as brand

    firm = firm or brand.firm_info()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=12 * mm, bottomMargin=20 * mm,
                            title=reply.subject or "Letter", author=firm.get("name", ""))
    width = A4[0] - 40 * mm
    base = ParagraphStyle("lt", fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=7)
    tight = ParagraphStyle("ltt", parent=base, spaceAfter=0)

    def fmt(text):
        out, pos = "", 0
        for m in PLACEHOLDER_RE.finditer(text):
            out += escape(text[pos:m.start()]) + f"<font color='#c0392b'><b>{escape(m.group(0))}</b></font>"
            pos = m.end()
        return out + escape(text[pos:])

    story = [brand.pdf_letterhead(width, "Letter", sub=_banner(reply), firm=firm), Spacer(1, 8 * mm)]
    refs = []
    if reply.our_reference:
        refs.append(f"Our ref: {reply.our_reference}")
    if reply.in_reply_to is not None and reply.in_reply_to.reference:
        refs.append(f"Your ref: {reply.in_reply_to.reference}")
    elif reply.case.zimra_reference:
        refs.append(f"Your ref: {reply.case.zimra_reference}")
    if refs:
        story.append(Paragraph(escape("  |  ".join(refs)), tight))
    story += [Spacer(1, 2 * mm), Paragraph(escape(_date_text(reply.letter_date)), base), Spacer(1, 2 * mm)]
    for ln in [x for x in (reply.addressee or "").split("\n") if x.strip()]:
        story.append(Paragraph(escape(ln), tight))
    story += [Spacer(1, 5 * mm), Paragraph(escape(reply.salutation or DEFAULT_SALUTATION), base),
              Paragraph(f"<b>{fmt((reply.subject or '').upper())}</b>", base)]
    bullets = []

    def flush():
        nonlocal bullets
        if bullets:
            story.append(ListFlowable([ListItem(Paragraph(fmt(t), base), leftIndent=12) for t in bullets], bulletType="bullet",
                                      start="•", leftIndent=12))
            bullets = []

    blocks = paragraphs(reply.body)
    for i, (kind, text) in enumerate(blocks):
        if kind == "b":
            bullets.append(text)
            continue
        flush()
        if kind == "n":
            last = i + 1 >= len(blocks) or blocks[i + 1][0] != "n"
            story.append(Paragraph(fmt(text), ParagraphStyle("ln", parent=base, leftIndent=14, firstLineIndent=-14, spaceAfter=7 if last else 3)))
        else:
            story.append(Paragraph(fmt(text), base))
    flush()
    story += [Spacer(1, 4 * mm), Paragraph("Yours faithfully", base), Spacer(1, 14 * mm)]
    if reply.signatory:
        story.append(Paragraph(f"<b>{escape(reply.signatory)}</b>", tight))
    if reply.signatory_title:
        story.append(Paragraph(escape(reply.signatory_title), tight))
    footer = brand.pdf_footer(firm)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
