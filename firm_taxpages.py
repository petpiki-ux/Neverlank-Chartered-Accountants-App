"""Routes for Firm Finance > Tax: tax computation, capital allowances, deferred tax, QPD and the tax rules. The maths is in
firm_tax.py; the figures reach the financial statements through firm_statements.build_column."""
import io
from datetime import date

from flask import render_template, redirect, url_for, request, flash, send_file

from extensions import db
from models import (
    FirmAssetClass, FirmTaxAssetRule, FirmTaxCategoryRule, FirmTaxLine, FirmDeferredItem, FirmYearEnd, FIRM_TAX_METHODS,
    FIRM_TAX_LINE_TYPES,
)
from finance import finance_required, _parse_date, _parse_float, category_names
from firm_finance import firmfs_bp, _pick_year_end
import firm_statements as fs
import firm_tax


def _tax_context():
    """(stm, ye, year_ends, tax) for the year-end named by ?ye=, or None with a redirect target."""
    ye, yes = _pick_year_end()
    if ye is None:
        flash("Add a year-end first: the tax pages work on one year-end at a time.", "info")
        return None
    stm = fs.build_firm_statements(ye.id, year_ends=yes)
    return stm, ye, yes, stm["cur"].get("tax")


def _tax_page(template, active, **extra):
    got = _tax_context()
    if got is None:
        return redirect(url_for("firmfs.index"))
    stm, ye, yes, tax = got
    return render_template(template, stm=stm, ye=ye, year_ends=yes, tax=tax, active=active, **extra)


def _back(ye_id, endpoint="firmfs.tax_computation"):
    return redirect(url_for(endpoint, ye=ye_id))


def _num(raw):
    """(value, ok): blank is 0, anything that is not a number is refused."""
    raw = (raw or "").strip()
    if raw == "":
        return 0.0, True
    v = _parse_float(raw)
    return (v, True) if v is not None else (0.0, False)


# ------------------------------------------------------------------ overview / computation

@firmfs_bp.route("/tax")
@finance_required
def tax_index():
    return redirect(url_for("firmfs.tax_computation", ye=request.args.get("ye", type=int)))


@firmfs_bp.route("/tax/computation")
@finance_required
def tax_computation():
    return _tax_page("firmfs/tax_computation.html", "tax_comp", line_types=FIRM_TAX_LINE_TYPES)


@firmfs_bp.route("/tax/lines/add", methods=["POST"])
@finance_required
def tax_line_add():
    ye = FirmYearEnd.query.get_or_404(request.form.get("ye_id", type=int) or 0)
    kind = request.form.get("item_type")
    desc = (request.form.get("description") or "").strip()[:200]
    amount, ok = _num(request.form.get("amount"))
    if kind not in dict(FIRM_TAX_LINE_TYPES) or not desc or not ok or amount <= 0:
        flash("Choose the type, describe the item and enter an amount above zero.", "danger")
    else:
        db.session.add(FirmTaxLine(year_end_id=ye.id, item_type=kind, description=desc, amount=amount))
        db.session.commit()
        flash("Line added to the tax computation.", "success")
    return _back(ye.id)


@firmfs_bp.route("/tax/lines/<int:line_id>/delete", methods=["POST"])
@finance_required
def tax_line_delete(line_id):
    ln = FirmTaxLine.query.get_or_404(line_id)
    ye_id = ln.year_end_id
    db.session.delete(ln)
    db.session.commit()
    flash("Line removed.", "info")
    return _back(ye_id)


@firmfs_bp.route("/tax/loss", methods=["POST"])
@finance_required
def tax_loss_save():
    ye = FirmYearEnd.query.get_or_404(request.form.get("ye_id", type=int) or 0)
    raw = (request.form.get("tax_loss_bf") or "").strip()
    if raw == "":
        ye.tax_loss_bf = None
        flash("The assessed loss brought forward now follows the previous year-end.", "success")
    else:
        v = _parse_float(raw)
        if v is None or v < 0:
            flash("Enter the assessed loss brought forward as a number (zero or more), or leave it blank.", "danger")
            return _back(ye.id)
        ye.tax_loss_bf = v
        flash("Assessed loss brought forward saved.", "success")
    db.session.commit()
    return _back(ye.id)


@firmfs_bp.route("/tax/mode", methods=["POST"])
@finance_required
def tax_mode_save():
    ye = FirmYearEnd.query.get_or_404(request.form.get("ye_id", type=int) or 0)
    ye.tax_manual = not request.form.get("from_tax_pages")
    db.session.commit()
    flash("Tax in the statements now comes from the Tax pages." if not ye.tax_manual else
          "Tax in the statements is now the figures you enter on the year-end page.", "success")
    return _back(ye.id)


# ------------------------------------------------------------------ capital allowances

@firmfs_bp.route("/tax/capital-allowances")
@finance_required
def tax_capital_allowances():
    return _tax_page("firmfs/tax_allowances.html", "tax_cap", methods=dict(FIRM_TAX_METHODS))


# ------------------------------------------------------------------ deferred tax

@firmfs_bp.route("/tax/deferred")
@finance_required
def tax_deferred():
    return _tax_page("firmfs/tax_deferred.html", "tax_def")


@firmfs_bp.route("/tax/deferred/add", methods=["POST"])
@finance_required
def tax_deferred_add():
    ye = FirmYearEnd.query.get_or_404(request.form.get("ye_id", type=int) or 0)
    desc = (request.form.get("description") or "").strip()[:200]
    acc, ok1 = _num(request.form.get("accounting_amount"))
    base, ok2 = _num(request.form.get("tax_base_amount"))
    if not desc or not ok1 or not ok2:
        flash("Describe the item and enter the carrying amount and the tax base as numbers (blank = 0).", "danger")
    else:
        db.session.add(FirmDeferredItem(year_end_id=ye.id, description=desc, accounting_amount=acc, tax_base_amount=base))
        db.session.commit()
        flash("Temporary difference added.", "success")
    return _back(ye.id, "firmfs.tax_deferred")


@firmfs_bp.route("/tax/deferred/<int:item_id>/delete", methods=["POST"])
@finance_required
def tax_deferred_delete(item_id):
    it = FirmDeferredItem.query.get_or_404(item_id)
    ye_id = it.year_end_id
    db.session.delete(it)
    db.session.commit()
    flash("Temporary difference removed.", "info")
    return _back(ye_id, "firmfs.tax_deferred")


# ------------------------------------------------------------------ QPD

@firmfs_bp.route("/tax/qpd")
@finance_required
def tax_qpd():
    got = _tax_context()
    if got is None:
        return redirect(url_for("firmfs.index"))
    stm, ye, yes, tax = got
    as_at = _parse_date(request.args.get("as_at")) or min(date.today(), ye.period_end)
    classes, assets = fs.load_register()
    est = firm_tax.qpd_estimate(ye, stm["cur"]["period_start"], as_at, classes, assets, (stm["pri"] or {}).get("tax"))
    return render_template("firmfs/tax_qpd.html", stm=stm, ye=ye, year_ends=yes, tax=tax, est=est, active="tax_qpd")


# ------------------------------------------------------------------ rules

@firmfs_bp.route("/tax/rules")
@finance_required
def tax_rules():
    ye, yes = _pick_year_end()
    rules = firm_tax.ensure_rules()
    classes = FirmAssetClass.query.order_by(FirmAssetClass.order, FirmAssetClass.name).all()
    cats = category_names(active_only=False)
    crules = firm_tax.category_rules()
    return render_template("firmfs/tax_rules.html", ye=ye, year_ends=yes, active="tax_rules", classes=classes, rules=rules,
                           methods=FIRM_TAX_METHODS, cats=cats, crules=crules,
                           cnotes={r.category: r.note for r in FirmTaxCategoryRule.query.all()})


@firmfs_bp.route("/tax/rules/assets", methods=["POST"])
@finance_required
def tax_rules_assets_save():
    rules = firm_tax.ensure_rules()
    errors = []
    for cid, rule in rules.items():
        f = request.form
        method = f.get(f"method_{cid}")
        if method not in dict(FIRM_TAX_METHODS):
            continue
        wt, ok1 = _num(f.get(f"wt_{cid}"))
        sia, ok2 = _num(f.get(f"sia_{cid}"))
        yrs, ok3 = _num(f.get(f"years_{cid}"))
        cap, ok4 = _num(f.get(f"cap_{cid}"))
        first_raw = (f.get(f"first_{cid}") or "").strip()
        first, ok5 = _num(first_raw) if first_raw else (None, True)
        if not (ok1 and ok2 and ok3 and ok4 and ok5) or min(wt, sia, yrs, cap) < 0 or wt > 100 or sia > 100 or (first is not None and not 0 <= first <= 100):
            errors.append(rule.asset_class.name if rule.asset_class else str(cid))
            continue
        changed = (rule.method != method or (rule.wt_rate_percent or 0) != wt or (rule.sia_percent or 0) != sia
                   or (rule.sia_years or 0) != int(yrs) or (rule.cost_cap or 0) != cap or rule.sia_first_percent != first
                   or bool(rule.prorata_first_year) != bool(f.get(f"prorata_{cid}")))
        rule.method, rule.wt_rate_percent, rule.sia_percent, rule.sia_years = method, wt, sia, int(yrs)
        rule.sia_first_percent = first
        rule.cost_cap = cap if cap > 0 else None
        rule.prorata_first_year = bool(f.get(f"prorata_{cid}"))
        rule.note = (f.get(f"note_{cid}") or "").strip()[:250] or None
        # a changed rule is no longer the one the Partner checked
        rule.verified = bool(f.get(f"verified_{cid}")) and not changed
    db.session.commit()
    if errors:
        flash("Not saved (rates must be numbers between 0 and 100, the rest zero or more): " + ", ".join(errors), "danger")
    else:
        flash("Capital allowance rules saved. A rule you changed shows as not verified until you tick it again after saving.", "success")
    return redirect(url_for("firmfs.tax_rules", ye=request.form.get("ye", type=int)))


@firmfs_bp.route("/tax/rules/categories", methods=["POST"])
@finance_required
def tax_rules_categories_save():
    errors = []
    for i, cat in enumerate(category_names(active_only=False)):
        pct, ok = _num(request.form.get(f"pct_{i}"))
        if request.form.get(f"name_{i}") != cat:
            continue
        if not ok or pct < 0 or pct > 100:
            errors.append(cat)
            continue
        row = FirmTaxCategoryRule.query.filter_by(category=cat).first()
        note = (request.form.get(f"note_{i}") or "").strip()[:250] or None
        if pct == 0 and not note:
            if row:
                db.session.delete(row)
            continue
        if row is None:
            row = FirmTaxCategoryRule(category=cat)
            db.session.add(row)
        row.disallowed_percent, row.note = pct, note
    db.session.commit()
    if errors:
        flash("Not saved (a percentage must be a number from 0 to 100): " + ", ".join(errors), "danger")
    else:
        flash("Expense deductibility saved.", "success")
    return redirect(url_for("firmfs.tax_rules", ye=request.form.get("ye", type=int)))


# ------------------------------------------------------------------ Excel

@firmfs_bp.route("/tax/export.xlsx")
@finance_required
def tax_export():
    got = _tax_context()
    if got is None:
        return redirect(url_for("firmfs.index"))
    stm, ye, yes, tax = got
    if tax is None:
        flash("This year-end is on manual tax: there is no tax pack to export.", "info")
        return redirect(url_for("firmfs.tax_computation", ye=ye.id))
    classes, assets = fs.load_register()
    est = firm_tax.qpd_estimate(ye, stm["cur"]["period_start"], min(date.today(), ye.period_end), classes, assets, (stm["pri"] or {}).get("tax"))
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    bold, head = Font(bold=True), PatternFill("solid", fgColor="E8E2D0")
    num = '#,##0.00;(#,##0.00);"-"'
    wb = Workbook()

    def sheet(title, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append([fs.FIRM_NAME]); ws["A1"].font = Font(bold=True, size=13)
        ws.append([title + " - " + stm["cur"]["label"]]); ws["A2"].font = Font(bold=True)
        ws.append([])
        ws.column_dimensions["A"].width = 58
        for col in "BCDEFGHIJK":
            ws.column_dimensions[col].width = 16
        return ws

    def row(ws, values, is_bold=False, header=False):
        ws.append(values)
        for c in ws[ws.max_row]:
            if isinstance(c.value, (int, float)):
                c.number_format = num
            if is_bold or header:
                c.font = bold
            if header:
                c.fill = head

    comp = tax["comp"]
    ws = sheet("Tax computation", first=True)
    row(ws, ["", "USD"], header=True)
    row(ws, ["Profit before tax", comp["profit_before_tax"]], True)
    for l in comp["addbacks"]:
        row(ws, ["Add: " + l["description"], l["amount"]])
    for l in comp["deductions"]:
        row(ws, ["Less: " + l["description"], -l["amount"]])
    row(ws, ["Taxable income before assessed loss", comp["taxable_income_before_losses"]], True)
    row(ws, ["Assessed loss brought forward", comp["tax_loss_brought_forward"]])
    row(ws, ["Less: assessed loss used", -comp["loss_utilised"]])
    row(ws, ["Taxable income", comp["taxable_income"]], True)
    row(ws, [f"Income tax at {ye.tax_rate_percent if ye.tax_rate_percent is not None else 0:g}%", tax["current_tax"]], True)
    row(ws, ["Assessed loss carried forward", tax["loss_cf"]])
    row(ws, ["Tax payable brought forward", tax["payable_bf"]])
    row(ws, ["Less: QPD payments in the period", -tax["qpd_paid"]])
    row(ws, ["Less: other income tax payments in the period", -tax["settled"]])
    row(ws, ["Income tax payable at the year-end", tax["payable"]], True)

    ws = sheet("Capital allowances")
    row(ws, ["Asset", "Class", "Method", "Tax cost", "Opening tax value", "Additions", "Allowance", "Recoupment", "Scrapping", "Closing tax value"], header=True)
    for r in tax["cap"]["rows"]:
        row(ws, [f"{r['asset_code']} {r['description']}", r["class_name"], dict(FIRM_TAX_METHODS).get(r["method"], r["method"]), r["tax_cost"],
                 r["opening"], r["additions"], r["allowance"], r["recoupment"], r["scrapping"], r["closing"]])
    t = tax["cap"]["total"]
    row(ws, ["Total", "", "", t["tax_cost"], t["opening"], t["additions"], t["allowance"], t["recoupment"], t["scrapping"], t["closing"]], True)

    ws = sheet("Deferred tax")
    row(ws, ["", "Carrying amount", "Tax base", "Temporary difference", "Deferred tax"], header=True)
    for r in tax["deferred"]["rows"]:
        row(ws, [r["description"], r["accounting_amount"], r["tax_base_amount"], r["temporary_difference"], r["deferred_tax"]])
    row(ws, ["Deferred tax liability/(asset) at the year-end", "", "", "", tax["deferred_closing"]], True)
    row(ws, ["Deferred tax liability/(asset) at the start", "", "", "", tax["deferred_opening"]])
    row(ws, ["Deferred tax charged/(credited) to profit or loss", "", "", "", tax["deferred_expense"]], True)

    ws = sheet("QPD")
    row(ws, ["Estimate as at", est["as_at"].isoformat()])
    row(ws, ["Estimated annual tax", est["annual_tax"]], True)
    row(ws, ["Instalment", "Due", "Cumulative %", "Instalment", "Cumulative due", "Paid by then", "Short"], header=True)
    for i in est["instalments"]:
        row(ws, [i["label"], i["due_date"].isoformat(), i["cumulative_pct"], i["amount"], i["cumulative_amount"], i["paid_by_then"], i["short"]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"firm_tax_{ye.period_end.isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
